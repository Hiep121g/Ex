"""Workbook model and formula engine for PyExcel.

Independent of the GUI: stores the raw text typed into each cell plus sheet
layout (formats, merged cells, row/column sizes, conditional formatting,
filters), parses and evaluates Excel-style formulas (cached, with
circular-reference detection), rewrites references when cells are copied or
rows/columns/cells are inserted or deleted, and reads/writes .pyxl (JSON),
.xlsx (via openpyxl) and .csv files.
"""

from __future__ import annotations

import base64
import calendar
import copy
import csv
import datetime as dt
import importlib.util
import io
import json
import math
import operator
import os
import random
import re
import statistics
import sys
import unicodedata
from collections import Counter, namedtuple
from decimal import ROUND_DOWN, ROUND_HALF_UP, ROUND_UP, Decimal, InvalidOperation
from fractions import Fraction

sys.setrecursionlimit(max(sys.getrecursionlimit(), 20000))

MAX_ROWS = 10000
MAX_COLS = 200
DEFAULT_COL_WIDTH = 80   # pixels at 96 DPI
DEFAULT_ROW_HEIGHT = 20  # pixels at 96 DPI
MIN_COL_WIDTH = 12
BORDER_STYLES = ("thin", "medium", "thick", "double", "dashed", "dotted")
BORDER_KEYS = ("bt", "bb", "bl", "br")  # top, bottom, left, right: [style, "#rrggbb"]
HAS_XLSX = importlib.util.find_spec("openpyxl") is not None
FILE_TAG = "PyExcel"

DIV0 = "#DIV/0!"
REF = "#REF!"
NAME = "#NAME?"
VALUE = "#VALUE!"
NA = "#N/A"
NUM = "#NUM!"
CIRC = "#CIRC!"
SYNTAX = "#ERROR!"


class CellError(Exception):
    """An Excel error value: raised to propagate through formulas, stored as a cell's value."""

    def __init__(self, code):
        super().__init__(code)
        self.code = code

    def __str__(self):
        return self.code

    def __eq__(self, other):
        return isinstance(other, CellError) and other.code == self.code

    def __hash__(self):
        return hash(self.code)


# ------------------------------------------------------------------ addresses

def col_name(c):
    """0 -> 'A', 25 -> 'Z', 26 -> 'AA'."""
    name = ""
    c += 1
    while c:
        c, rem = divmod(c - 1, 26)
        name = chr(65 + rem) + name
    return name


def col_index(name):
    n = 0
    for ch in name.upper():
        n = n * 26 + ord(ch) - 64
    return n - 1


def cell_name(r, c):
    return f"{col_name(c)}{r + 1}"


def range_name(r1, c1, r2, c2):
    return cell_name(r1, c1) if (r1, c1) == (r2, c2) else f"{cell_name(r1, c1)}:{cell_name(r2, c2)}"


def parse_range(text):
    """'A1:C4' or 'B2' -> (r1, c1, r2, c2); None when invalid."""
    pos = parse_address(text)
    if pos is None:
        return None
    return pos * 2 if len(pos) == 2 else pos


_ADDR_RE = re.compile(r"\$?([A-Za-z]{1,3})\$?(\d+)")


def rect_minus(rect, cut):
    """Parts of rectangle `rect` (r1, c1, r2, c2) outside rectangle `cut`, as up to four rectangles."""
    r1, c1, r2, c2 = rect
    a1, b1, a2, b2 = cut
    if a1 > r2 or a2 < r1 or b1 > c2 or b2 < c1:
        return [tuple(rect)]
    parts = []
    if r1 < a1:
        parts.append((r1, c1, a1 - 1, c2))
    if a2 < r2:
        parts.append((a2 + 1, c1, r2, c2))
    top, bottom = max(r1, a1), min(r2, a2)
    if c1 < b1:
        parts.append((top, c1, bottom, b1 - 1))
    if b2 < c2:
        parts.append((top, b2 + 1, bottom, c2))
    return parts


def parse_address(text):
    """'B3' -> (2, 1); 'A1:C4' -> (0, 0, 3, 2); None when invalid."""
    parts = text.strip().split(":")
    if len(parts) > 2:
        return None
    points = []
    for part in parts:
        m = _ADDR_RE.fullmatch(part.strip())
        if not m:
            return None
        r, c = int(m.group(2)) - 1, col_index(m.group(1))
        if not (0 <= r < MAX_ROWS and 0 <= c < MAX_COLS):
            return None
        points.append((r, c))
    if len(points) == 1:
        return points[0]
    (ra, ca), (rb, cb) = points
    return min(ra, rb), min(ca, cb), max(ra, rb), max(ca, cb)


_BARE_SHEET_RE = re.compile(r"[^\W\d][\w.]*")
_INVALID_SHEET_CHARS = set("[]:*?/\\")


def quote_sheet(name):
    """Sheet name as it must be written in a formula ('My Sheet'!A1)."""
    if (_BARE_SHEET_RE.fullmatch(name) and not re.fullmatch(r"[A-Za-z]{1,3}\d+", name)
            and name.upper() not in ("TRUE", "FALSE")):
        return name
    return "'" + name.replace("'", "''") + "'"


def absolute_ref(sheet, rect):
    """'Sheet1!$A$1:$B$4' (or 'Sheet1!$A$1' for one cell), as names refer to cells."""
    r1, c1, r2, c2 = rect
    first = f"${col_name(c1)}${r1 + 1}"
    return f"{quote_sheet(sheet)}!{first}" + ("" if (r1, c1) == (r2, c2) else f":${col_name(c2)}${r2 + 1}")


def name_from_label(text):
    """A defined name made from a heading, as Formulas > Create from Selection does ('Đơn giá' -> 'Đơn_giá')."""
    name = re.sub(r"[^\w.]", "", re.sub(r"\s+", "_", str(text).strip()))
    if name and not (name[0].isalpha() or name[0] == "_"):
        name = "_" + name
    if _REF_LIKE_RE.match(name) or name.upper() in ("TRUE", "FALSE"):
        name += "_"
    return name


# ------------------------------------------------------------------ values

def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def is_formula(raw):
    return isinstance(raw, str) and raw.startswith("=") and len(raw) > 1


_EPOCH = dt.datetime(1899, 12, 30)


def date_to_serial(d):
    if not isinstance(d, dt.datetime):
        d = dt.datetime(d.year, d.month, d.day)
    delta = d - _EPOCH
    return delta.days + delta.seconds / 86400 + delta.microseconds / 86400e6


def serial_to_datetime(x):
    return _EPOCH + dt.timedelta(days=float(x))


_NUM_RE = re.compile(r"[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?")
_THOUSANDS_RE = re.compile(r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?")
_DATE_DMY_RE = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")
_DATE_ISO_RE = re.compile(r"(\d{4})-(\d{1,2})-(\d{1,2})")


def parse_literal(raw):
    """Interpret typed (non-formula) text. Returns (value, implied number format)."""
    if raw.startswith("'"):
        return raw[1:], None
    s = raw.strip()
    upper = s.upper()
    if upper in ("TRUE", "FALSE"):
        return upper == "TRUE", None
    pct = s.endswith("%")
    body = s[:-1].strip() if pct else s
    if _NUM_RE.fullmatch(body) or _THOUSANDS_RE.fullmatch(body):
        v = float(body.replace(",", ""))
        if pct:
            return v / 100, ("0.00%" if "." in body else "0%")
        if "," in body:
            return v, ("#,##0.00" if "." in body else "#,##0")
        return v, None
    if not pct:
        m = _DATE_DMY_RE.fullmatch(s)
        ymd = (m.group(3), m.group(2), m.group(1)) if m else None
        if not m:
            m = _DATE_ISO_RE.fullmatch(s)
            ymd = m.groups() if m else None
        if ymd:
            try:
                return date_to_serial(dt.date(*map(int, ymd))), "dd/mm/yyyy"
            except ValueError:
                pass
    return raw, None


def number_text(v):
    """A number as text with up to 15 significant digits (Excel's value-to-text rule)."""
    if float(v).is_integer() and abs(v) < 1e15:
        return str(int(v))
    return f"{v:.15g}".replace("e+", "E+").replace("e-", "E-")


def to_number(v):
    if v is None:
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if is_number(v):
        return float(v)
    if isinstance(v, str):
        lit = parse_literal(v)[0]
        if is_number(lit):
            return lit
    raise CellError(VALUE)


def to_text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if is_number(v):
        return number_text(v)
    return str(v)


def to_bool(v):
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    if is_number(v):
        return v != 0
    s = str(v).strip().upper()
    if s in ("TRUE", "FALSE"):
        return s == "TRUE"
    raise CellError(VALUE)


_CMP_OPS = {"=": operator.eq, "<>": operator.ne, "<": operator.lt,
            ">": operator.gt, "<=": operator.le, ">=": operator.ge}


def sort_key(v):
    """Excel ordering: numbers < text (case-insensitive) < booleans < errors."""
    if isinstance(v, bool):
        return (2, v)
    if is_number(v):
        return (0, v)
    if isinstance(v, CellError):
        return (3, v.code)
    return (1, str(v).lower())


# Vietnamese alphabetical order (with the Latin letters Vietnamese lacks), and the tones in dictionary order:
# none, huyền, hỏi, ngã, sắc, nặng
_VN_LETTERS = {ch: i for i, ch in enumerate("aăâbcdđeêfghijklmnoôơpqrstuưvwxyz")}
_VN_TONES = {"\u0300": 1, "\u0309": 2, "\u0303": 3, "\u0301": 4, "\u0323": 5}
_DIGITS_RE = re.compile(r"(\d+)")


def _text_order(text):
    """(letters, tones) ordering a piece of text alphabetically the Vietnamese way: "d" < "đ" < "e",
    "a" < "ă" < "â" < "b", and tones only break ties ("anh" < "ánh" < "bình")."""
    letters, tones = [], []
    for ch in unicodedata.normalize("NFD", text.lower()):
        if ch in _VN_TONES and tones:
            tones[-1] = _VN_TONES[ch]
        elif unicodedata.combining(ch) and letters:  # breve, circumflex, horn: part of the letter (ă, â, ơ...)
            letters[-1] += ch
        else:
            letters.append(ch)
            tones.append(0)
    ranks = []
    for letter in letters:
        letter = unicodedata.normalize("NFC", letter)
        if letter in _VN_LETTERS:
            ranks.append(200 + _VN_LETTERS[letter])
        else:  # spaces and ASCII punctuation sort before letters, other scripts after them
            ranks.append(ord(letter[0]) if ord(letter[0]) < 200 else 1000 + ord(letter[0]))
    return tuple(ranks), tuple(tones)


def natural_key(v):
    """Order for sorting in the UI - sorted rows, filter lists, PivotTable items - unlike Excel's:
    numbers by value, also when typed as text ("2" before "10"), dates typed as text by date, and text
    "naturally", runs of digits by value ("SP2" before "SP10", "BCR-300x30" before "BCR-1000x50") in
    Vietnamese alphabetical order. Formulas keep Excel's rules (see compare / sort_key)."""
    if isinstance(v, bool):
        return (2, v)
    if isinstance(v, CellError):
        return (3, v.code)
    if is_number(v):
        return (0, float(v))
    text = str(v).strip()
    literal = parse_literal(text)[0] if text else ""
    if is_number(literal):
        return (0, float(literal))
    parts = tuple((0, int(chunk)) if chunk.isdigit() else (1, _text_order(chunk))
                  for chunk in _DIGITS_RE.split(text) if chunk)
    return (1, parts, text.lower())


def _blank_like(v):
    if isinstance(v, bool):
        return False
    if isinstance(v, str):
        return ""
    return 0.0


def compare(op, a, b):
    if a is None:
        a = _blank_like(b)
    if b is None:
        b = _blank_like(a)
    return _CMP_OPS[op](sort_key(a), sort_key(b))


def _dup_key(v):
    """Key under which duplicate values are counted (text compares case-insensitively)."""
    if isinstance(v, str):
        return ("s", v.lower())
    return ("b", v) if isinstance(v, bool) else ("n", float(v))


def _finite(x):
    if isinstance(x, float) and not math.isfinite(x):
        raise CellError(NUM)
    return x


def binary_op(op, a, b):
    if op == "&":
        return to_text(a) + to_text(b)
    if op in _CMP_OPS:
        return compare(op, a, b)
    x, y = to_number(a), to_number(b)
    if op == "+":
        res = x + y
    elif op == "-":
        res = x - y
    elif op == "*":
        res = x * y
    elif op == "/":
        if y == 0:
            raise CellError(DIV0)
        res = x / y
    elif op == "^":
        if x == 0 and y < 0:
            raise CellError(DIV0)
        try:
            res = x ** y
        except OverflowError:
            raise CellError(NUM) from None
        if isinstance(res, complex):
            raise CellError(NUM)
    else:
        raise CellError(SYNTAX)
    return _finite(res)


# ------------------------------------------------------------------ display formats

def format_general(v, digits=None):
    """'General' number format; `digits` limits significant digits to fit narrow columns."""
    if v == 0:
        return "0"
    if digits is None:
        if float(v).is_integer() and abs(v) < 1e11:
            return str(int(v))
        digits = 10
    s = f"{v:.{digits}g}"
    if "e" in s:
        mantissa, exp = s.split("e")
        s = f"{mantissa}E{exp[0]}{exp[1:].lstrip('0').zfill(2)}"
    return s


_NUM_CORE_RE = re.compile(r"[#0][#0,]*(?:\.[#0]+)?|\.[#0]+")
_SCI_RE = re.compile(r"[#0]+(\.0+)?E([+-])(0+)", re.I)
_DATE_TOKEN_RE = re.compile(r'"[^"]*"|yyyy|yy|dddd|ddd|dd|d|mmmm|mmm|mm|m|hh|h|ss|.', re.I | re.S)
_DATE_PARTS = {"yyyy", "yy", "dddd", "ddd", "dd", "d", "mmmm", "mmm", "mm", "m", "hh", "h", "ss"}
WEEKDAYS = ("Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật")
WEEKDAYS_SHORT = ("T2", "T3", "T4", "T5", "T6", "T7", "CN")


def is_date_format(fmt):
    low = re.sub(r'"[^"]*"', "", fmt).lower()
    return "yy" in low or "dd" in low or "hh" in low or "mmm" in low


def _round_half_up(x, digits):
    try:
        return float(Decimal(repr(float(x))).quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP))
    except InvalidOperation:
        return float(x)


def split_sections(fmt):
    """Split a number format on ';' (positive;negative;zero;text), ignoring quoted text."""
    parts, cur, quoted = [], [], False
    for ch in fmt:
        if ch == '"':
            quoted = not quoted
        if ch == ";" and not quoted:
            parts.append("".join(cur))
            cur = []
        else:
            cur.append(ch)
    parts.append("".join(cur))
    return parts


def _affix(s):
    """Literal text around the digits: drop quotes, escapes and Excel's padding codes (_x, *x)."""
    return re.sub(r"_.|\*.", "", s).replace('"', "").replace("\\", "")


def format_number(v, fmt):
    """Apply a number format such as '#,##0.00', '0%', '$#,##0', '0.00E+00', '# ?/?', 'dd/mm/yyyy'
    or a sectioned one like '#,##0;(#,##0);"-"'."""
    sections = split_sections(fmt)
    if len(sections) > 1:
        if v < 0:
            return _format_section(-v, sections[1], signed=False)
        if v == 0 and len(sections) > 2:
            return _format_section(0.0, sections[2])
    return _format_section(v, sections[0])


def _format_section(v, fmt, signed=True):
    if fmt in ("", "@") or fmt.lower() == "general":
        return format_general(v)
    if is_date_format(fmt):
        return format_date(v, fmt)
    sign = "-" if signed and v < 0 else ""
    if "?/" in fmt:
        whole, rest = int(abs(v)), abs(v) - int(abs(v))
        m = re.search(r"/(\?+|\d+)", fmt)
        frac = Fraction(rest).limit_denominator(10 ** len(m.group(1)) - 1)
        if frac == 1:
            whole, frac = whole + 1, Fraction(0)
        if frac == 0:
            return f"{sign}{whole}"
        text = f"{frac.numerator}/{frac.denominator}"
        return f"{sign}{whole} {text}" if whole else f"{sign}{text}"
    m = _SCI_RE.search(fmt)
    if m:
        decimals = len(m.group(1)) - 1 if m.group(1) else 0
        mantissa, exp = f"{abs(v):.{decimals}E}".split("E")
        digits = exp[1:].lstrip("0").zfill(len(m.group(3)))
        exp_sign = "-" if exp[0] == "-" else ("+" if m.group(2) == "+" else "")
        return f"{sign}{_affix(fmt[:m.start()])}{mantissa}E{exp_sign}{digits}{_affix(fmt[m.end():])}"
    m = _NUM_CORE_RE.search(fmt)
    if not m:
        return _affix(fmt)
    core = m.group(0)
    prefix, suffix = _affix(fmt[:m.start()]), _affix(fmt[m.end():])
    if "%" in suffix:
        v = v * 100
    decimals = len(core.split(".", 1)[1]) if "." in core else 0
    x = _round_half_up(abs(v), decimals)
    body = f"{x:,.{decimals}f}" if "," in core else f"{x:.{decimals}f}"
    if x == 0:
        sign = ""
    return f"{sign}{prefix}{body}{suffix}"


def format_date(v, fmt):
    try:
        d = serial_to_datetime(v + 1 / 172800)  # half a second, so 59.9999s doesn't show as :59
    except (OverflowError, ValueError):
        return "#" * 8
    tokens = _DATE_TOKEN_RE.findall(fmt)
    out, last = [], None
    for i, tok in enumerate(tokens):
        low = tok.lower()
        if tok.startswith('"'):
            out.append(tok[1:-1])
            continue
        if low not in _DATE_PARTS:
            out.append(tok)
            continue
        if low in ("mm", "m"):
            nxt = next((t.lower() for t in tokens[i + 1:] if t.lower() in _DATE_PARTS), None)
            val = d.minute if last in ("hh", "h") or nxt == "ss" else d.month
            out.append(f"{val:02d}" if low == "mm" else str(val))
        elif low == "dddd":
            out.append(WEEKDAYS[d.weekday()])
        elif low == "ddd":
            out.append(WEEKDAYS_SHORT[d.weekday()])
        elif low == "mmmm":
            out.append(f"tháng {d.month}")
        elif low == "mmm":
            out.append(f"thg {d.month}")
        else:
            val = {"yyyy": d.year, "yy": d.year % 100, "dd": d.day, "d": d.day,
                   "hh": d.hour, "h": d.hour, "ss": d.second}[low]
            out.append(f"{val:04d}" if low == "yyyy" else f"{val:02d}" if len(low) == 2 else str(val))
        last = low
    return "".join(out)


def format_value(value, fmt=None):
    """Text shown for a cell value (without column-width fitting)."""
    if value is None:
        return ""
    if isinstance(value, CellError):
        return value.code
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if is_number(value):
        return format_number(value, fmt) if fmt else format_general(value)
    return str(value)


def adjust_decimals(fmt, value, delta):
    """Number format with one decimal place more (delta=1) or fewer (delta=-1)."""
    if fmt in (None, "@"):
        s = format_general(value) if is_number(value) else "0"
        d = len(s.split(".")[1]) if "." in s and "E" not in s else 0
        fmt = "0" + ("." + "0" * d if d else "")
    out = []
    for section in split_sections(fmt):
        m = None if is_date_format(section) or "?/" in section else _NUM_CORE_RE.search(section)
        if m is None:
            out.append(section)
            continue
        int_part, _, dec = m.group(0).partition(".")
        nd = max(0, min(15, len(dec) + delta))
        out.append(section[:m.start()] + (int_part or "0") + ("." + "0" * nd if nd else "") + section[m.end():])
    return ";".join(out)


def value_to_raw(v):
    """Raw cell text that reproduces a computed value (used by 'paste values')."""
    if v is None or v == "":
        return None
    if isinstance(v, CellError):
        return v.code
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if is_number(v):
        return number_text(v)
    s = str(v)
    if s.startswith(("=", "'")) or parse_literal(s) != (s, None):
        return "'" + s
    return s


def blend(c1, c2, t):
    """Colour between '#rrggbb' c1 (t=0) and c2 (t=1)."""
    a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
    return "#%02x%02x%02x" % tuple(round(x + (y - x) * t) for x, y in zip(a, b))


# ------------------------------------------------------------------ tokenizer

_SHEET_PAT = r"(?:'(?:[^']|'')+'|[^\W\d][\w.]*)"
_TOKEN_RE = re.compile(rf"""
    (?P<ws>\s+)
  | (?P<string>"(?:[^"]|"")*")
  | (?P<error>\#(?:DIV/0!|REF!|NAME\?|VALUE!|N/A|NUM!|CIRC!|ERROR!))
  | (?P<ref>(?:(?P<sheet>{_SHEET_PAT})!)?(?P<cabs>\$?)(?P<col>[A-Za-z]{{1,3}})(?P<rabs>\$?)(?P<row>\d+)(?![\w(]))
  | (?P<number>(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?)
  | (?P<ident>[^\W\d][\w.]*)
  | (?P<op><>|<=|>=|[-+*/^&=<>%:(),;])
""", re.VERBOSE)
_TOKEN_KINDS = ("ws", "string", "error", "ref", "number", "ident", "op")

RefParts = namedtuple("RefParts", "sheet_text sheet cabs col rabs row")


class Token:
    __slots__ = ("kind", "text", "ref")

    def __init__(self, kind, text, ref=None):
        self.kind, self.text, self.ref = kind, text, ref


def tokenize(text):
    tokens, pos = [], 0
    while pos < len(text):
        m = _TOKEN_RE.match(text, pos)
        if m is None:
            raise CellError(SYNTAX)
        kind = next(k for k in _TOKEN_KINDS if m.group(k) is not None)
        ref = None
        if kind == "ref":
            st = m.group("sheet")
            sheet = (st[1:-1].replace("''", "'") if st.startswith("'") else st) if st else None
            ref = RefParts(st, sheet, bool(m.group("cabs")), col_index(m.group("col")),
                           bool(m.group("rabs")), int(m.group("row")) - 1)
        tokens.append(Token(kind, m.group(0), ref))
        pos = m.end()
    return tokens


def _format_ref(p):
    prefix = f"{p.sheet_text}!" if p.sheet_text else ""
    return f"{prefix}{'$' if p.cabs else ''}{col_name(p.col)}{'$' if p.rabs else ''}{p.row + 1}"


def _ref_items(tokens):
    """Group tokens into ('tok', t), ('ref', t) and ('range', t1, t2) items."""
    i, n = 0, len(tokens)
    while i < n:
        t = tokens[i]
        if (t.kind == "ref" and i + 2 < n and tokens[i + 1].kind == "op"
                and tokens[i + 1].text == ":" and tokens[i + 2].kind == "ref"):
            yield ("range", t, tokens[i + 2])
            i += 3
        else:
            yield ("ref" if t.kind == "ref" else "tok", t)
            i += 1


# ------------------------------------------------------------------ parser

class _Parser:
    """Recursive-descent parser producing tuple-based ASTs (precedence as in Excel)."""

    def __init__(self, tokens):
        self.toks = [t for t in tokens if t.kind != "ws"]
        self.pos = 0

    def parse(self):
        if not self.toks:
            raise CellError(SYNTAX)
        node = self.comparison()
        if self.pos != len(self.toks):
            raise CellError(SYNTAX)
        return node

    def _at(self, *ops):
        t = self.toks[self.pos] if self.pos < len(self.toks) else None
        return t is not None and t.kind == "op" and t.text in ops

    def _expect(self, op):
        if not self._at(op):
            raise CellError(SYNTAX)
        self.pos += 1

    def _binary(self, ops, operand):
        node = operand()
        while self._at(*ops):
            op = self.toks[self.pos].text
            self.pos += 1
            node = ("bin", op, node, operand())
        return node

    def comparison(self):
        return self._binary(("=", "<>", "<", ">", "<=", ">="), self.concat)

    def concat(self):
        return self._binary(("&",), self.additive)

    def additive(self):
        return self._binary(("+", "-"), self.multiplicative)

    def multiplicative(self):
        return self._binary(("*", "/"), self.power)

    def power(self):
        return self._binary(("^",), self.unary)

    def unary(self):
        if self._at("-"):
            self.pos += 1
            return ("neg", self.unary())
        if self._at("+"):
            self.pos += 1
            return self.unary()
        node = self.primary()
        while self._at("%"):
            self.pos += 1
            node = ("pct", node)
        return node

    def primary(self):
        if self.pos >= len(self.toks):
            raise CellError(SYNTAX)
        t = self.toks[self.pos]
        self.pos += 1
        if t.kind == "number":
            return ("num", float(t.text))
        if t.kind == "string":
            return ("str", t.text[1:-1].replace('""', '"'))
        if t.kind == "error":
            return ("err", t.text.upper())
        if t.kind == "ref":
            p = t.ref
            if self._at(":"):
                self.pos += 1
                if self.pos >= len(self.toks) or self.toks[self.pos].kind != "ref":
                    raise CellError(SYNTAX)
                q = self.toks[self.pos].ref
                self.pos += 1
                return ("range", p.sheet, min(p.row, q.row), min(p.col, q.col),
                        max(p.row, q.row), max(p.col, q.col))
            return ("ref", p.sheet, p.row, p.col)
        if t.kind == "ident":
            name = t.text.upper()
            if self._at("("):
                self.pos += 1
                return ("func", name, self._args())
            if name in ("TRUE", "FALSE"):
                return ("bool", name == "TRUE")
            return ("name", name)  # a defined name, looked up when evaluated
        if t.kind == "op" and t.text == "(":
            node = self.comparison()
            self._expect(")")
            return node
        raise CellError(SYNTAX)

    def _args(self):
        args = []
        if self._at(")"):
            self.pos += 1
            return args
        while True:
            args.append(("empty",) if self._at(",", ";", ")") else self.comparison())
            if self._at(",", ";"):
                self.pos += 1
                continue
            self._expect(")")
            return args


def parse_formula(raw):
    """AST for a formula string ('=...'); syntax errors become an error node."""
    try:
        return _Parser(tokenize(raw[1:])).parse()
    except CellError as err:
        return ("err", err.code)


# ------------------------------------------------------------------ formula text rewriting

def normalize_formula(raw, names=None):
    """Upper-case functions and references, write defined names as they were defined (`names`: the
    workbook's names), use ',' separators and close open parentheses."""
    if not is_formula(raw):
        return raw
    try:
        tokens = tokenize(raw[1:])
    except CellError:
        return raw
    out, depth = [], 0
    for i, t in enumerate(tokens):
        if t.kind == "ident":
            after = next((n for n in tokens[i + 1:] if n.kind != "ws"), None)
            called = after is not None and after.kind == "op" and after.text == "("
            upper = t.text.upper()
            if not called and names and upper in names:
                out.append(names[upper]["name"])
            else:
                out.append(upper if called or upper in ("TRUE", "FALSE") or not names else t.text)
        elif t.kind == "ref":
            out.append(_format_ref(t.ref))
        elif t.kind == "op" and t.text == ";":
            out.append(",")
        else:
            out.append(t.text)
            if t.kind == "op":
                depth += {"(": 1, ")": -1}.get(t.text, 0)
    return "=" + "".join(out) + ")" * max(depth, 0)


def shift_formula(raw, dr, dc):
    """Move relative references when a formula is copied dr rows / dc columns away."""
    if not is_formula(raw) or (dr == 0 and dc == 0):
        return raw
    try:
        tokens = tokenize(raw[1:])
    except CellError:
        return raw

    def moved(p):
        row = p.row if p.rabs else p.row + dr
        col = p.col if p.cabs else p.col + dc
        if not (0 <= row < MAX_ROWS and 0 <= col < MAX_COLS):
            return None
        return p._replace(row=row, col=col)

    out = []
    for item in _ref_items(tokens):
        if item[0] == "tok":
            out.append(item[1].text)
        elif item[0] == "ref":
            p = moved(item[1].ref)
            out.append(_format_ref(p) if p else REF)
        else:
            p, q = moved(item[1].ref), moved(item[2].ref)
            out.append(f"{_format_ref(p)}:{_format_ref(q)}" if p and q else REF)
    return "=" + "".join(out)


def formula_references(text):
    """(sheet or None, r1, c1, r2, c2) for each reference in a formula being typed."""
    if not text.startswith("="):
        return []
    try:
        tokens = tokenize(text[1:])
    except CellError:
        return []
    refs = []
    for item in _ref_items(tokens):
        if item[0] == "ref":
            p = item[1].ref
            refs.append((p.sheet, p.row, p.col, p.row, p.col))
        elif item[0] == "range":
            p, q = item[1].ref, item[2].ref
            refs.append((p.sheet, min(p.row, q.row), min(p.col, q.col), max(p.row, q.row), max(p.col, q.col)))
    return refs


def _rewrite_axis(raw, home, target, axis, move, move_span, band=None):
    """Rewrite references into sheet `target` after rows (axis 0) / columns (axis 1) moved.

    With `band` = (lo, hi) only cells whose other coordinate lies in the band
    moved (inserting/deleting cells instead of whole rows/columns)."""
    try:
        tokens = tokenize(raw[1:])
    except CellError:
        return raw

    def idx(p):
        return p.row if axis == 0 else p.col

    def in_band(p):
        other = p.col if axis == 0 else p.row
        return band is None or band[0] <= other <= band[1]

    def put(p, i):
        return p._replace(row=i) if axis == 0 else p._replace(col=i)

    out = []
    for item in _ref_items(tokens):
        if item[0] == "tok":
            out.append(item[1].text)
            continue
        p = item[1].ref
        q = item[2].ref if item[0] == "range" else p
        if (p.sheet or home).lower() != target or not (in_band(p) and in_band(q)):
            out.append(item[1].text if item[0] == "ref" else f"{item[1].text}:{item[2].text}")
        elif item[0] == "ref":
            n = move(idx(p))
            out.append(_format_ref(put(p, n)) if n is not None else REF)
        else:
            span = move_span(min(idx(p), idx(q)), max(idx(p), idx(q)))
            out.append(REF if span is None else
                       f"{_format_ref(put(p, span[0]))}:{_format_ref(put(q, span[1]))}")
    return "=" + "".join(out)


def _rename_in_formula(raw, old, new):
    try:
        tokens = tokenize(raw[1:])
    except CellError:
        return raw
    out = []
    for t in tokens:
        if t.kind == "ref" and t.ref.sheet and t.ref.sheet.lower() == old.lower():
            out.append(_format_ref(t.ref._replace(sheet_text=quote_sheet(new), sheet=new)))
        else:
            out.append(t.text)
    return "=" + "".join(out)


def _number_raw(v, fmt):
    if fmt and is_date_format(fmt):
        return format_date(v, "dd/mm/yyyy")
    if fmt and "%" in fmt:
        return number_text(round(v * 100, 10)) + "%"
    return number_text(v)


def fill_series(raws):
    """For auto-fill: a function k -> raw text continuing the source cells, or None to copy them.

    k is the offset from the first source cell. Two or more numbers continue
    their step, a single date steps by one day, and text ending in a number
    ("Item 1") counts up.
    """
    if not raws or any(r is None or is_formula(r) for r in raws):
        return None
    parsed = [parse_literal(r) for r in raws]
    values = [v for v, _ in parsed]
    fmt = parsed[0][1]
    if all(is_number(v) for v in values):
        if len(values) == 1:
            if not (fmt and is_date_format(fmt)):
                return None
            step = 1.0
        else:
            step = (values[-1] - values[0]) / (len(values) - 1)
        first = values[0]
        return lambda k: _number_raw(first + step * k, fmt)
    if len(raws) == 1 and isinstance(values[0], str):
        m = re.fullmatch(r"(.*?)(\d+)(\D*)", raws[0], re.S)
        if m:
            prefix, digits, suffix = m.groups()
            return lambda k: f"{prefix}{max(0, int(digits) + k):0{len(digits)}d}{suffix}"
    return None


# ------------------------------------------------------------------ ranges

class RangeRef:
    """A rectangular block of cells, handed to functions without being evaluated."""

    __slots__ = ("wb", "sheet", "r1", "c1", "r2", "c2")

    def __init__(self, wb, sheet, r1, c1, r2, c2):
        self.wb, self.sheet = wb, sheet
        self.r1, self.c1, self.r2, self.c2 = r1, c1, r2, c2

    @property
    def nrows(self):
        return self.r2 - self.r1 + 1

    @property
    def ncols(self):
        return self.c2 - self.c1 + 1

    def value(self, i, j):
        r, c = self.r1 + i, self.c1 + j
        if r >= MAX_ROWS or c >= MAX_COLS:
            return None
        return self.wb.get_value(self.sheet, r, c)

    def rows(self):
        return [[self.value(i, j) for j in range(self.ncols)] for i in range(self.nrows)]

    def flat(self):
        return [v for row in self.rows() for v in row]

    def stored(self):
        """(row, col) of the cells holding something, row by row; on large sparse ranges only
        stored cells are visited."""
        cells = self.wb.sheets[self.sheet].cells
        if self.nrows * self.ncols <= len(cells):
            return [(r, c) for r in range(self.r1, min(self.r2 + 1, MAX_ROWS))
                    for c in range(self.c1, min(self.c2 + 1, MAX_COLS)) if (r, c) in cells]
        return sorted(k for k in cells if self.r1 <= k[0] <= self.r2 and self.c1 <= k[1] <= self.c2)

    def values(self):
        """Non-empty values row by row."""
        for r, c in self.stored():
            v = self.wb.get_value(self.sheet, r, c)
            if v is not None:
                yield v

    def line(self):
        """Cells of a one-row or one-column range, in order; None for a 2-D block."""
        if self.ncols == 1:
            return [self.value(i, 0) for i in range(self.nrows)]
        if self.nrows == 1:
            return [self.value(0, j) for j in range(self.ncols)]
        return None

    def part(self, i, by_row):
        """Row i (by_row) or column i of the range, as a range."""
        if by_row:
            return RangeRef(self.wb, self.sheet, self.r1 + i, self.c1, self.r1 + i, self.c2)
        return RangeRef(self.wb, self.sheet, self.r1, self.c1 + i, self.r2, self.c1 + i)


def scalar(x):
    """Single value of an argument: a one-cell range is dereferenced, larger ones are an error."""
    if isinstance(x, RangeRef):
        if x.nrows == 1 and x.ncols == 1:
            v = x.value(0, 0)
            if isinstance(v, CellError):
                raise CellError(v.code)
            return v
        raise CellError(VALUE)
    return x


# ------------------------------------------------------------------ functions

_FUNCS = {}
_LAZY_FUNCS = {}  # take (workbook, argument nodes, sheet) and evaluate only what they need

_FUNC_FORMATS = {"TODAY": "dd/mm/yyyy", "DATE": "dd/mm/yyyy", "NOW": "dd/mm/yyyy hh:mm", "EDATE": "dd/mm/yyyy",
                 "EOMONTH": "dd/mm/yyyy", "WORKDAY": "dd/mm/yyyy", "DATEVALUE": "dd/mm/yyyy", "TIME": "hh:mm"}


def _register(*names, lazy=False):
    def deco(func):
        for name in names:
            (_LAZY_FUNCS if lazy else _FUNCS)[name] = func
        return func
    return deco


def function_names():
    return sorted(set(_FUNCS) | set(_LAZY_FUNCS))


def _num(x):
    return to_number(scalar(x))


def _int(x):
    return int(_num(x))


def _text(x):
    return to_text(scalar(x))


def _bool(x):
    return to_bool(scalar(x))


def _range_arg(x):
    if not isinstance(x, RangeRef):
        raise CellError(VALUE)
    return x


def _raise_if_error(v):
    if isinstance(v, CellError):
        raise CellError(v.code)
    return v


def _peek(x):
    """Value of a one-cell argument without raising the error it may hold."""
    if isinstance(x, RangeRef) and x.nrows == 1 and x.ncols == 1:
        return x.value(0, 0)
    return x


def _numbers(args):
    """Numbers for SUM/AVERAGE/...: ranges contribute numeric cells only, literals are coerced."""
    for a in args:
        if isinstance(a, RangeRef):
            for v in a.values():
                if is_number(_raise_if_error(v)):
                    yield float(v)
        elif a is not None:
            yield to_number(a)


def _flat(args):
    for a in args:
        if isinstance(a, RangeRef):
            for v in a.flat():
                yield _raise_if_error(v)
        else:
            yield a


def _quantize(x, digits, mode):
    try:
        return float(Decimal(repr(float(x))).quantize(Decimal(1).scaleb(-digits), rounding=mode))
    except InvalidOperation:
        return float(x)


# math and statistics
@_register("SUM")
def _f_sum(*args):
    return math.fsum(_numbers(args))


@_register("AVERAGE")
def _f_average(*args):
    nums = list(_numbers(args))
    if not nums:
        raise CellError(DIV0)
    return math.fsum(nums) / len(nums)


@_register("MIN")
def _f_min(*args):
    nums = list(_numbers(args))
    return min(nums) if nums else 0.0


@_register("MAX")
def _f_max(*args):
    nums = list(_numbers(args))
    return max(nums) if nums else 0.0


@_register("PRODUCT")
def _f_product(*args):
    nums = list(_numbers(args))
    return math.prod(nums) if nums else 0.0


@_register("MEDIAN")
def _f_median(*args):
    nums = list(_numbers(args))
    if not nums:
        raise CellError(NUM)
    return float(statistics.median(nums))


@_register("STDEV", "STDEV.S")
def _f_stdev(*args):
    nums = list(_numbers(args))
    if len(nums) < 2:
        raise CellError(DIV0)
    return statistics.stdev(nums)


@_register("VAR", "VAR.S")
def _f_var(*args):
    nums = list(_numbers(args))
    if len(nums) < 2:
        raise CellError(DIV0)
    return statistics.variance(nums)


@_register("STDEVP", "STDEV.P")
def _f_stdevp(*args):
    nums = list(_numbers(args))
    if not nums:
        raise CellError(DIV0)
    return statistics.pstdev(nums)


@_register("VARP", "VAR.P")
def _f_varp(*args):
    nums = list(_numbers(args))
    if not nums:
        raise CellError(DIV0)
    return statistics.pvariance(nums)


def _kth(args, k, largest):
    nums = sorted(_numbers([args]), reverse=largest)
    k = math.ceil(_num(k))
    if not 1 <= k <= len(nums):
        raise CellError(NUM)
    return nums[k - 1]


@_register("LARGE")
def _f_large(array, k):
    return _kth(array, k, True)


@_register("SMALL")
def _f_small(array, k):
    return _kth(array, k, False)


def _rank(x, ref, order, average):
    v = _num(x)
    nums = list(_numbers([_range_arg(ref)]))
    if v not in nums:
        raise CellError(NA)
    ascending = order is not None and _num(order) != 0
    better = sum(1 for n in nums if (n < v if ascending else n > v))
    if not average:
        return float(better + 1)
    ties = nums.count(v)
    return better + (ties + 1) / 2


@_register("RANK", "RANK.EQ")
def _f_rank(x, ref, order=None):
    return _rank(x, ref, order, False)


@_register("RANK.AVG")
def _f_rank_avg(x, ref, order=None):
    return _rank(x, ref, order, True)


@_register("MODE", "MODE.SNGL")
def _f_mode(*args):
    nums = list(_numbers(args))
    counts = Counter(nums)
    best = max(counts.values(), default=0)
    if best < 2:
        raise CellError(NA)
    return next(n for n in nums if counts[n] == best)  # the first one met wins ties, as in Excel


@_register("PERCENTILE", "PERCENTILE.INC")
def _f_percentile(array, k):
    nums = sorted(_numbers([array]))
    k = _num(k)
    if not nums or not 0 <= k <= 1:
        raise CellError(NUM)
    pos = k * (len(nums) - 1)
    lo = math.floor(pos)
    hi = min(lo + 1, len(nums) - 1)
    return nums[lo] + (nums[hi] - nums[lo]) * (pos - lo)


@_register("QUARTILE", "QUARTILE.INC")
def _f_quartile(array, quart):
    q = math.trunc(_num(quart))
    if not 0 <= q <= 4:
        raise CellError(NUM)
    return _f_percentile(array, q / 4)


@_register("COUNT")
def _f_count(*args):
    n = 0
    for a in args:
        if isinstance(a, RangeRef):
            n += sum(1 for v in a.values() if is_number(v))
        elif is_number(a) or isinstance(a, bool):
            n += 1
    return n


@_register("COUNTA")
def _f_counta(*args):
    return sum(sum(1 for _ in a.values()) if isinstance(a, RangeRef) else int(a is not None) for a in args)


@_register("COUNTBLANK")
def _f_countblank(rng):
    rng = _range_arg(rng)
    return rng.nrows * rng.ncols - sum(1 for v in rng.values() if v != "")


@_register("ROUND")
def _f_round(x, digits=0):
    return _quantize(_num(x), _int(digits), ROUND_HALF_UP)


@_register("ROUNDUP")
def _f_roundup(x, digits=0):
    return _quantize(_num(x), _int(digits), ROUND_UP)


@_register("ROUNDDOWN", "TRUNC")
def _f_rounddown(x, digits=0):
    return _quantize(_num(x), _int(digits), ROUND_DOWN)


@_register("INT")
def _f_int(x):
    return float(math.floor(_num(x)))


@_register("MOD")
def _f_mod(a, b):
    a, b = _num(a), _num(b)
    if b == 0:
        raise CellError(DIV0)
    return a - b * math.floor(a / b)


@_register("ABS")
def _f_abs(x):
    return abs(_num(x))


@_register("SIGN")
def _f_sign(x):
    v = _num(x)
    return float((v > 0) - (v < 0))


@_register("SQRT")
def _f_sqrt(x):
    v = _num(x)
    if v < 0:
        raise CellError(NUM)
    return math.sqrt(v)


@_register("POWER")
def _f_power(a, b):
    return binary_op("^", _num(a), _num(b))


@_register("EXP")
def _f_exp(x):
    return math.exp(_num(x))


@_register("LN")
def _f_ln(x):
    v = _num(x)
    if v <= 0:
        raise CellError(NUM)
    return math.log(v)


@_register("LOG")
def _f_log(x, base=10):
    v, b = _num(x), _num(base)
    if v <= 0 or b <= 0 or b == 1:
        raise CellError(NUM)
    return math.log(v, b)


@_register("LOG10")
def _f_log10(x):
    return _f_log(x, 10)


@_register("PI")
def _f_pi():
    return math.pi


@_register("CEILING")
def _f_ceiling(x, significance=1):
    v, s = _num(x), _num(significance)
    return 0.0 if s == 0 else math.ceil(v / s) * s


@_register("FLOOR")
def _f_floor(x, significance=1):
    v, s = _num(x), _num(significance)
    if s == 0:
        raise CellError(DIV0)
    return math.floor(v / s) * s


@_register("RAND")
def _f_rand():
    return random.random()


@_register("RANDBETWEEN")
def _f_randbetween(a, b):
    lo, hi = math.ceil(_num(a)), math.floor(_num(b))
    if lo > hi:
        raise CellError(NUM)
    return float(random.randint(lo, hi))


@_register("SUMPRODUCT")
def _f_sumproduct(*ranges):
    if not ranges:
        raise CellError(VALUE)
    ranges = [_range_arg(r) for r in ranges]
    shape = (ranges[0].nrows, ranges[0].ncols)
    if any((r.nrows, r.ncols) != shape for r in ranges):
        raise CellError(VALUE)
    grids = [r.rows() for r in ranges]
    total = 0.0
    for i in range(shape[0]):
        for j in range(shape[1]):
            prod = 1.0
            for g in grids:
                v = _raise_if_error(g[i][j])
                prod *= v if is_number(v) else 0.0
            total += prod
    return total


@_register("SUMSQ")
def _f_sumsq(*args):
    return math.fsum(x * x for x in _numbers(args))


@_register("MROUND")
def _f_mround(x, multiple):
    v, m = _num(x), _num(multiple)
    if m == 0:
        return 0.0
    if v * m < 0:
        raise CellError(NUM)
    return _quantize(v / m, 0, ROUND_HALF_UP) * m


@_register("QUOTIENT")
def _f_quotient(a, b):
    a, b = _num(a), _num(b)
    if b == 0:
        raise CellError(DIV0)
    return float(math.trunc(a / b))


def _whole_numbers(args):
    nums = [math.trunc(x) for x in _numbers(args)]
    if any(x < 0 for x in nums):
        raise CellError(NUM)
    return nums


@_register("GCD")
def _f_gcd(*args):
    return float(math.gcd(*_whole_numbers(args)))


@_register("LCM")
def _f_lcm(*args):
    return float(math.lcm(*_whole_numbers(args)))


@_register("FACT")
def _f_fact(x):
    n = math.trunc(_num(x))
    if n < 0 or n > 170:
        raise CellError(NUM)
    return float(math.factorial(n))


@_register("COMBIN")
def _f_combin(n, k):
    n, k = math.trunc(_num(n)), math.trunc(_num(k))
    if n < 0 or k < 0 or k > n:
        raise CellError(NUM)
    return float(math.comb(n, k))


def _away_from_zero(v, parity):
    """Round away from zero to the nearest even (parity 0) or odd (parity 1) integer."""
    n = math.ceil(abs(v))
    if n % 2 != parity:
        n += 1
    return float(math.copysign(n, v)) if v else float(parity)


@_register("EVEN")
def _f_even(x):
    return _away_from_zero(_num(x), 0)


@_register("ODD")
def _f_odd(x):
    return _away_from_zero(_num(x), 1)


def _unary_math(name, func):
    _register(name)(lambda x: func(_num(x)))


for _name, _func in (("SIN", math.sin), ("COS", math.cos), ("TAN", math.tan), ("ASIN", math.asin),
                     ("ACOS", math.acos), ("ATAN", math.atan), ("DEGREES", math.degrees),
                     ("RADIANS", math.radians)):
    _unary_math(_name, _func)


@_register("ATAN2")
def _f_atan2(x, y):
    x, y = _num(x), _num(y)
    if x == 0 and y == 0:
        raise CellError(DIV0)
    return math.atan2(y, x)  # Excel takes x first


_SUBTOTALS = {1: "AVERAGE", 2: "COUNT", 3: "COUNTA", 4: "MAX", 5: "MIN", 6: "PRODUCT", 7: "STDEV",
              8: "STDEVP", 9: "SUM", 10: "VAR", 11: "VARP"}
_SUBTOTAL_RE = re.compile(r"=\s*SUBTOTAL\s*\(", re.I)


@_register("SUBTOTAL")
def _f_subtotal(kind, *refs):
    """Aggregate that skips rows hidden by a filter (and, for codes 101-111, by hand) and
    cells that are themselves SUBTOTALs."""
    k = _int(kind)
    name = _SUBTOTALS.get(k - 100 if k > 100 else k)
    if name is None or not refs:
        raise CellError(VALUE)
    values = []
    for ref in refs:
        rng = _range_arg(ref)
        sh = rng.wb.sheets[rng.sheet]
        filtered = sh.filter["hidden"] if sh.filter else ()
        for r, c in rng.stored():
            if r in filtered or (k > 100 and r in sh.hidden_rows) or _SUBTOTAL_RE.match(sh.cells[(r, c)]):
                continue
            values.append(_raise_if_error(rng.wb.get_value(rng.sheet, r, c)))
    if name == "COUNTA":
        return len(values)
    nums = [float(v) for v in values if is_number(v)]
    if name == "COUNT":
        return len(nums)
    return _FUNCS[name](*nums)


# logic
def _logicals(args):
    for a in args:
        if isinstance(a, RangeRef):
            for v in a.values():
                v = _raise_if_error(v)
                if isinstance(v, bool):
                    yield v
                elif is_number(v):
                    yield v != 0
        elif a is not None:
            yield to_bool(a)


@_register("AND")
def _f_and(*args):
    vals = list(_logicals(args))
    if not vals:
        raise CellError(VALUE)
    return all(vals)


@_register("OR")
def _f_or(*args):
    vals = list(_logicals(args))
    if not vals:
        raise CellError(VALUE)
    return any(vals)


@_register("XOR")
def _f_xor(*args):
    vals = list(_logicals(args))
    if not vals:
        raise CellError(VALUE)
    return sum(vals) % 2 == 1


@_register("NOT")
def _f_not(x):
    return not _bool(x)


@_register("TRUE")
def _f_true():
    return True


@_register("FALSE")
def _f_false():
    return False


@_register("IF", lazy=True)
def _f_if(wb, nodes, sheet):
    if len(nodes) not in (2, 3):
        raise CellError(VALUE)
    if to_bool(scalar(wb._eval(nodes[0], sheet))):
        return wb._eval(nodes[1], sheet)
    return wb._eval(nodes[2], sheet) if len(nodes) == 3 else False


@_register("IFS", lazy=True)
def _f_ifs(wb, nodes, sheet):
    if not nodes or len(nodes) % 2:
        raise CellError(VALUE)
    for i in range(0, len(nodes), 2):
        if to_bool(scalar(wb._eval(nodes[i], sheet))):
            return wb._eval(nodes[i + 1], sheet)
    raise CellError(NA)


@_register("SWITCH", lazy=True)
def _f_switch(wb, nodes, sheet):
    if len(nodes) < 3:
        raise CellError(VALUE)
    value = scalar(wb._eval(nodes[0], sheet))
    pairs = nodes[1:]
    for i in range(0, len(pairs) - 1, 2):
        if compare("=", value, scalar(wb._eval(pairs[i], sheet))):
            return wb._eval(pairs[i + 1], sheet)
    if len(pairs) % 2:
        return wb._eval(pairs[-1], sheet)
    raise CellError(NA)


@_register("IFERROR", lazy=True)
def _f_iferror(wb, nodes, sheet):
    if len(nodes) != 2:
        raise CellError(VALUE)
    try:
        return scalar(wb._eval(nodes[0], sheet))
    except CellError:
        return wb._eval(nodes[1], sheet)


@_register("IFNA", lazy=True)
def _f_ifna(wb, nodes, sheet):
    if len(nodes) != 2:
        raise CellError(VALUE)
    try:
        return scalar(wb._eval(nodes[0], sheet))
    except CellError as err:
        if err.code != NA:
            raise
        return wb._eval(nodes[1], sheet)


def _error_code(wb, nodes, sheet):
    if len(nodes) != 1:
        raise CellError(VALUE)
    try:
        scalar(wb._eval(nodes[0], sheet))
    except CellError as err:
        return err.code
    return None


@_register("ISERROR", lazy=True)
def _f_iserror(wb, nodes, sheet):
    return _error_code(wb, nodes, sheet) is not None


@_register("ISERR", lazy=True)
def _f_iserr(wb, nodes, sheet):
    return _error_code(wb, nodes, sheet) not in (None, NA)


@_register("ISNA", lazy=True)
def _f_isna(wb, nodes, sheet):
    return _error_code(wb, nodes, sheet) == NA


@_register("ISBLANK")
def _f_isblank(x):
    return _peek(x) is None


@_register("ISNUMBER")
def _f_isnumber(x):
    return is_number(_peek(x))


@_register("ISTEXT")
def _f_istext(x):
    return isinstance(_peek(x), str)


@_register("ISLOGICAL")
def _f_islogical(x):
    return isinstance(_peek(x), bool)


@_register("ISNONTEXT")
def _f_isnontext(x):
    return not isinstance(_peek(x), str)


@_register("ISEVEN")
def _f_iseven(x):
    return math.trunc(_num(x)) % 2 == 0


@_register("ISODD")
def _f_isodd(x):
    return math.trunc(_num(x)) % 2 == 1


@_register("HYPERLINK")
def _f_hyperlink(link, name=None):
    """Shows `name` (or the link); clicking the cell opens the link (see hyperlink_at)."""
    return scalar(link) if name is None else scalar(name)


@_register("ISFORMULA")
def _f_isformula(ref):
    rng = _range_arg(ref)
    return is_formula(rng.wb.get_raw(rng.sheet, rng.r1, rng.c1))


@_register("NA")
def _f_na():
    raise CellError(NA)


@_register("ROW", lazy=True)
def _f_row(wb, nodes, sheet):
    if not nodes:
        return float(wb._stack[-1][1] + 1)
    ref = wb._eval(nodes[0], sheet)
    return float(_range_arg(ref).r1 + 1)


@_register("COLUMN", lazy=True)
def _f_column(wb, nodes, sheet):
    if not nodes:
        return float(wb._stack[-1][2] + 1)
    ref = wb._eval(nodes[0], sheet)
    return float(_range_arg(ref).c1 + 1)


@_register("ROWS")
def _f_rows(rng):
    return _range_arg(rng).nrows


@_register("COLUMNS")
def _f_columns(rng):
    return _range_arg(rng).ncols


@_register("CHOOSE")
def _f_choose(index, *values):
    i = _int(index)
    if not 1 <= i <= len(values):
        raise CellError(VALUE)
    return values[i - 1]


# text
@_register("CONCAT", "CONCATENATE")
def _f_concat(*args):
    return "".join(to_text(v) for v in _flat(args))


@_register("TEXTJOIN")
def _f_textjoin(delimiter, ignore_empty, *args):
    parts = [to_text(v) for v in _flat(args)]
    if _bool(ignore_empty):
        parts = [p for p in parts if p != ""]
    return _text(delimiter).join(parts)


@_register("LEN")
def _f_len(x):
    return len(_text(x))


@_register("UPPER")
def _f_upper(x):
    return _text(x).upper()


@_register("LOWER")
def _f_lower(x):
    return _text(x).lower()


@_register("PROPER")
def _f_proper(x):
    return _text(x).title()


@_register("TRIM")
def _f_trim(x):
    return " ".join(_text(x).split())


@_register("LEFT")
def _f_left(x, n=1):
    n = _int(n)
    if n < 0:
        raise CellError(VALUE)
    return _text(x)[:n]


@_register("RIGHT")
def _f_right(x, n=1):
    n, s = _int(n), _text(x)
    if n < 0:
        raise CellError(VALUE)
    return s[max(0, len(s) - n):]


@_register("MID")
def _f_mid(x, start, n):
    start, n = _int(start), _int(n)
    if start < 1 or n < 0:
        raise CellError(VALUE)
    return _text(x)[start - 1:start - 1 + n]


@_register("SUBSTITUTE")
def _f_substitute(x, old, new, instance=None):
    s, old, new = _text(x), _text(old), _text(new)
    if not old:
        return s
    if instance is None:
        return s.replace(old, new)
    k = _int(instance)
    if k < 1:
        raise CellError(VALUE)
    pos = -1
    for _ in range(k):
        pos = s.find(old, pos + 1)
        if pos < 0:
            return s
    return s[:pos] + new + s[pos + len(old):]


@_register("REPLACE")
def _f_replace(x, start, n, new):
    s, start, n = _text(x), _int(start), _int(n)
    if start < 1 or n < 0:
        raise CellError(VALUE)
    return s[:start - 1] + _text(new) + s[start - 1 + n:]


@_register("REPT")
def _f_rept(x, n):
    n, s = _int(n), _text(x)
    if n < 0 or len(s) * n > 32767:
        raise CellError(VALUE)
    return s * n


@_register("FIND")
def _f_find(sub, x, start=1):
    i = _text(x).find(_text(sub), _int(start) - 1)
    if i < 0 or _int(start) < 1:
        raise CellError(VALUE)
    return i + 1


@_register("SEARCH")
def _f_search(sub, x, start=1):
    i = _text(x).lower().find(_text(sub).lower(), _int(start) - 1)
    if i < 0 or _int(start) < 1:
        raise CellError(VALUE)
    return i + 1


@_register("EXACT")
def _f_exact(a, b):
    return _text(a) == _text(b)


@_register("VALUE")
def _f_value(x):
    return to_number(scalar(x))


@_register("TEXT")
def _f_text(x, fmt):
    v = scalar(x)
    try:
        n = to_number(v)
    except CellError:
        return to_text(v)
    return format_number(n, _text(fmt))


@_register("CHAR")
def _f_char(x):
    n = _int(x)
    if not 1 <= n <= 0x10FFFF:
        raise CellError(VALUE)
    return chr(n)


@_register("CODE")
def _f_code(x):
    s = _text(x)
    if not s:
        raise CellError(VALUE)
    return ord(s[0])


def _split_at(x, delimiter, instance):
    """Text before and after the n-th delimiter (counting from the end when n < 0)."""
    s, d, n = _text(x), _text(delimiter), _int(instance)
    if n == 0 or not d:
        raise CellError(VALUE)
    pos, end = -1, len(s)
    for _ in range(abs(n)):
        if n > 0:
            pos = s.find(d, pos + 1)
        else:
            pos = s.rfind(d, 0, end)
            end = pos + len(d) - 1
        if pos < 0:
            raise CellError(NA)
    return s[:pos], s[pos + len(d):]


@_register("TEXTBEFORE")
def _f_textbefore(x, delimiter, instance=1):
    return _split_at(x, delimiter, instance)[0]


@_register("TEXTAFTER")
def _f_textafter(x, delimiter, instance=1):
    return _split_at(x, delimiter, instance)[1]


# dates
@_register("TODAY")
def _f_today():
    return date_to_serial(dt.date.today())


@_register("NOW")
def _f_now():
    return date_to_serial(dt.datetime.now())


@_register("DATE")
def _f_date(y, m, d):
    y, m, d = _int(y), _int(m), _int(d)
    if 0 <= y < 1900:
        y += 1900
    y += (m - 1) // 12
    m = (m - 1) % 12 + 1
    return date_to_serial(dt.date(y, m, 1) + dt.timedelta(days=d - 1))


@_register("YEAR")
def _f_year(x):
    return serial_to_datetime(_num(x)).year


@_register("MONTH")
def _f_month(x):
    return serial_to_datetime(_num(x)).month


@_register("DAY")
def _f_day(x):
    return serial_to_datetime(_num(x)).day


@_register("HOUR")
def _f_hour(x):
    return serial_to_datetime(_num(x) + 1 / 172800).hour


@_register("MINUTE")
def _f_minute(x):
    return serial_to_datetime(_num(x) + 1 / 172800).minute


@_register("WEEKDAY")
def _f_weekday(x, kind=1):
    wd = serial_to_datetime(_num(x)).weekday()  # Monday = 0
    return wd + 1 if _int(kind) == 2 else (wd + 1) % 7 + 1


@_register("DAYS")
def _f_days(end, start):
    return math.floor(_num(end)) - math.floor(_num(start))


@_register("SECOND")
def _f_second(x):
    return serial_to_datetime(_num(x) + 1 / 172800).second


@_register("TIME")
def _f_time(h, m, s):
    total = _int(h) * 3600 + _int(m) * 60 + _int(s)
    if total < 0:
        raise CellError(NUM)
    return (total % 86400) / 86400


@_register("DATEVALUE")
def _f_datevalue(x):
    value, fmt = parse_literal(_text(x))
    if not (fmt and is_date_format(fmt)):
        raise CellError(VALUE)
    return float(math.floor(value))


def _date_of(x):
    return serial_to_datetime(math.floor(_num(x))).date()


def _plus_months(d, months):
    """(year, month) `months` after the month of date d."""
    y, m = divmod(d.month - 1 + months, 12)
    if not 1 <= d.year + y <= 9999:
        raise CellError(NUM)
    return d.year + y, m + 1


@_register("EDATE")
def _f_edate(start, months):
    d = _date_of(start)
    y, m = _plus_months(d, math.trunc(_num(months)))
    return date_to_serial(dt.date(y, m, min(d.day, calendar.monthrange(y, m)[1])))


@_register("EOMONTH")
def _f_eomonth(start, months):
    y, m = _plus_months(_date_of(start), math.trunc(_num(months)))
    return date_to_serial(dt.date(y, m, calendar.monthrange(y, m)[1]))


def _is_weekend(serial):
    return (serial + 5) % 7 >= 5  # serial 0 (30/12/1899) was a Saturday


def _holidays(arg):
    return set() if arg is None else {math.floor(v) for v in _numbers([arg])}


@_register("NETWORKDAYS")
def _f_networkdays(start, end, holidays=None):
    a, b = math.floor(_num(start)), math.floor(_num(end))
    sign = 1 if a <= b else -1
    a, b = min(a, b), max(a, b)
    weeks, rest = divmod(b - a + 1, 7)
    count = weeks * 5 + sum(1 for d in range(b - rest + 1, b + 1) if not _is_weekend(d))
    count -= sum(1 for d in _holidays(holidays) if a <= d <= b and not _is_weekend(d))
    return float(sign * count)


@_register("WORKDAY")
def _f_workday(start, days, holidays=None):
    d, n = math.floor(_num(start)), math.trunc(_num(days))
    if abs(n) > 1_000_000:
        raise CellError(NUM)
    skip, step = _holidays(holidays), 1 if n >= 0 else -1
    while n:
        d += step
        if not _is_weekend(d) and d not in skip:
            n -= step
    return float(d)


@_register("DATEDIF")
def _f_datedif(start, end, unit):
    a, b, unit = _date_of(start), _date_of(end), _text(unit).upper()
    if a > b:
        raise CellError(NUM)
    months = (b.year - a.year) * 12 + b.month - a.month - (b.day < a.day)

    def shifted(d, n):  # d moved n months on, the day clamped to the month's length
        y, m = _plus_months(d, n)
        return dt.date(y, m, min(d.day, calendar.monthrange(y, m)[1]))

    if unit == "D":
        return float((b - a).days)
    if unit == "M":
        return float(months)
    if unit == "Y":
        return float(months // 12)
    if unit == "YM":
        return float(months % 12)
    if unit == "MD":
        return float((b - shifted(a, months)).days)
    if unit == "YD":
        return float((b - shifted(a, months // 12 * 12)).days)
    raise CellError(NUM)


@_register("WEEKNUM")
def _f_weeknum(x, kind=1):
    d, kind = _date_of(x), _int(kind)
    if kind not in (1, 2):
        raise CellError(NUM)
    jan1 = dt.date(d.year, 1, 1)
    offset = (jan1.weekday() + 1) % 7 if kind == 1 else jan1.weekday()  # days of week 1 before Jan 1
    return float((d.timetuple().tm_yday - 1 + offset) // 7 + 1)


@_register("ISOWEEKNUM")
def _f_isoweeknum(x):
    return float(_date_of(x).isocalendar()[1])


# lookup
def _match_index(items, key, mode):
    """Position of `key` in items: mode 0 exact, 1 largest <= key, -1 smallest >= key."""
    if key is None:
        raise CellError(NA)
    if mode == 0:
        for i, v in enumerate(items):
            if v is not None and not isinstance(v, CellError) and compare("=", v, key):
                return i
        raise CellError(NA)
    rank, best = sort_key(key)[0], None
    for i, v in enumerate(items):
        if v is None or isinstance(v, CellError) or sort_key(v)[0] != rank:
            continue
        if compare("<=" if mode > 0 else ">=", v, key):
            best = i
        else:
            break
    if best is None:
        raise CellError(NA)
    return best


@_register("VLOOKUP")
def _f_vlookup(key, table, col, approx=True):
    table, ci = _range_arg(table), _int(col)
    if ci < 1:
        raise CellError(VALUE)
    if ci > table.ncols:
        raise CellError(REF)
    i = _match_index([table.value(r, 0) for r in range(table.nrows)], scalar(key), 1 if _bool(approx) else 0)
    return _raise_if_error(table.value(i, ci - 1))


@_register("HLOOKUP")
def _f_hlookup(key, table, row, approx=True):
    table, ri = _range_arg(table), _int(row)
    if ri < 1:
        raise CellError(VALUE)
    if ri > table.nrows:
        raise CellError(REF)
    j = _match_index([table.value(0, c) for c in range(table.ncols)], scalar(key), 1 if _bool(approx) else 0)
    return _raise_if_error(table.value(ri - 1, j))


@_register("INDEX")
def _f_index(table, row, col=None):
    table, r = _range_arg(table), _int(row)
    if col is None:
        r, c = (1, r) if table.nrows == 1 else (r, 1)
    else:
        c = _int(col)
    if not (1 <= r <= table.nrows and 1 <= c <= table.ncols):
        raise CellError(REF)
    return _raise_if_error(table.value(r - 1, c - 1))


@_register("MATCH")
def _f_match(key, rng, mode=1):
    rng = _range_arg(rng)
    if rng.nrows != 1 and rng.ncols != 1:
        raise CellError(NA)
    m = _int(mode)
    return _match_index(rng.flat(), scalar(key), (m > 0) - (m < 0)) + 1


def _wildcard_re(text):
    """Regex for an Excel pattern where * is any run of characters and ? any one character."""
    return re.compile("".join(".*" if ch == "*" else "." if ch == "?" else re.escape(ch) for ch in text), re.S)


def _xmatch_index(items, key, match_mode, search_mode):
    """Position of `key` for XLOOKUP/XMATCH: match_mode 0 exact, -1 exact or next smaller,
    1 exact or next larger, 2 wildcard; a negative search_mode searches from the end."""
    if match_mode not in (-1, 0, 1, 2) or search_mode not in (-2, -1, 1, 2):
        raise CellError(VALUE)
    if key is None:
        raise CellError(NA)
    order = range(len(items)) if search_mode > 0 else range(len(items) - 1, -1, -1)
    if match_mode == 2:
        pattern = _wildcard_re(to_text(key).lower())
        for i in order:
            if isinstance(items[i], str) and pattern.fullmatch(items[i].lower()):
                return i
        raise CellError(NA)
    rank, best = sort_key(key)[0], None
    for i in order:
        v = items[i]
        if v is None or isinstance(v, CellError):
            continue
        if compare("=", v, key):
            return i
        if (match_mode and sort_key(v)[0] == rank and compare("<" if match_mode < 0 else ">", v, key)
                and (best is None or compare(">" if match_mode < 0 else "<", v, items[best]))):
            best = i
    if best is None:
        raise CellError(NA)
    return best


def _optional_int(wb, nodes, i, sheet, default):
    if len(nodes) <= i or nodes[i][0] == "empty":
        return default
    return int(to_number(scalar(wb._eval(nodes[i], sheet))))


@_register("XLOOKUP", lazy=True)
def _f_xlookup(wb, nodes, sheet):
    if not 3 <= len(nodes) <= 6:
        raise CellError(VALUE)
    key = scalar(wb._eval(nodes[0], sheet))
    look = _range_arg(wb._eval(nodes[1], sheet))
    result = _range_arg(wb._eval(nodes[2], sheet))
    items = look.line()
    by_row = look.ncols == 1
    if items is None or (result.nrows if by_row else result.ncols) != len(items):
        raise CellError(VALUE)
    match_mode = _optional_int(wb, nodes, 4, sheet, 0)
    search_mode = _optional_int(wb, nodes, 5, sheet, 1)
    try:
        i = _xmatch_index(items, key, match_mode, search_mode)
    except CellError as err:
        if err.code != NA or len(nodes) < 4 or nodes[3][0] == "empty":
            raise
        return wb._eval(nodes[3], sheet)
    return result.part(i, by_row)


@_register("XMATCH")
def _f_xmatch(key, rng, match_mode=None, search_mode=None):
    items = _range_arg(rng).line()
    if items is None:
        raise CellError(VALUE)
    match_mode = 0 if match_mode is None else _int(match_mode)
    search_mode = 1 if search_mode is None else _int(search_mode)
    return _xmatch_index(items, scalar(key), match_mode, search_mode) + 1


@_register("OFFSET")
def _f_offset(ref, rows, cols, height=None, width=None):
    ref = _range_arg(ref)
    r1, c1 = ref.r1 + _int(rows), ref.c1 + _int(cols)
    h = ref.nrows if height is None else _int(height)
    w = ref.ncols if width is None else _int(width)
    if h == 0 or w == 0:
        raise CellError(REF)
    if h < 0:  # a negative height or width reaches up / left
        r1, h = r1 + h + 1, -h
    if w < 0:
        c1, w = c1 + w + 1, -w
    if r1 < 0 or c1 < 0 or r1 + h > MAX_ROWS or c1 + w > MAX_COLS:
        raise CellError(REF)
    return RangeRef(ref.wb, ref.sheet, r1, c1, r1 + h - 1, c1 + w - 1)


_R1C1_RE = re.compile(r"R(\d+)C(\d+)(?::R(\d+)C(\d+))?", re.I)


@_register("INDIRECT", lazy=True)
def _f_indirect(wb, nodes, sheet):
    if len(nodes) not in (1, 2):
        raise CellError(VALUE)
    text = to_text(scalar(wb._eval(nodes[0], sheet))).strip()
    a1 = len(nodes) < 2 or nodes[1][0] == "empty" or to_bool(scalar(wb._eval(nodes[1], sheet)))
    sheet_ref = None
    if "!" in text:
        sheet_ref, text = text.rsplit("!", 1)
        if len(sheet_ref) >= 2 and sheet_ref[0] == sheet_ref[-1] == "'":
            sheet_ref = sheet_ref[1:-1].replace("''", "'")
    if a1:
        rect = parse_range(text)
    else:
        m = _R1C1_RE.fullmatch(text)
        if m:
            ra, ca = int(m.group(1)) - 1, int(m.group(2)) - 1
            rb, cb = (int(m.group(3)) - 1, int(m.group(4)) - 1) if m.group(3) else (ra, ca)
            rect = (min(ra, rb), min(ca, cb), max(ra, rb), max(ca, cb))
        else:
            rect = None
    if rect is None:
        raise CellError(REF)
    return wb._range(sheet_ref, sheet, *rect)


# conditional aggregation
def _criteria(crit):
    """Predicate for COUNTIF/SUMIF-style criteria such as 5, ">5", "<>x" or "a*"."""
    crit = scalar(crit)
    op, target = "=", crit
    if isinstance(crit, str):
        m = re.match(r"(<=|>=|<>|<|>|=)?(.*)", crit, re.S)
        op = m.group(1) or "="
        target = parse_literal(m.group(2))[0] if m.group(2) else ""
    if target is None:
        target = ""
    cmp = _CMP_OPS[op]
    if is_number(target) or isinstance(target, bool):
        want_bool = isinstance(target, bool)

        def pred(v):
            if not (is_number(v) or isinstance(v, bool)) or isinstance(v, bool) != want_bool:
                return op == "<>"
            return cmp(v, target)
        return pred
    t = target.lower()
    pattern = None
    if op in ("=", "<>") and ("*" in t or "?" in t):
        pattern = _wildcard_re(t)

    def pred(v):
        if op in ("=", "<>"):
            if v is None:
                s = ""
            elif isinstance(v, str):
                s = v.lower()
            else:
                return op == "<>"
            hit = pattern.fullmatch(s) is not None if pattern else s == t
            return hit == (op == "=")
        return isinstance(v, str) and cmp(v.lower(), t)
    return pred


def _aligned(base, rng):
    """`rng` resized to the shape of `base`, keeping its top-left corner (as Excel does)."""
    return RangeRef(rng.wb, rng.sheet, rng.r1, rng.c1, rng.r1 + base.nrows - 1, rng.c1 + base.ncols - 1)


def _matching_cells(pairs):
    """(first range, positions) where every (range, criteria) pair matches."""
    if not pairs or len(pairs) % 2:
        raise CellError(VALUE)
    ranges = [_range_arg(pairs[i]) for i in range(0, len(pairs), 2)]
    preds = [_criteria(pairs[i]) for i in range(1, len(pairs), 2)]
    shape = (ranges[0].nrows, ranges[0].ncols)
    if any((r.nrows, r.ncols) != shape for r in ranges):
        raise CellError(VALUE)
    hits = [(i, j) for i in range(shape[0]) for j in range(shape[1])
            if all(p(r.value(i, j)) for r, p in zip(ranges, preds))]
    return ranges[0], hits


def _numbers_at(rng, positions):
    for i, j in positions:
        v = _raise_if_error(rng.value(i, j))
        if is_number(v):
            yield v


@_register("COUNTIF")
def _f_countif(rng, crit):
    return len(_matching_cells((rng, crit))[1])


@_register("COUNTIFS")
def _f_countifs(*pairs):
    return len(_matching_cells(pairs)[1])


@_register("SUMIF")
def _f_sumif(rng, crit, sum_rng=None):
    base, hits = _matching_cells((rng, crit))
    target = _aligned(base, _range_arg(sum_rng)) if sum_rng is not None else base
    return math.fsum(_numbers_at(target, hits))


@_register("SUMIFS")
def _f_sumifs(sum_rng, *pairs):
    base, hits = _matching_cells(pairs)
    return math.fsum(_numbers_at(_aligned(base, _range_arg(sum_rng)), hits))


@_register("AVERAGEIF")
def _f_averageif(rng, crit, avg_rng=None):
    base, hits = _matching_cells((rng, crit))
    target = _aligned(base, _range_arg(avg_rng)) if avg_rng is not None else base
    nums = list(_numbers_at(target, hits))
    if not nums:
        raise CellError(DIV0)
    return math.fsum(nums) / len(nums)


@_register("AVERAGEIFS")
def _f_averageifs(avg_rng, *pairs):
    base, hits = _matching_cells(pairs)
    nums = list(_numbers_at(_aligned(base, _range_arg(avg_rng)), hits))
    if not nums:
        raise CellError(DIV0)
    return math.fsum(nums) / len(nums)


@_register("MAXIFS")
def _f_maxifs(max_rng, *pairs):
    base, hits = _matching_cells(pairs)
    return max(_numbers_at(_aligned(base, _range_arg(max_rng)), hits), default=0.0)


@_register("MINIFS")
def _f_minifs(min_rng, *pairs):
    base, hits = _matching_cells(pairs)
    return min(_numbers_at(_aligned(base, _range_arg(min_rng)), hits), default=0.0)


# ------------------------------------------------------------------ AutoFilter criteria

# A filter column's criterion is either a set of displayed texts to keep (the checklist) or a dict:
#   {"type": "custom", "join": "and" | "or", "conds": [[op, text], ...]}   op in FILTER_OPS
#   {"type": "top", "n": 10, "bottom": bool, "percent": bool}
#   {"type": "average", "above": bool}
#   {"type": "dynamic", "kind": one of DYNAMIC_FILTERS}                      dates relative to today
#   {"type": "color", "what": "bg" | "fg", "color": "#rrggbb" | None}         None: no fill / automatic
FILTER_OPS = ("=", "<>", ">", ">=", "<", "<=", "begins", "nbegins", "ends", "nends", "contains", "ncontains")
DYNAMIC_FILTERS = ("tomorrow", "today", "yesterday", "nextweek", "thisweek", "lastweek", "nextmonth", "thismonth",
                   "lastmonth", "nextyear", "thisyear", "lastyear")
_TEXT_OPS = {"begins": "{}*", "ends": "*{}", "contains": "*{}*"}


def filter_condition(op, target, value, text):
    """Whether a cell (its value and displayed text) meets one custom AutoFilter condition. Text
    comparisons ignore case and take Excel's wildcards; a number or date target compares numbers."""
    target = str(target)
    negate = op.startswith("n") and op[1:] in _TEXT_OPS
    base = op[1:] if negate else op
    if base in _TEXT_OPS:
        return (_wildcard_re(_TEXT_OPS[base].format(target.lower())).fullmatch(text.lower()) is not None) != negate
    goal = parse_literal(target)[0] if target.strip() else ""
    if op in ("=", "<>"):
        if is_number(goal):
            hit = is_number(value) and abs(value - goal) < 1e-9
        elif isinstance(goal, str) and ("*" in goal or "?" in goal):
            hit = _wildcard_re(goal.lower()).fullmatch(text.lower()) is not None
        else:
            hit = text.lower() == target.strip().lower()
        return hit == (op == "=")
    if is_number(goal):
        return is_number(value) and _CMP_OPS[op](value, goal)
    return isinstance(value, str) and value != "" and _CMP_OPS[op](value.lower(), str(goal).lower())


def dynamic_date_span(kind, today):
    """(first, last) date that a relative date filter such as "thismonth" keeps."""
    step = {"last": -1, "this": 0, "next": 1}
    if kind in ("today", "yesterday", "tomorrow"):
        day = today + dt.timedelta(days={"yesterday": -1, "today": 0, "tomorrow": 1}[kind])
        return day, day
    when, unit = kind[:4], kind[4:]
    n = step[when]
    if unit == "week":  # Sunday to Saturday, as in Excel
        start = today - dt.timedelta(days=(today.weekday() + 1) % 7) + dt.timedelta(weeks=n)
        return start, start + dt.timedelta(days=6)
    if unit == "month":
        y, m = divmod(today.month - 1 + n, 12)
        year, month = today.year + y, m + 1
        return dt.date(year, month, 1), dt.date(year, month, calendar.monthrange(year, month)[1])
    return dt.date(today.year + n, 1, 1), dt.date(today.year + n, 12, 31)


# ------------------------------------------------------------------ workbook

CHART_TYPES = ("column", "bar", "line", "pie")
PICTURE = "picture"  # a picture floating over the cells, kept with the charts


def new_picture_id():
    return os.urandom(6).hex()
CHART_SIZE = (480, 288)  # Excel's default 5 x 3 inches
VALIDATION_TYPES = ("any", "whole", "decimal", "list", "date", "length", "custom")
VALIDATION_OPS = ("between", "notbetween", "=", "<>", ">", "<", ">=", "<=")


def _shift_freeze(freeze, axis, at, move_span):
    """Frozen panes after rows (axis 0) or columns (axis 1) were inserted or deleted at `at`."""
    a0, a1 = freeze[0:2] if axis == 0 else freeze[2:4]
    if a1 > a0:
        span = move_span(a0, a1 - 1)
        if span is None:
            a0 = a1 = 0
        else:  # lines inserted right at the top of the frozen block join it, as in Excel
            a0, a1 = (a0 if at == a0 < span[0] else span[0]), span[1] + 1
    elif a0:
        span = move_span(a0, a0)
        a0 = a1 = span[0] if span else 0
    freeze = (a0, a1) + tuple(freeze[2:4]) if axis == 0 else tuple(freeze[0:2]) + (a0, a1)
    return freeze if freeze[1] > freeze[0] or freeze[3] > freeze[2] else None


MAX_OUTLINE = 7  # outline levels Excel allows
_DEFINED_NAME_RE = re.compile(r"[^\W\d][\w.]*\Z")
_REF_LIKE_RE = re.compile(r"(?:[A-Za-z]{1,3}\d+|[RrCc]\d*|[Rr]\d+[Cc]\d+)\Z")
# SUBTOTAL function numbers -> (label after each group's value, label of the grand total row)
SUBTOTAL_NAMES = {9: ("Tổng", "Tổng Cuối"), 3: ("Đếm", "Đếm chung"), 1: ("Trung bình", "Trung bình chung"),
                  4: ("Lớn nhất", "Lớn nhất chung"), 5: ("Nhỏ nhất", "Nhỏ nhất chung"), 6: ("Tích", "Tích chung"),
                  2: ("Đếm số", "Đếm số chung"), 7: ("Độ lệch chuẩn", "Độ lệch chuẩn chung"),
                  8: ("Độ lệch chuẩn TT", "Độ lệch chuẩn TT chung"), 10: ("Phương sai", "Phương sai chung"),
                  11: ("Phương sai TT", "Phương sai TT chung")}


def outline_groups(levels):
    """[(level, first, last)] of the groups an outline {index: level} makes, outer levels first."""
    groups = []
    for level in range(1, max(levels.values(), default=0) + 1):
        run = None
        for i in sorted(levels):
            if levels[i] < level:
                continue
            if run and i == run[1] + 1:
                run[1] = i
            else:
                if run:
                    groups.append((level, *run))
                run = [i, i]
        if run:
            groups.append((level, *run))
    return groups


class Sheet:
    """One worksheet. Format dicts are never mutated in place, only replaced."""

    _STATE = ("cells", "formats", "col_widths", "row_heights", "custom_rows", "hidden_rows",
              "hidden_cols", "merges", "cond_formats", "filter", "tab_color", "freeze", "validations", "notes",
              "charts", "pivots", "page_setup", "row_levels", "col_levels", "collapsed_rows", "collapsed_cols",
              "links")
    __slots__ = ("name", "_merge_map") + _STATE

    def __init__(self, name):
        self.name = name
        self.cells = {}         # (row, col) -> raw text as typed
        self.formats = {}       # (row, col) -> {"font", "size", "bold", "italic", "underline", "strike",
                                #   "fg", "bg", "align", "valign", "wrap", "indent", "rotate", "num",
                                #   "bt", "bb", "bl", "br"}
        self.col_widths = {}    # col -> width in pixels at 96 DPI
        self.row_heights = {}   # row -> height in pixels at 96 DPI (rows that differ from the default)
        self.custom_rows = set()  # rows sized by hand, which automatic fitting leaves alone
        self.hidden_rows = set()
        self.hidden_cols = set()
        self.merges = []        # [(r1, c1, r2, c2)]
        self.cond_formats = []  # [{"range": (r1, c1, r2, c2), "type": ..., ...}] - later rules win
        self.filter = None      # {"range": (r1, c1, r2, c2), "criteria": {col: set(texts)}, "hidden": set(rows)}
        self.tab_color = None
        self.freeze = None      # (r0, r1, c0, c1): rows r0..r1-1 and columns c0..c1-1 stay put when scrolling
        self.validations = []   # [{"range": (r1, c1, r2, c2), "type": ..., ...}] - later rules win, see VALIDATION_TYPES
        self.notes = {}         # (row, col) -> note text
        self.links = {}         # (row, col) -> {"target": URL, file, mailto: or "#Sheet!A1", "tip": ScreenTip}
        self.charts = []        # [{"type": CHART_TYPES[i], "range": (r1, c1, r2, c2), "title": str,
                                #   "anchor": (row, col, dx, dy), "size": (w, h)}] - offsets and size in pixels at 96 DPI;
                                #   pictures float the same way: {"type": PICTURE, "image": base64, "format": "png",
                                #   "id": str, "name": str, "anchor", "size"} (no range)
        self.pivots = []        # PivotTables whose reports sit on this sheet (see pivot.new_pivot)
        self.page_setup = None  # printing options that differ from printing.DEFAULT_SETUP
        self.row_levels = {}    # row -> outline (group) level 1..MAX_OUTLINE; summary rows sit below groups
        self.col_levels = {}    # col -> outline level; summary columns sit right of groups
        self.collapsed_rows = set()  # summary rows of collapsed groups (Excel's "collapsed" flag)
        self.collapsed_cols = set()
        self._merge_map = None

    def merge_at(self, r, c):
        """The merged block containing (r, c), or None."""
        if not self.merges:
            return None
        if self._merge_map is None:
            self._merge_map = {}
            for m in self.merges:
                for rr in range(m[0], m[2] + 1):
                    for cc in range(m[1], m[3] + 1):
                        self._merge_map[(rr, cc)] = m
        return self._merge_map.get((r, c))

    def set_merges(self, merges):
        self.merges = [tuple(m) for m in merges]
        self._merge_map = None

    def row_hidden(self, r):
        return r in self.hidden_rows or (self.filter is not None and r in self.filter["hidden"])

    def copy_state(self):
        return {key: copy.deepcopy(getattr(self, key)) for key in self._STATE}

    @classmethod
    def from_state(cls, name, state):
        sheet = cls(name)
        for key, value in state.items():
            setattr(sheet, key, copy.deepcopy(value))
        sheet.set_merges(sheet.merges)
        return sheet


class Workbook:
    def __init__(self, sheet_names=("Sheet1",)):
        self.sheets = {}
        self.active = None
        self._cache = {}       # (sheet, r, c) -> computed value
        self._ast = {}         # (sheet, r, c) -> parsed formula
        self._cf_cache = {}    # (sheet, r, c) -> conditional-format style
        self._cf_stats = {}    # (sheet, rule index, name) -> statistics of a rule's range
        self._evaluating = set()
        self._stack = []
        self.names = {}        # NAME (upper case) -> {"name": as typed, "ref": "=Sheet1!$A$1:$B$4", "comment": ""}
        self._names_ast = {}
        for name in sheet_names:
            self.add_sheet(name)

    # sheets
    def resolve_sheet(self, name):
        if name in self.sheets:
            return name
        low = name.lower()
        return next((n for n in self.sheets if n.lower() == low), None)

    def check_sheet_name(self, name, ignore=None):
        """Error message (shown in the UI) when `name` can't be used, else None."""
        if not name.strip():
            return "Tên trang tính không được để trống."
        if len(name) > 31:
            return "Tên trang tính tối đa 31 ký tự."
        if any(ch in _INVALID_SHEET_CHARS for ch in name) or name.startswith("'") or name.endswith("'"):
            return "Tên trang tính không được chứa các ký tự [ ] : * ? / \\ hoặc bắt đầu/kết thúc bằng dấu '."
        other = self.resolve_sheet(name)
        if other is not None and other != ignore:
            return "Đã có trang tính với tên này."
        return None

    def add_sheet(self, name=None, index=None):
        if name is None:
            i = len(self.sheets) + 1
            while self.resolve_sheet(f"Sheet{i}"):
                i += 1
            name = f"Sheet{i}"
        err = self.check_sheet_name(name)
        if err:
            raise ValueError(err)
        items = list(self.sheets.items())
        items.insert(len(items) if index is None else index, (name, Sheet(name)))
        self.sheets = dict(items)
        self.invalidate()
        return name

    def rename_sheet(self, old, new):
        err = self.check_sheet_name(new, ignore=old)
        if err:
            raise ValueError(err)
        for sh in self.sheets.values():
            for key, raw in list(sh.cells.items()):
                if is_formula(raw):
                    sh.cells[key] = _rename_in_formula(raw, old, new)
        for d in self.names.values():
            d["ref"] = _rename_in_formula(d["ref"], old, new)
        self._names_ast.clear()
        self.sheets = {(new if n == old else n): s for n, s in self.sheets.items()}
        self.sheets[new].name = new
        for sh in self.sheets.values():
            for p in sh.pivots:
                if p["source"]["sheet"].lower() == old.lower():
                    p["source"]["sheet"] = new
        self._ast.clear()
        self.invalidate()

    def delete_sheet(self, name):
        if len(self.sheets) <= 1:
            raise ValueError("Sổ làm việc phải có ít nhất một trang tính.")
        del self.sheets[name]
        self._ast.clear()
        self.invalidate()

    def move_sheet(self, name, index):
        items = [(n, s) for n, s in self.sheets.items() if n != name]
        items.insert(max(0, min(index, len(items))), (name, self.sheets[name]))
        self.sheets = dict(items)

    def copy_sheet(self, name, index=None):
        """Duplicate a sheet as 'Name (2)'; returns the new name."""
        i = 2
        while self.resolve_sheet(f"{name[:26]} ({i})"):
            i += 1
        new = self.add_sheet(f"{name[:26]} ({i})", index)
        self.sheets[new] = Sheet.from_state(new, self.sheets[name].copy_state())
        self.invalidate()
        return new

    # cells
    def get_raw(self, sheet, r, c):
        return self.sheets[sheet].cells.get((r, c), "")

    def set_raw(self, sheet, r, c, raw):
        cells = self.sheets[sheet].cells
        if raw:
            cells[(r, c)] = raw
        else:
            cells.pop((r, c), None)
        self._ast.pop((sheet, r, c), None)
        self.invalidate()

    def get_format(self, sheet, r, c):
        return self.sheets[sheet].formats.get((r, c))

    def set_format(self, sheet, r, c, fmt):
        formats = self.sheets[sheet].formats
        old = formats.get((r, c)) or {}
        if fmt:
            formats[(r, c)] = dict(fmt)
        else:
            formats.pop((r, c), None)
        if "@" in (old.get("num"), (fmt or {}).get("num")) and old.get("num") != (fmt or {}).get("num"):
            self.invalidate()  # the Text format changes how the cell's content is read

    def invalidate(self):
        self._cache.clear()
        self._cf_cache.clear()
        self._cf_stats.clear()

    def used_extent(self, sheet):
        """(rows, cols) spanned by non-empty cells."""
        cells = self.sheets[sheet].cells
        if not cells:
            return 0, 0
        return max(r for r, _ in cells) + 1, max(c for _, c in cells) + 1

    # evaluation
    def _parsed(self, key, raw):
        ast = self._ast.get(key)
        if ast is None:
            ast = self._ast[key] = parse_formula(raw)
        return ast

    def get_value(self, sheet, r, c):
        key = (sheet, r, c)
        try:
            return self._cache[key]
        except KeyError:
            pass
        sh = self.sheets[sheet]
        raw = sh.cells.get((r, c))
        if raw is None:
            return None
        fmt = sh.formats.get((r, c))
        if fmt and fmt.get("num") == "@":  # Text format: keep exactly what was typed
            value = raw[1:] if raw.startswith("'") else raw
        elif not is_formula(raw):
            value = parse_literal(raw)[0]
        else:
            if key in self._evaluating:
                raise CellError(CIRC)
            self._evaluating.add(key)
            self._stack.append(key)
            try:
                value = scalar(self._eval(self._parsed(key, raw), sheet))
                if value is None:
                    value = 0.0
            except CellError as err:
                value = CellError(err.code)
            except RecursionError:
                value = CellError(REF)
            finally:
                self._evaluating.discard(key)
                self._stack.pop()
        self._cache[key] = value
        return value

    def cell_display(self, sheet, r, c):
        """(value, number format) used to draw a cell."""
        value = self.get_value(sheet, r, c)
        fmt = self.sheets[sheet].formats.get((r, c))
        numfmt = fmt.get("num") if fmt else None
        if numfmt is None and is_number(value):
            numfmt = self.implied_format(sheet, r, c)
        return value, numfmt

    def implied_format(self, sheet, r, c):
        """Format suggested by the content itself, e.g. a typed date or =TODAY()."""
        raw = self.sheets[sheet].cells.get((r, c))
        if raw is None:
            return None
        if not is_formula(raw):
            return parse_literal(raw)[1]
        ast = self._parsed((sheet, r, c), raw)
        return _FUNC_FORMATS.get(ast[1]) if ast[0] == "func" else None

    def _range(self, sheet_ref, home, r1, c1, r2, c2):
        name = home if sheet_ref is None else self.resolve_sheet(sheet_ref)
        if name is None or r1 < 0 or c1 < 0 or r2 >= MAX_ROWS or c2 >= MAX_COLS:
            raise CellError(REF)
        return RangeRef(self, name, r1, c1, r2, c2)

    def _eval(self, node, sheet):
        kind = node[0]
        if kind in ("num", "str", "bool"):
            return node[1]
        if kind == "ref":
            return self._range(node[1], sheet, node[2], node[3], node[2], node[3])
        if kind == "range":
            return self._range(node[1], sheet, *node[2:])
        if kind == "bin":
            return binary_op(node[1], scalar(self._eval(node[2], sheet)), scalar(self._eval(node[3], sheet)))
        if kind == "neg":
            return -to_number(scalar(self._eval(node[1], sheet)))
        if kind == "pct":
            return to_number(scalar(self._eval(node[1], sheet))) / 100
        if kind == "func":
            return self._call(node[1], node[2], sheet)
        if kind == "name":
            return self._eval(self._name_ast(node[1]), sheet)
        if kind == "err":
            raise CellError(node[1])
        if kind == "empty":
            return None
        raise CellError(SYNTAX)

    def _name_ast(self, name):
        ast = self._names_ast.get(name)
        if ast is None:
            d = self.names.get(name)
            ast = self._names_ast[name] = parse_formula(d["ref"]) if d else ("err", NAME)
        return ast

    def _call(self, name, nodes, sheet):
        lazy = _LAZY_FUNCS.get(name)
        if lazy is not None:
            return lazy(self, nodes, sheet)
        func = _FUNCS.get(name)
        if func is None:
            raise CellError(NAME)
        args = [self._eval(n, sheet) for n in nodes]
        try:
            result = func(*args)
        except CellError:
            raise
        except TypeError:  # wrong number of arguments
            raise CellError(VALUE) from None
        except ZeroDivisionError:
            raise CellError(DIV0) from None
        except (ValueError, OverflowError, InvalidOperation, statistics.StatisticsError):
            raise CellError(NUM) from None
        return _finite(result)

    def evaluate(self, text, sheet):
        """Value of a literal or '=formula' outside any cell (conditional-format operands)."""
        if not is_formula(text):
            return parse_literal(text)[0] if text else None
        self._stack.append((sheet, 0, 0))
        try:
            return scalar(self._eval(parse_formula(text), sheet))
        except CellError as err:
            return CellError(err.code)
        finally:
            self._stack.pop()

    def display_text(self, sheet, r, c):
        return format_value(*self.cell_display(sheet, r, c))

    def hyperlink_at(self, sheet, r, c):
        """{"target", "tip"} a click on the cell opens: its hyperlink, or a =HYPERLINK(link, ...) formula's link."""
        sh = self.sheets[sheet]
        link = sh.links.get((r, c))
        if link:
            return link
        raw = sh.cells.get((r, c), "")
        if is_formula(raw) and "HYPERLINK" in raw.upper():
            ast = self._parsed((sheet, r, c), raw)
            if ast[0] == "func" and ast[1] == "HYPERLINK" and ast[2]:
                try:
                    target = scalar(self._eval(ast[2][0], sheet))
                except CellError:
                    return None
                if target not in (None, ""):
                    return {"target": str(target), "tip": ""}
        return None

    # conditional formatting
    def cond_format(self, sheet, r, c):
        """Style from the conditional-format rules covering (r, c): may hold "bg", "fg",
        "bold", "italic" and "databar" (fraction, colour); None when no rule applies."""
        rules = self.sheets[sheet].cond_formats
        if not rules:
            return None
        key = (sheet, r, c)
        if key in self._cf_cache:
            return self._cf_cache[key]
        style = {}
        for i, rule in enumerate(rules):
            r1, c1, r2, c2 = rule["range"]
            if r1 <= r <= r2 and c1 <= c <= c2:
                hit = self._rule_style(sheet, i, rule, r, c)
                if hit:
                    style.update(hit)
        self._cf_cache[key] = style or None
        return style or None

    def _rule_values(self, sheet, rule):
        r1, c1, r2, c2 = rule["range"]
        return [v for v in RangeRef(self, sheet, r1, c1, r2, c2).values() if not isinstance(v, CellError)]

    def _cf_stat(self, sheet, i, rule, name):
        key = (sheet, i, name)
        if key not in self._cf_stats:
            values = self._rule_values(sheet, rule)
            nums = sorted((float(v) for v in values if is_number(v)), reverse=True)
            if name == "counts":
                stat = Counter(_dup_key(v) for v in values if v != "")
            elif name == "sorted":
                stat = nums
            elif name == "avg":
                stat = math.fsum(nums) / len(nums) if nums else 0.0
            elif name == "minmax":
                stat = (nums[-1], nums[0]) if nums else (0.0, 0.0)
            elif name == "median":
                stat = statistics.median(nums) if nums else 0.0
            else:  # an operand of a "cell" rule
                stat = self.evaluate(str(rule.get(name, "")), sheet)
            self._cf_stats[key] = stat
        return self._cf_stats[key]

    def _rule_style(self, sheet, i, rule, r, c):
        kind = rule["type"]
        v = self.get_value(sheet, r, c)
        if v is None or v == "" or isinstance(v, CellError):
            return None
        style = rule.get("style")
        if kind == "cell":
            a = self._cf_stat(sheet, i, rule, "v1")
            if isinstance(a, CellError):
                return None
            op = rule["op"]
            if op in ("between", "notbetween"):
                b = self._cf_stat(sheet, i, rule, "v2")
                if isinstance(b, CellError):
                    return None
                lo, hi = (a, b) if compare("<=", a, b) else (b, a)
                inside = compare(">=", v, lo) and compare("<=", v, hi)
                return style if inside == (op == "between") else None
            return style if compare(op, v, a) else None
        if kind == "text":
            return style if str(rule.get("text", "")).lower() in to_text(v).lower() else None
        if kind in ("dup", "unique"):
            repeated = self._cf_stat(sheet, i, rule, "counts")[_dup_key(v)] > 1
            return style if repeated == (kind == "dup") else None
        if not is_number(v):
            return None
        if kind in ("top", "bottom"):
            nums = self._cf_stat(sheet, i, rule, "sorted")
            n = int(rule.get("n", 10))
            if rule.get("percent"):
                n = int(len(nums) * n / 100)
            n = max(1, min(n, len(nums)))
            hit = v >= nums[n - 1] if kind == "top" else v <= nums[-n]
            return style if hit else None
        if kind in ("above", "below"):
            avg = self._cf_stat(sheet, i, rule, "avg")
            return style if (v > avg if kind == "above" else v < avg) else None
        lo, hi = self._cf_stat(sheet, i, rule, "minmax")
        if kind == "databar":
            lo, hi = min(lo, 0.0), max(hi, 0.0)
            frac = 1.0 if hi == lo else (v - lo) / (hi - lo)
            return {"databar": (max(0.0, min(1.0, frac)), rule.get("color", "#638ec6"))}
        if kind == "scale":
            colors = rule["colors"]
            if hi == lo:
                return {"bg": colors[len(colors) // 2]}
            if len(colors) == 2:
                return {"bg": blend(colors[0], colors[1], (v - lo) / (hi - lo))}
            mid = self._cf_stat(sheet, i, rule, "median")
            if v <= mid:
                return {"bg": blend(colors[0], colors[1], 1.0 if mid == lo else (v - lo) / (mid - lo))}
            return {"bg": blend(colors[1], colors[2], (v - mid) / (hi - mid))}
        return None

    def set_cond_formats(self, sheet, rules):
        self.sheets[sheet].cond_formats = rules
        self.invalidate()

    # data validation
    def validation_at(self, sheet, r, c):
        """The data-validation rule covering (r, c), or None."""
        for rule in reversed(self.sheets[sheet].validations):
            r1, c1, r2, c2 = rule["range"]
            if r1 <= r <= r2 and c1 <= c <= c2:
                return rule
        return None

    def _rule_eval(self, sheet, rule, text, r, c):
        """Value of a rule operand, with relative references moved from the rule's first cell to (r, c)."""
        text = str(text or "").strip()
        if is_formula(text):
            text = shift_formula(text, r - rule["range"][0], c - rule["range"][1])
        self._stack.append((sheet, r, c))
        try:
            if not is_formula(text):
                return parse_literal(text)[0] if text else None
            return self._eval(parse_formula(text), sheet)
        except CellError as err:
            return CellError(err.code)
        finally:
            self._stack.pop()

    def validation_options(self, sheet, r, c, rule=None):
        """Texts offered by a list rule: its comma-separated items or the cells of its source range."""
        rule = rule or self.validation_at(sheet, r, c)
        source = str(rule.get("source", "")).strip()
        if not is_formula(source):
            return [item.strip() for item in re.split(r"[,;]", source) if item.strip()]
        value = self._rule_eval(sheet, rule, source, r, c)
        if isinstance(value, RangeRef):
            texts = (self.display_text(value.sheet, rr, cc) for rr in range(value.r1, value.r2 + 1)
                     for cc in range(value.c1, value.c2 + 1))
            return [t for t in texts if t != ""]
        return [] if value is None or isinstance(value, CellError) else [to_text(value)]

    def validate(self, sheet, r, c, raw):
        """The rule that typing `raw` into (r, c) breaks, or None when the entry is allowed."""
        rule = self.validation_at(sheet, r, c)
        if rule is None or rule.get("type", "any") == "any":
            return None
        if not raw:
            return None if rule.get("blank", True) else rule
        old = self.get_raw(sheet, r, c)
        self.set_raw(sheet, r, c, raw)  # judge the value the cell would really get
        try:
            ok = self._valid_value(sheet, rule, r, c, raw)
        finally:
            self.set_raw(sheet, r, c, old)
        return None if ok else rule

    def _valid_value(self, sheet, rule, r, c, raw):
        value, kind = self.get_value(sheet, r, c), rule["type"]
        if isinstance(value, CellError):
            return False
        if kind == "custom":  # the formula must come out TRUE (or a non-zero number)
            try:
                result = scalar(self._rule_eval(sheet, rule, rule.get("v1"), r, c))
            except CellError:
                return False
            return (isinstance(result, bool) or is_number(result)) and bool(result)
        if kind == "list":
            options = self.validation_options(sheet, r, c, rule)
            text = self.display_text(sheet, r, c).lower()
            typed = (raw[1:] if raw.startswith("'") else raw).strip().lower()
            return any(o.lower() in (text, typed) or (is_number(value) and parse_literal(o)[0] == value)
                       for o in options)
        if kind == "length":
            x = float(len(to_text(value)))
        elif is_number(value):
            x = float(value)
            if kind == "whole" and not x.is_integer():
                return False
        else:
            return False
        a = self._rule_eval(sheet, rule, rule.get("v1"), r, c)
        b = self._rule_eval(sheet, rule, rule.get("v2"), r, c)
        try:
            a = to_number(scalar(a))
            b = to_number(scalar(b)) if rule.get("op", "between") in ("between", "notbetween") else a
        except CellError:
            return False
        op = rule.get("op", "between")
        if op in ("between", "notbetween"):
            return (min(a, b) <= x <= max(a, b)) == (op == "between")
        return _CMP_OPS[op](x, a)

    # charts
    def chart_layout(self, sheet, chart):
        """How a chart reads its range, like Excel: series run down the columns of a block at least as
        tall as it is wide (else along its rows); a first row/column holding text gives the series names /
        category labels. Returns (rect clipped to the data, by_cols, has_names, has_categories)."""
        r1, c1, r2, c2 = chart["range"]
        rows, cols = self.used_extent(sheet)
        r2, c2 = min(r2, max(r1, rows - 1)), min(c2, max(c1, cols - 1))
        by_cols = r2 - r1 >= c2 - c1
        along, across = (r2 - r1 + 1, c2 - c1 + 1) if by_cols else (c2 - c1 + 1, r2 - r1 + 1)

        def value(i, j):
            return self.get_value(sheet, *((r1 + i, c1 + j) if by_cols else (r1 + j, c1 + i)))

        def is_text(v):
            return isinstance(v, str) and v != ""

        names = along > 1 and any(is_text(value(0, j)) for j in range(across))
        first = 1 if names else 0
        cats = across > 1 and (any(is_text(value(i, 0)) for i in range(first, along))
                               or (names and value(0, 0) in (None, "")))
        return (r1, c1, r2, c2), by_cols, names, cats

    def chart_data(self, sheet, chart):
        """(category labels, [(series name, [number or None])]) plotted by a chart."""
        (r1, c1, r2, c2), by_cols, names, cats = self.chart_layout(sheet, chart)
        along, across = (r2 - r1 + 1, c2 - c1 + 1) if by_cols else (c2 - c1 + 1, r2 - r1 + 1)
        along, across = min(along, 1000), min(across, 50)

        def pos(i, j):
            return (r1 + i, c1 + j) if by_cols else (r1 + j, c1 + i)

        first = 1 if names else 0
        labels = ([self.display_text(sheet, *pos(i, 0)) for i in range(first, along)] if cats
                  else [str(k + 1) for k in range(along - first)])
        series = []
        for j in range(1 if cats else 0, across):
            name = self.display_text(sheet, *pos(0, j)) if names else f"Chuỗi {len(series) + 1}"
            values = [self.get_value(sheet, *pos(i, j)) for i in range(first, along)]
            series.append((name, [float(v) if is_number(v) else None for v in values]))
        numeric = [item for item in series if any(v is not None for v in item[1])]
        return labels, numeric or series  # a text column in the middle of the data is not a series

    # filtering
    def filter_values(self, sheet, col):
        """Distinct display texts in a filter column (data rows only), sorted like Excel."""
        f = self.sheets[sheet].filter
        r1, _, r2, _ = f["range"]
        texts = {}
        for r in range(r1 + 1, r2 + 1):
            texts.setdefault(self.display_text(sheet, r, col), self.get_value(sheet, r, col))
        return sorted(texts, key=lambda t: (t == "", natural_key(texts[t]) if texts[t] is not None else (9,)))

    def cell_color(self, sheet, r, c, what="bg"):
        """Fill ("bg") or font ("fg") colour a cell shows, conditional formatting included; None when
        it has none of its own."""
        style = self.cond_format(sheet, r, c) or {}
        fmt = self.sheets[sheet].formats.get((r, c)) or {}
        return style.get(what) or fmt.get(what)

    def filter_colors(self, sheet, col, what="bg"):
        """Colours of a filter column's data cells, in order of appearance (None: no fill / automatic),
        when they are not all the same - like Excel, which offers colour filters only then."""
        r1, _, r2, _ = self.sheets[sheet].filter["range"]
        colors = list(dict.fromkeys(self.cell_color(sheet, r, col, what) for r in range(r1 + 1, r2 + 1)))
        return colors if len(colors) > 1 else []

    def filter_kind(self, sheet, col):
        """"number", "date" or "text": what most data cells of a filter column hold (picks Excel's
        Number / Date / Text Filters menu)."""
        r1, _, r2, _ = self.sheets[sheet].filter["range"]
        counts = Counter()
        for r in range(r1 + 1, r2 + 1):
            value = self.get_value(sheet, r, col)
            if is_number(value):
                fmt = self.cell_display(sheet, r, col)[1]
                counts["date" if fmt and is_date_format(fmt) else "number"] += 1
            elif value not in (None, ""):
                counts["text"] += 1
        return counts.most_common(1)[0][0] if counts else "text"

    def filter_match(self, sheet, r, col, crit, stats=None):
        """Whether row r passes a filter column's criterion (see FILTER_OPS above). `stats` caches the
        sorted numbers of each column for top-10 and average filters."""
        if not isinstance(crit, dict):
            return self.display_text(sheet, r, col) in crit
        kind = crit["type"]
        if kind == "color":
            return self.cell_color(sheet, r, col, crit.get("what", "bg")) == crit.get("color")
        value = self.get_value(sheet, r, col)
        if kind == "custom":
            text = self.display_text(sheet, r, col)
            hits = [filter_condition(op, target, value, text) for op, target in crit["conds"]]
            return all(hits) if crit.get("join", "and") == "and" else any(hits)
        if not is_number(value):
            return False
        if kind == "dynamic":
            first, last = dynamic_date_span(crit["kind"], dt.date.today())
            return first <= serial_to_datetime(math.floor(value)).date() <= last
        stats = {} if stats is None else stats
        if col not in stats:
            r1, _, r2, _ = self.sheets[sheet].filter["range"]
            stats[col] = sorted(float(v) for v in (self.get_value(sheet, rr, col) for rr in range(r1 + 1, r2 + 1))
                                if is_number(v))
        nums = stats[col]
        if kind == "average":
            avg = math.fsum(nums) / len(nums)
            return value > avg if crit.get("above", True) else value < avg
        n = int(crit.get("n", 10))
        if crit.get("percent"):
            n = int(len(nums) * n / 100)
        n = max(1, min(n, len(nums)))
        return value <= nums[n - 1] if crit.get("bottom") else value >= nums[-n]

    def apply_filter(self, sheet):
        """Recompute which data rows the filter criteria hide."""
        f = self.sheets[sheet].filter
        if not f:
            return
        r1, _, r2, _ = f["range"]
        hidden, stats = set(), {}
        for r in range(r1 + 1, r2 + 1):
            for col, crit in f["criteria"].items():
                if not self.filter_match(sheet, r, col, crit, stats):
                    hidden.add(r)
                    break
        f["hidden"] = hidden
        self.invalidate()  # SUBTOTAL results depend on which rows are filtered out

    # structure
    def insert_cells(self, sheet, r1, c1, r2, c2, shift="down"):
        if shift == "down":
            self._shift(sheet, 0, r1, r2 - r1 + 1, (c1, c2))
        else:
            self._shift(sheet, 1, c1, c2 - c1 + 1, (r1, r2))

    def delete_cells(self, sheet, r1, c1, r2, c2, shift="up"):
        if shift == "up":
            self._shift(sheet, 0, r1, -(r2 - r1 + 1), (c1, c2))
        else:
            self._shift(sheet, 1, c1, -(c2 - c1 + 1), (r1, r2))

    def insert_rows(self, sheet, at, count=1):
        self._shift(sheet, 0, at, count)

    def delete_rows(self, sheet, at, count=1):
        self._shift(sheet, 0, at, -count)

    def insert_cols(self, sheet, at, count=1):
        self._shift(sheet, 1, at, count)

    def delete_cols(self, sheet, at, count=1):
        self._shift(sheet, 1, at, -count)

    def _shift(self, sheet, axis, at, delta, band=None):
        """Insert (delta > 0) or delete (delta < 0) rows (axis 0) or columns (axis 1).

        With `band` = (lo, hi) only the cells whose other coordinate is inside the band
        move, which is how "insert/delete cells, shift down/right/up/left" works."""
        limit = MAX_ROWS if axis == 0 else MAX_COLS
        end = at - delta  # first index after a deleted block
        lo, hi = band if band else (0, max(MAX_ROWS, MAX_COLS))

        def move(i):
            if delta > 0:
                if i < at:
                    return i
                return i + delta if i + delta < limit else None
            if i >= end:
                return i + delta
            return None if i >= at else i

        def move_span(i1, i2):
            if delta > 0:
                n1, n2 = move(i1), move(i2)
                return None if n1 is None else (n1, limit - 1 if n2 is None else n2)
            n1 = i1 + delta if i1 >= end else (at if i1 >= at else i1)
            n2 = i2 + delta if i2 >= end else (at - 1 if i2 >= at else i2)
            return (n1, n2) if n1 <= n2 else None

        target = sheet.lower()
        for sh in self.sheets.values():
            for key, raw in list(sh.cells.items()):
                if is_formula(raw):
                    sh.cells[key] = _rewrite_axis(raw, sh.name, target, axis, move, move_span, band)
        for d in self.names.values():
            d["ref"] = _rewrite_axis(d["ref"], sheet, target, axis, move, move_span, band)
        self._names_ast.clear()

        def remap(d):
            out = {}
            for (r, c), v in d.items():
                idx, other = (r, c) if axis == 0 else (c, r)
                if not lo <= other <= hi:
                    out[(r, c)] = v
                elif (n := move(idx)) is not None:
                    out[(n, c) if axis == 0 else (r, n)] = v
            return out

        def move_rect(rect):
            """A block (merge, rule range, filter range) after the shift; None if deleted."""
            r1, c1, r2, c2 = rect
            a1, a2, o1, o2 = (r1, r2, c1, c2) if axis == 0 else (c1, c2, r1, r2)
            if not (lo <= o1 and o2 <= hi):
                return tuple(rect)
            span = move_span(a1, a2)
            if span is None:
                return None
            return (span[0], c1, span[1], c2) if axis == 0 else (r1, span[0], r2, span[1])

        def move_set(indices):
            return {n for i in indices if (n := move(i)) is not None}

        sh = self.sheets[sheet]
        sh.cells, sh.formats, sh.notes = remap(sh.cells), remap(sh.formats), remap(sh.notes)
        sh.links = remap(sh.links)
        sh.set_merges(m for m in map(move_rect, sh.merges) if m and (m[0], m[1]) != (m[2], m[3]))
        sh.cond_formats = [dict(rule, range=rect) for rule in sh.cond_formats
                           if (rect := move_rect(rule["range"])) is not None]
        sh.validations = [dict(rule, range=rect) for rule in sh.validations
                          if (rect := move_rect(rule["range"])) is not None]
        for other in self.sheets.values():  # pivots reading this sheet, wherever their reports are
            for p in other.pivots:
                if p["source"]["sheet"].lower() == target:
                    p["source"]["range"] = move_rect(p["source"]["range"]) or p["source"]["range"]
        if band is None and sh.page_setup:
            ps = dict(sh.page_setup)
            if ps.get("area"):
                ps["area"] = move_rect(ps["area"])
            if axis == 0 and ps.get("titles"):
                span = move_span(*ps["titles"])
                ps["titles"] = tuple(span) if span else None
            key = "breaks" if axis == 0 else "col_breaks"
            if ps.get(key):
                ps[key] = sorted(move_set(ps[key]))
            sh.page_setup = ps
        if band is None:
            for p in sh.pivots:
                r, c = p["anchor"]
                n = move(r if axis == 0 else c)
                n = min(at, limit - 1) if n is None else n
                p["anchor"] = (n, c) if axis == 0 else (r, n)
                if p["rect"]:
                    p["rect"] = move_rect(p["rect"])
        charts = []
        for chart in sh.charts:
            rect = move_rect(chart["range"]) if chart.get("range") else None
            if rect is None and chart.get("range"):  # its data is gone
                continue
            r, c, dx, dy = chart["anchor"]
            if band is None:  # charts sit on whole rows/columns: they move when those do
                n = move(r if axis == 0 else c)
                n = min(at, limit - 1) if n is None else n
                r, c = (n, c) if axis == 0 else (r, n)
            charts.append(dict(chart, anchor=(r, c, dx, dy), **({"range": rect} if rect else {})))
        sh.charts = charts
        if sh.filter:
            rect = move_rect(sh.filter["range"])
            if rect is None:
                sh.filter = None
            else:
                sh.filter["range"] = rect
                if band is None and axis == 0:
                    sh.filter["hidden"] = move_set(sh.filter["hidden"])
                elif band is None:
                    sh.filter["criteria"] = {n: v for c, v in sh.filter["criteria"].items()
                                             if (n := move(c)) is not None}
        if band is None and sh.freeze:
            sh.freeze = _shift_freeze(sh.freeze, axis, at, move_span)
        if band is None and axis == 0:
            sh.row_heights = {n: h for r, h in sh.row_heights.items() if (n := move(r)) is not None}
            sh.custom_rows, sh.hidden_rows = move_set(sh.custom_rows), move_set(sh.hidden_rows)
        elif band is None:
            sh.col_widths = {n: w for c, w in sh.col_widths.items() if (n := move(c)) is not None}
            sh.hidden_cols = move_set(sh.hidden_cols)
        if band is None:
            levels = {n: v for i, v in (sh.row_levels if axis == 0 else sh.col_levels).items()
                      if (n := move(i)) is not None}
            if delta > 0:  # rows inserted inside a group join it, as in Excel
                inner = min(levels.get(at - 1, 0), levels.get(at + delta, 0))
                levels.update({i: inner for i in range(at, min(at + delta, limit)) if inner})
            if axis == 0:
                sh.row_levels, sh.collapsed_rows = levels, move_set(sh.collapsed_rows)
            else:
                sh.col_levels, sh.collapsed_cols = levels, move_set(sh.collapsed_cols)
        self._ast.clear()
        self.invalidate()

    # outline (Data > Group)
    def _outline(self, sheet, axis):
        sh = self.sheets[sheet]
        if axis == 0:
            return sh.row_levels, sh.hidden_rows, sh.collapsed_rows
        return sh.col_levels, sh.hidden_cols, sh.collapsed_cols

    def group(self, sheet, axis, i1, i2, delta=1):
        """Group (delta 1) or ungroup (delta -1) rows (axis 0) or columns i1..i2 by one outline level.
        Returns False when there was nothing to ungroup."""
        levels, hidden, marks = self._outline(sheet, axis)
        changed = False
        for i in range(i1, i2 + 1):
            n = max(0, min(MAX_OUTLINE, levels.get(i, 0) + delta))
            changed |= n != levels.get(i, 0)
            if n:
                levels[i] = n
            else:
                levels.pop(i, None)
        ends = {b + 1 for _, _, b in outline_groups(levels)}
        marks.intersection_update(ends)
        return changed

    def clear_outline(self, sheet, axis):
        """Ungroup everything on an axis, showing what collapsed groups had hidden."""
        levels, hidden, marks = self._outline(sheet, axis)
        for _, a, b in outline_groups(levels):
            if b + 1 in marks:
                hidden.difference_update(range(a, b + 1))
        levels.clear()
        marks.clear()

    def set_collapsed(self, sheet, axis, level, start, end, collapsed):
        """Collapse (hide) or expand a group; expanding keeps the collapsed groups inside it hidden."""
        levels, hidden, marks = self._outline(sheet, axis)
        if collapsed:
            hidden.update(range(start, end + 1))
            marks.add(end + 1)
            return
        marks.discard(end + 1)
        hidden.difference_update(range(start, end + 1))
        for lv, a, b in outline_groups(levels):
            if lv > level and start <= a and b <= end and b + 1 in marks:
                hidden.update(range(a, b + 1))

    def show_level(self, sheet, axis, n):
        """The outline's level buttons: show what is above level n, collapse every group of level n or deeper."""
        levels, hidden, marks = self._outline(sheet, axis)
        for i, lv in levels.items():
            if lv >= n:
                hidden.add(i)
            else:
                hidden.discard(i)
        marks.clear()
        marks.update(b + 1 for lv, _, b in outline_groups(levels) if lv >= n)

    def remove_subtotals(self, sheet, rect):
        """Data > Subtotal > Remove All: delete the rows holding SUBTOTAL formulas in a table and its
        row outline. Returns the table's rectangle afterwards."""
        r1, c1, r2, c2 = rect
        sh = self.sheets[sheet]
        rows = [r for r in range(r2, r1, -1)
                if any(_SUBTOTAL_RE.match(sh.cells.get((r, c), "")) for c in range(c1, c2 + 1))]
        for r in rows:
            self.delete_rows(sheet, r)
        r2 -= len(rows)
        levels, hidden, marks = self._outline(sheet, 0)
        for r in range(r1, r2 + 2):
            levels.pop(r, None)
            marks.discard(r)
        hidden.difference_update(range(r1, r2 + 1))
        return r1, c1, r2, c2

    def add_subtotals(self, sheet, rect, by, func, columns, replace=True):
        """Data > Subtotal on a table whose first row holds headers: after each run of equal values in
        column `by` insert a row "<value> <name>" with =SUBTOTAL(func, ...) under `columns`, then a grand
        total row, and outline the rows (level 1: grand total, 2: subtotals, 3: everything).
        Returns the table's rectangle afterwards."""
        if replace:
            rect = self.remove_subtotals(sheet, rect)
        r1, c1, r2, c2 = rect
        sh = self.sheets[sheet]
        name, grand = SUBTOTAL_NAMES.get(func, SUBTOTAL_NAMES[9])
        runs = []
        for r in range(r1 + 1, r2 + 1):
            key = self.display_text(sheet, r, by)
            if runs and runs[-1][2] == key:
                runs[-1][1] = r
            else:
                runs.append([r, r, key])
        if not runs:
            return rect
        bold = {"bold": True}

        def total_row(row, label, first, last):
            self.set_raw(sheet, row, by, label)
            self.set_format(sheet, row, by, {**(self.get_format(sheet, last, by) or {}), **bold})
            for c in columns:
                self.set_raw(sheet, row, c, f"=SUBTOTAL({func},{range_name(first, c, last, c)})")
                num = (self.get_format(sheet, last, c) or {}).get("num")
                self.set_format(sheet, row, c, {**bold, **({"num": num} if num else {})})
        offset, groups = 0, []
        for start, end, key in runs:
            start, end = start + offset, end + offset
            self.insert_rows(sheet, end + 1)
            total_row(end + 1, f"{key} {name}".strip(), start, end)
            groups.append((start, end))
            offset += 1
        last = r2 + offset
        self.insert_rows(sheet, last + 1)
        total_row(last + 1, grand, r1 + 1, last)
        sh.row_levels.pop(last + 1, None)  # the grand total row is not part of the outline it closes
        self.group(sheet, 0, r1 + 1, last)
        for start, end in groups:
            self.group(sheet, 0, start, end)
        self.invalidate()
        return r1, c1, last + 1, c2

    def snapshot(self):
        return [(s.name, s.copy_state()) for s in self.sheets.values()], copy.deepcopy(self.names)

    def restore(self, snap):
        sheets, names = snap
        self.sheets = {name: Sheet.from_state(name, state) for name, state in sheets}
        self.names = copy.deepcopy(names)
        self._names_ast.clear()
        self._ast.clear()
        self.invalidate()

    # defined names (Formulas > Name Manager)
    def check_name(self, name, ignore=None):
        """Error message (shown in the UI) when `name` can't be a defined name, else None."""
        if not name:
            return "Tên không được để trống."
        if len(name) > 255 or not _DEFINED_NAME_RE.match(name):
            return ("Tên phải bắt đầu bằng chữ cái hoặc dấu gạch dưới và chỉ gồm chữ, số, dấu chấm, gạch dưới "
                    "(không có khoảng trắng).")
        if _REF_LIKE_RE.match(name) or name.upper() in ("TRUE", "FALSE"):
            return "Tên không được giống địa chỉ ô (như A1, R1C1) hay TRUE/FALSE."
        if name.upper() in self.names and name.upper() != (ignore or "").upper():
            return "Đã có tên này trong sổ làm việc."
        return None

    def define_name(self, name, ref, comment="", old=None):
        """Add a name, or replace `old`; ref is a formula such as '=Sheet1!$A$1:$A$9' or '=0.1'."""
        err = self.check_name(name, ignore=old)
        if err:
            raise ValueError(err)
        if old:
            self.names.pop(old.upper(), None)
        ref = ref.strip()
        self.names[name.upper()] = {"name": name, "ref": normalize_formula(ref if ref.startswith("=") else "=" + ref),
                                    "comment": comment}
        self._names_ast.clear()
        self.invalidate()

    def sorted_names(self):
        """The defined names' entries in natural Vietnamese alphabetical order."""
        return sorted(self.names.values(), key=lambda d: natural_key(d["name"]))

    def delete_name(self, name):
        self.names.pop(name.upper(), None)
        self._names_ast.clear()
        self.invalidate()

    def name_target(self, name):
        """(sheet, (r1, c1, r2, c2)) when a name refers to a block of cells, else None."""
        d = self.names.get(name.upper())
        ast = parse_formula(d["ref"]) if d else None
        if not ast or ast[0] not in ("ref", "range") or ast[1] is None:
            return None
        sheet = self.resolve_sheet(ast[1])
        if sheet is None:
            return None
        return (sheet, ast[2:4] * 2) if ast[0] == "ref" else (sheet, ast[2:6])

    def name_value(self, name):
        """Short text of what a name gives now, for the Name Manager."""
        try:
            value = self._eval(self._name_ast(name.upper()), next(iter(self.sheets)))
        except CellError as err:
            return err.code
        except RecursionError:
            return REF
        if isinstance(value, RangeRef):
            if value.nrows * value.ncols == 1:
                return format_value(value.value(0, 0), None)
            values = [format_value(v, None) for _, v in zip(range(7), value.values())]
            return "{" + "; ".join(values[:6]) + ("; ..." if len(values) > 6 else "") + "}"
        return format_value(value, None)

    # serialisation
    def to_dict(self, active=None):
        def sheet_dict(s):
            d = {
                "name": s.name,
                "cells": {cell_name(r, c): raw for (r, c), raw in sorted(s.cells.items())},
                "formats": {cell_name(r, c): f for (r, c), f in sorted(s.formats.items())},
                "col_widths": {col_name(c): w for c, w in sorted(s.col_widths.items())},
                "row_heights": {str(r + 1): h for r, h in sorted(s.row_heights.items())},
                "custom_rows": sorted(r + 1 for r in s.custom_rows),
                "hidden_rows": sorted(r + 1 for r in s.hidden_rows),
                "hidden_cols": [col_name(c) for c in sorted(s.hidden_cols)],
                "merges": [range_name(*m) for m in s.merges],
                "cond_formats": [dict(rule, range=range_name(*rule["range"])) for rule in s.cond_formats],
                "tab_color": s.tab_color,
            }
            if s.freeze:
                d["freeze"] = list(s.freeze)
            if s.row_levels or s.col_levels:
                d["outline"] = {"rows": {str(r + 1): v for r, v in sorted(s.row_levels.items())},
                                "cols": {col_name(c): v for c, v in sorted(s.col_levels.items())},
                                "collapsed_rows": sorted(r + 1 for r in s.collapsed_rows),
                                "collapsed_cols": [col_name(c) for c in sorted(s.collapsed_cols)]}
            if s.validations:
                d["validations"] = [dict(rule, range=range_name(*rule["range"])) for rule in s.validations]
            if s.notes:
                d["notes"] = {cell_name(r, c): text for (r, c), text in sorted(s.notes.items())}
            if s.links:
                d["links"] = {cell_name(r, c): link for (r, c), link in sorted(s.links.items())}
            if s.charts:
                d["charts"] = [dict(ch, anchor=list(ch["anchor"]), size=list(ch["size"]),
                                    **({"range": range_name(*ch["range"])} if ch.get("range") else {}))
                               for ch in s.charts]
            if s.pivots:
                d["pivots"] = [_pivot_to_dict(p) for p in s.pivots]
            if s.page_setup:
                ps = dict(s.page_setup)
                if ps.get("area"):
                    ps["area"] = range_name(*ps["area"])
                if ps.get("titles"):
                    ps["titles"] = [ps["titles"][0] + 1, ps["titles"][1] + 1]
                for key in ("breaks", "col_breaks"):
                    if ps.get(key):
                        ps[key] = [i + 1 for i in ps[key]]
                d["page_setup"] = ps
            if s.filter:
                d["filter"] = {"range": range_name(*s.filter["range"]),
                               "criteria": {col_name(c): v if isinstance(v, dict) else sorted(v)
                                            for c, v in s.filter["criteria"].items()}}
            return d
        out = {"format": FILE_TAG, "version": 2, "active": active,
               "sheets": [sheet_dict(s) for s in self.sheets.values()]}
        if self.names:
            out["names"] = [dict(d) for _, d in sorted(self.names.items())]
        return out

    @classmethod
    def from_dict(cls, data):
        if not isinstance(data, dict) or data.get("format") != FILE_TAG:
            raise ValueError("Đây không phải tệp PyExcel hợp lệ.")
        wb = cls(())
        for sd in data.get("sheets", []):
            name = wb.add_sheet(sd["name"])  # add_sheet replaces wb.sheets, so look up afterwards
            sh = wb.sheets[name]
            for addr, raw in sd.get("cells", {}).items():
                pos = parse_address(addr)
                if pos and len(pos) == 2 and raw:
                    sh.cells[pos] = str(raw)
            for addr, fmt in sd.get("formats", {}).items():
                pos = parse_address(addr)
                if pos and len(pos) == 2 and fmt:
                    sh.formats[pos] = dict(fmt)
            for letter, width in sd.get("col_widths", {}).items():
                sh.col_widths[col_index(letter)] = int(width)
            sh.row_heights = {int(r) - 1: int(h) for r, h in sd.get("row_heights", {}).items()}
            sh.custom_rows = {int(r) - 1 for r in sd.get("custom_rows", [])}
            sh.hidden_rows = {int(r) - 1 for r in sd.get("hidden_rows", [])}
            sh.hidden_cols = {col_index(c) for c in sd.get("hidden_cols", [])}
            outline = sd.get("outline") or {}
            sh.row_levels = {int(r) - 1: int(v) for r, v in outline.get("rows", {}).items()}
            sh.col_levels = {col_index(c): int(v) for c, v in outline.get("cols", {}).items()}
            sh.collapsed_rows = {int(r) - 1 for r in outline.get("collapsed_rows", [])}
            sh.collapsed_cols = {col_index(c) for c in outline.get("collapsed_cols", [])}
            sh.set_merges(m for m in map(parse_range, sd.get("merges", [])) if m)
            sh.cond_formats = [dict(rule, range=parse_range(rule["range"])) for rule in sd.get("cond_formats", [])
                               if parse_range(rule.get("range", ""))]
            sh.tab_color = sd.get("tab_color")
            freeze = sd.get("freeze")
            if isinstance(freeze, list) and len(freeze) == 4:
                sh.freeze = tuple(int(v) for v in freeze)
            sh.validations = [dict(rule, range=parse_range(rule["range"])) for rule in sd.get("validations", [])
                              if parse_range(rule.get("range", ""))]
            for addr, text in sd.get("notes", {}).items():
                pos = parse_address(addr)
                if pos and len(pos) == 2 and text:
                    sh.notes[pos] = str(text)
            for addr, link in sd.get("links", {}).items():
                pos = parse_address(addr)
                if pos and len(pos) == 2 and isinstance(link, dict) and link.get("target"):
                    sh.links[pos] = {"target": str(link["target"]), "tip": str(link.get("tip") or "")}
            sh.pivots = [p for p in map(_pivot_from_dict, sd.get("pivots", [])) if p]
            ps = sd.get("page_setup")
            if isinstance(ps, dict):
                ps = dict(ps)
                ps["area"] = parse_range(ps["area"]) if ps.get("area") else None
                ps["titles"] = (int(ps["titles"][0]) - 1, int(ps["titles"][1]) - 1) if ps.get("titles") else None
                for key in ("breaks", "col_breaks"):
                    ps[key] = [int(i) - 1 for i in ps.get(key, [])]
                sh.page_setup = ps
            for ch in sd.get("charts", []):
                if ch.get("type") == PICTURE and ch.get("image"):
                    sh.charts.append(dict(ch, id=ch.get("id") or new_picture_id(),
                                          anchor=tuple(ch.get("anchor", (0, 0, 0, 0))),
                                          size=tuple(ch.get("size", CHART_SIZE))))
                    continue
                rect = parse_range(ch.get("range", ""))
                if rect and ch.get("type") in CHART_TYPES:
                    sh.charts.append(dict(ch, range=rect, anchor=tuple(ch.get("anchor", (0, 0, 0, 0))),
                                          size=tuple(ch.get("size", CHART_SIZE))))
            f = sd.get("filter")
            if f and parse_range(f.get("range", "")):
                sh.filter = {"range": parse_range(f["range"]), "hidden": set(),
                             "criteria": {col_index(c): dict(v) if isinstance(v, dict) else set(v)
                                          for c, v in f.get("criteria", {}).items()}}
        if not wb.sheets:
            wb.add_sheet()
        for name in wb.sheets:  # after every sheet exists, so cross-sheet formulas resolve
            wb.apply_filter(name)
        for d in data.get("names", []):
            if isinstance(d, dict) and wb.check_name(str(d.get("name", ""))) is None and d.get("ref"):
                wb.define_name(d["name"], str(d["ref"]), str(d.get("comment") or ""))
        wb.active = data.get("active") if data.get("active") in wb.sheets else None
        return wb


def _pivot_to_dict(p):
    return dict(p, source={"sheet": p["source"]["sheet"], "range": range_name(*p["source"]["range"])},
                anchor=cell_name(*p["anchor"]), rect=range_name(*p["rect"]) if p["rect"] else None,
                items={f: sorted(v) for f, v in p["items"].items()}, values=[dict(v) for v in p["values"]])


def _pivot_from_dict(d):
    """A pivot read back from a file, or None when it is unusable."""
    try:
        rect, anchor = parse_range(d["source"]["range"]), parse_address(d["anchor"])
        if not rect or not anchor or len(anchor) != 2:
            return None
        p = dict(d, source={"sheet": str(d["source"]["sheet"]), "range": rect}, anchor=anchor,
                 rect=parse_range(d["rect"]) if d.get("rect") else None,
                 items={f: set(v) for f, v in d.get("items", {}).items()},
                 values=[dict(v) for v in d.get("values", [])])
        for key in ("filters", "cols", "rows"):
            p[key] = [str(f) for f in d.get(key, [])]
        for key, default in (("sort", {}), ("group", {}), ("grand_rows", True), ("grand_cols", True),
                             ("subtotals", True), ("layout", "compact"), ("name", "PivotTable")):
            p.setdefault(key, default)
        return p
    except (KeyError, TypeError, AttributeError):
        return None


PIVOT_STORE = "_PyExcelPivots"  # very hidden sheet keeping pivot definitions in .xlsx files


# ------------------------------------------------------------------ files

def load_workbook(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _load_xlsx(path)
    if ext in (".csv", ".txt", ".tsv"):
        return _load_csv(path)
    with open(path, encoding="utf-8") as f:
        return Workbook.from_dict(json.load(f))


def save_workbook(wb, path, active=None):
    """Write the workbook; CSV keeps only the `active` sheet's displayed values."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        _save_xlsx(wb, path, active)
    elif ext == ".csv":
        _save_csv(wb, active if active in wb.sheets else next(iter(wb.sheets)), path)
    else:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(wb.to_dict(active), f, ensure_ascii=False, indent=1)


def _load_csv(path):
    with open(path, "rb") as f:
        data = f.read()
    for enc in ("utf-8-sig", "cp1258", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    wb = Workbook()
    sh = wb.sheets["Sheet1"]
    for r, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if r >= MAX_ROWS:
            break
        for c, val in enumerate(row[:MAX_COLS]):
            if val != "":
                sh.cells[(r, c)] = val
    return wb


def _save_csv(wb, sheet, path):
    rows, cols = wb.used_extent(sheet)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        for r in range(rows):
            writer.writerow([format_value(*wb.cell_display(sheet, r, c)) for c in range(cols)])


_XL_OPS = {">": "greaterThan", "<": "lessThan", ">=": "greaterThanOrEqual", "<=": "lessThanOrEqual",
           "=": "equal", "<>": "notEqual", "between": "between", "notbetween": "notBetween"}
_XL_BORDER = {"hair": "thin", "mediumDashed": "dashed", "dashDot": "dashed", "mediumDashDot": "dashed",
              "dashDotDot": "dotted", "mediumDashDotDot": "dashed", "slantDashDot": "dashed"}


def _xl_color(color):
    rgb = getattr(color, "rgb", None)
    if getattr(color, "type", None) == "rgb" and isinstance(rgb, str) and len(rgb) == 8:
        return "#" + rgb[2:].lower()
    return None


def _xl_hex(color):
    return "FF" + color[1:].upper() if color else None


def _numfmt_from_xlsx(cell):
    nf = cell.number_format or "General"
    if nf in ("General", "@"):
        return None if nf == "General" else "@"
    low = nf.lower()
    if getattr(cell, "is_date", False):
        if "h" in low:
            return "dd/mm/yyyy hh:mm" if ("d" in low or "y" in low) else "hh:mm"
        return "dd/mm/yyyy"
    sections = []
    for section in split_sections(nf)[:3]:
        section = re.sub(r"\[[^\]]*\]|_.|\*.", "", section)
        section = re.sub(r"\\(.)", r"\1", section).strip()
        # unquote literal text that can't be mistaken for format codes: '#,##0 "₫"' -> '#,##0 ₫'
        sections.append(re.sub(r'"([^"]*)"', lambda m: m.group(0) if re.search(r"[0#?.,%EeDdMmYyHhSs@;/]", m.group(1))
                               else m.group(1), section))
    first = sections[0]
    if _NUM_CORE_RE.search(first) or "?/" in first or _SCI_RE.search(first):
        return ";".join(sections)
    return None


def _raw_from_xlsx(cell):
    v = cell.value
    if v is None:
        return None
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return number_text(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return number_text(date_to_serial(v))
    if isinstance(v, dt.time):
        return number_text((v.hour * 3600 + v.minute * 60 + v.second) / 86400)
    text = str(getattr(v, "text", v))  # ArrayFormula objects keep the formula in .text
    if cell.data_type == "f" and is_formula(text):
        return normalize_formula(text.replace("_xlfn.", "").replace("_xlws.", ""))
    return text if parse_literal(text) == (text, None) else "'" + text


def _format_from_xlsx(cell):
    fmt = {}
    font = cell.font
    if font is not None:
        if font.name and font.name != "Calibri":
            fmt["font"] = font.name
        if font.sz and float(font.sz) != 11:
            size = float(font.sz)
            fmt["size"] = int(size) if size.is_integer() else size
        for key, on in (("bold", font.b), ("italic", font.i), ("strike", font.strike)):
            if on:
                fmt[key] = True
        if font.u:
            fmt["underline"] = "double" if "double" in str(font.u) else "single"
        fg = _xl_color(font.color)
        if fg and fg != "#000000":
            fmt["fg"] = fg
    if cell.fill is not None and cell.fill.fill_type == "solid":
        bg = _xl_color(cell.fill.fgColor)
        if bg:
            fmt["bg"] = bg
    al = cell.alignment
    if al is not None:
        if al.horizontal in ("left", "center", "right"):
            fmt["align"] = al.horizontal
        if al.vertical in ("top", "center"):
            fmt["valign"] = al.vertical
        if al.wrap_text:
            fmt["wrap"] = True
        if al.indent:
            fmt["indent"] = int(al.indent)
        rot = int(al.text_rotation or 0)
        if rot:
            fmt["rotate"] = rot if rot <= 90 or rot == 255 else 90 - rot
    border = cell.border
    if border is not None:
        for key, side in zip(BORDER_KEYS, (border.top, border.bottom, border.left, border.right)):
            if side is not None and side.style:
                style = side.style if side.style in BORDER_STYLES else _XL_BORDER.get(side.style, "thin")
                fmt[key] = [style, _xl_color(side.color) or "#000000"]
    numfmt = _numfmt_from_xlsx(cell)
    if numfmt:
        fmt["num"] = numfmt
    return fmt


def _cf_value_from_xlsx(formula):
    f = str(formula)
    if len(f) >= 2 and f[0] == f[-1] == '"':
        return f[1:-1].replace('""', '"')
    if _NUM_RE.fullmatch(f):
        return f
    return "=" + f


def _cf_rule_from_xlsx(rule):
    style = {}
    dxf = rule.dxf
    if dxf is not None:
        if dxf.fill is not None:
            bg = _xl_color(dxf.fill.bgColor) or _xl_color(dxf.fill.fgColor)
            if bg:
                style["bg"] = bg
        if dxf.font is not None:
            fg = _xl_color(dxf.font.color)
            if fg:
                style["fg"] = fg
            if dxf.font.b:
                style["bold"] = True
    kind = rule.type
    if kind == "cellIs":
        op = {v: k for k, v in _XL_OPS.items()}.get(rule.operator)
        if not op or not rule.formula:
            return None
        spec = {"type": "cell", "op": op, "v1": _cf_value_from_xlsx(rule.formula[0]), "style": style}
        if "between" in op and len(rule.formula) > 1:
            spec["v2"] = _cf_value_from_xlsx(rule.formula[1])
        return spec
    if kind == "containsText" and rule.text:
        return {"type": "text", "text": rule.text, "style": style}
    if kind == "top10":
        return {"type": "bottom" if rule.bottom else "top", "n": int(rule.rank or 10),
                "percent": bool(rule.percent), "style": style}
    if kind == "aboveAverage":
        return {"type": "below" if rule.aboveAverage is False else "above", "style": style}
    if kind in ("duplicateValues", "uniqueValues"):
        return {"type": "dup" if kind == "duplicateValues" else "unique", "style": style}
    if kind == "dataBar" and rule.dataBar is not None:
        return {"type": "databar", "color": _xl_color(rule.dataBar.color) or "#638ec6"}
    if kind == "colorScale" and rule.colorScale is not None:
        colors = [_xl_color(c) for c in rule.colorScale.color]
        if len(colors) in (2, 3) and all(colors):
            return {"type": "scale", "colors": colors}
    return None


_XL_VALIDATION = {"whole": "whole", "decimal": "decimal", "list": "list", "date": "date",
                  "textLength": "length", "custom": "custom"}
_XL_ERROR_STYLES = {"stop": "stop", "warning": "warning", "information": "info"}


def _validation_operand(formula, kind):
    if formula is None:
        return ""
    text = _cf_value_from_xlsx(formula)
    if kind == "date" and _NUM_RE.fullmatch(text):
        return format_date(float(text), "dd/mm/yyyy")
    return text


def _validations_from_xlsx(ws):
    rules = []
    for dv in ws.data_validations.dataValidation:
        kind = _XL_VALIDATION.get(dv.type)
        if kind is None:
            continue
        rule = {"type": kind, "blank": bool(dv.allow_blank), "style": _XL_ERROR_STYLES.get(dv.errorStyle, "stop"),
                "error_title": dv.errorTitle or "", "error": dv.error or "",
                "prompt_title": dv.promptTitle or "", "prompt": dv.prompt or ""}
        if kind == "list":
            f = str(dv.formula1 or "")
            rule["source"] = f[1:-1].replace('""', '"') if f.startswith('"') else "=" + f
            rule["dropdown"] = not dv.showDropDown  # Excel's flag is inverted: set means "hide the arrow"
        else:
            rule["op"] = {v: k for k, v in _XL_OPS.items()}.get(dv.operator or "between", "between")
            rule["v1"] = "=" + str(dv.formula1) if kind == "custom" else _validation_operand(dv.formula1, kind)
            if rule["op"] in ("between", "notbetween"):
                rule["v2"] = _validation_operand(dv.formula2, kind)
        for rg in dv.sqref.ranges:
            if rg.min_row <= MAX_ROWS and rg.min_col <= MAX_COLS:
                rect = (rg.min_row - 1, rg.min_col - 1, min(rg.max_row, MAX_ROWS) - 1, min(rg.max_col, MAX_COLS) - 1)
                rules.append(dict(rule, range=rect))
    return rules


def _xl_validation(rule):
    from openpyxl.worksheet.datavalidation import DataValidation

    kind = rule["type"]

    def operand(text):
        text = str(text or "").strip()
        if is_formula(text):
            return text[1:]
        value = parse_literal(text)[0] if text else 0
        return number_text(value) if is_number(value) else '"' + str(value).replace('"', '""') + '"'

    dv = DataValidation(type={v: k for k, v in _XL_VALIDATION.items()}[kind], allow_blank=rule.get("blank", True),
                        showErrorMessage=True, errorStyle={v: k for k, v in _XL_ERROR_STYLES.items()}.get(
                            rule.get("style", "stop"), "stop"),
                        errorTitle=rule.get("error_title") or None, error=rule.get("error") or None,
                        promptTitle=rule.get("prompt_title") or None, prompt=rule.get("prompt") or None,
                        showInputMessage=bool(rule.get("prompt")))
    if kind == "list":
        source = str(rule.get("source", "")).strip()
        dv.formula1 = source[1:] if is_formula(source) else '"' + source.replace('"', '""') + '"'
        dv.showDropDown = not rule.get("dropdown", True)
    elif kind == "custom":
        dv.formula1 = operand(rule.get("v1"))
    else:
        dv.operator = _XL_OPS.get(rule.get("op", "between"), "between")
        dv.formula1 = operand(rule.get("v1"))
        if rule.get("op", "between") in ("between", "notbetween"):
            dv.formula2 = operand(rule.get("v2"))
    dv.add(range_name(*rule["range"]))
    return dv


_EMU_PER_PX = 9525


def _link_from_xlsx(hl):
    target = hl.target or (f"#{hl.location}" if hl.location else "")
    return {"target": target, "tip": hl.tooltip or ""} if target else None


def _xl_link(cell, link):
    from openpyxl.worksheet.hyperlink import Hyperlink

    target = link["target"]
    ref = cell.coordinate
    if target.startswith("#"):
        cell.hyperlink = Hyperlink(ref=ref, location=target[1:], tooltip=link.get("tip") or None)
    else:
        cell.hyperlink = Hyperlink(ref=ref, target=target, tooltip=link.get("tip") or None)


def _xl_picture(picture):
    """An openpyxl image of a picture without needing Pillow (openpyxl's own Image class opens it)."""
    from openpyxl.drawing.image import Image
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D

    class RawImage(Image):
        def __init__(self, data, fmt, size):  # noqa: super().__init__ would need Pillow
            self.ref, self._raw, self.format = None, data, fmt
            self.width, self.height = size
            self.anchor = "A1"

        def _data(self):
            return self._raw

    data = base64.b64decode(picture["image"])
    img = RawImage(data, picture.get("format", "png"), picture["size"])
    r, c, dx, dy = picture["anchor"]
    w, h = picture["size"]
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=c, colOff=dx * _EMU_PER_PX, row=r, rowOff=dy * _EMU_PER_PX),
                               ext=XDRPositiveSize2D(w * _EMU_PER_PX, h * _EMU_PER_PX))
    return img


def _pictures_from_xlsx(ws):
    """Pictures of a sheet (openpyxl only reads them when Pillow is installed)."""
    pictures = []
    for img in getattr(ws, "_images", []):
        try:
            data = img._data()
        except Exception:  # unreadable picture data
            continue
        fmt = (img.format or "png").lower().replace("jpg", "jpeg")
        if fmt not in ("png", "gif", "jpeg"):
            continue
        anchor = img.anchor
        start = getattr(anchor, "_from", None)
        if start is None:
            continue
        ext = getattr(anchor, "ext", None)
        size = ((round(ext.width / _EMU_PER_PX), round(ext.height / _EMU_PER_PX)) if ext is not None and ext.width
                else (img.width, img.height))
        pictures.append({"type": PICTURE, "image": base64.b64encode(data).decode("ascii"), "format": fmt,
                         "id": new_picture_id(), "name": "", "size": (max(8, size[0]), max(8, size[1])),
                         "anchor": (start.row, start.col, round(start.colOff / _EMU_PER_PX),
                                    round(start.rowOff / _EMU_PER_PX))})
    return pictures


def _charts_from_xlsx(ws, sh):
    """Charts whose series all come from their own sheet, each as one block of cells (approximately
    the bounding box of its series names, values and categories)."""
    charts = []
    for xl in getattr(ws, "_charts", []):
        kind = {"BarChart": "bar" if getattr(xl, "barDir", "col") == "bar" else "column", "LineChart": "line",
                "PieChart": "pie"}.get(type(xl).__name__)
        refs = []
        for ser in xl.series:
            for holder in (ser.tx, ser.val, ser.cat):
                ref = holder and (getattr(holder, "strRef", None) or getattr(holder, "numRef", None))
                if ref is not None and ref.f:
                    refs.append(ref.f)
        rects = []
        for f in refs:
            sheet_part, _, addr = f.rpartition("!")
            sheet_name = sheet_part[1:-1].replace("''", "'") if sheet_part.startswith("'") else sheet_part
            rect = parse_range(addr.replace("$", ""))
            if rect is None or (sheet_name and sheet_name != ws.title):
                rects = []
                break
            rects.append(rect)
        if kind is None or not rects:
            continue
        anchor = xl.anchor
        start = getattr(anchor, "_from", None)
        if start is None:
            continue
        ext, end = getattr(anchor, "ext", None), getattr(anchor, "to", None)
        if ext is not None and ext.width:
            size = (round(ext.width / _EMU_PER_PX), round(ext.height / _EMU_PER_PX))
        elif end is not None:
            size = (sum(sh.col_widths.get(c, DEFAULT_COL_WIDTH) for c in range(start.col, end.col))
                    + round((end.colOff - start.colOff) / _EMU_PER_PX),
                    sum(sh.row_heights.get(r, DEFAULT_ROW_HEIGHT) for r in range(start.row, end.row))
                    + round((end.rowOff - start.rowOff) / _EMU_PER_PX))
        else:
            size = CHART_SIZE
        title = ""
        rich = xl.title and xl.title.tx and xl.title.tx.rich
        if rich:
            title = "\n".join("".join(run.t or "" for run in (p.r or [])) for p in rich.p).strip()
        charts.append({"type": kind, "title": title,
                       "range": (min(r[0] for r in rects), min(r[1] for r in rects),
                                 max(r[2] for r in rects), max(r[3] for r in rects)),
                       "anchor": (start.row, start.col, round(start.colOff / _EMU_PER_PX),
                                  round(start.rowOff / _EMU_PER_PX)),
                       "size": (max(60, size[0]), max(40, size[1]))})
    return charts


def _xl_chart(wb, sheet, ws, chart):
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D

    (r1, c1, r2, c2), by_cols, names, cats = wb.chart_layout(sheet, chart)
    names, cats = int(names), int(cats)  # how many header lines to skip
    kind = chart["type"]
    xl = PieChart() if kind == "pie" else LineChart() if kind == "line" else BarChart()
    if kind in ("column", "bar"):
        xl.type = "col" if kind == "column" else "bar"
    if by_cols:
        data = Reference(ws, min_col=c1 + 1 + cats, min_row=r1 + 1, max_col=c2 + 1, max_row=r2 + 1)
        labels = Reference(ws, min_col=c1 + 1, min_row=r1 + 1 + names, max_row=r2 + 1) if cats else None
    else:
        data = Reference(ws, min_col=c1 + 1, min_row=r1 + 1 + cats, max_col=c2 + 1, max_row=r2 + 1)
        labels = Reference(ws, min_col=c1 + 1 + names, max_col=c2 + 1, min_row=r1 + 1) if cats else None
    xl.add_data(data, titles_from_data=bool(names), from_rows=not by_cols)
    if labels is not None:
        xl.set_categories(labels)
    if chart.get("title"):
        xl.title = chart["title"]
    r, c, dx, dy = chart["anchor"]
    w, h = chart["size"]
    xl.anchor = OneCellAnchor(_from=AnchorMarker(col=c, colOff=dx * _EMU_PER_PX, row=r, rowOff=dy * _EMU_PER_PX),
                              ext=XDRPositiveSize2D(w * _EMU_PER_PX, h * _EMU_PER_PX))
    return xl


_XL_FILTER_OPS = {"=": "equal", "<>": "notEqual", ">": "greaterThan", ">=": "greaterThanOrEqual",
                  "<": "lessThan", "<=": "lessThanOrEqual"}
_XL_DYNAMIC = {kind: kind.replace("week", "Week").replace("month", "Month").replace("year", "Year")
               for kind in DYNAMIC_FILTERS}


def _filter_from_xlsx(fc, book):
    """A filter column's criterion from Excel's filterColumn, or None for kinds PyExcel lacks."""
    if fc.filters is not None:
        allowed = {str(v) for v in fc.filters.filter}
        if fc.filters.blank:
            allowed.add("")
        return allowed
    if fc.customFilters is not None:
        conds = []
        for cf in fc.customFilters.customFilter:
            op = {v: k for k, v in _XL_FILTER_OPS.items()}.get(cf.operator or "equal", "=")
            val = str(cf.val if cf.val is not None else "")
            core = val.strip("*")
            if op in ("=", "<>") and len(val) > 1 and "*" not in core and "?" not in core and val != core:
                kind = "contains" if val.startswith("*") and val.endswith("*") else "ends" if val.startswith("*") else "begins"
                op, val = (kind if op == "=" else "n" + kind), core
            conds.append([op, val])
        return {"type": "custom", "join": "and" if fc.customFilters._and else "or", "conds": conds}
    if fc.top10 is not None:
        return {"type": "top", "n": int(fc.top10.val or 10), "bottom": fc.top10.top is False,
                "percent": bool(fc.top10.percent)}
    if fc.dynamicFilter is not None:
        kind = fc.dynamicFilter.type
        if kind in ("aboveAverage", "belowAverage"):
            return {"type": "average", "above": kind == "aboveAverage"}
        kind = {v: k for k, v in _XL_DYNAMIC.items()}.get(kind)
        return {"type": "dynamic", "kind": kind} if kind else None
    if fc.colorFilter is not None and fc.colorFilter.dxfId is not None:
        try:
            dxf = book._differential_styles[fc.colorFilter.dxfId]
        except (IndexError, TypeError):
            return None
        cell = fc.colorFilter.cellColor is not False
        color = (_xl_color(dxf.fill.bgColor) or _xl_color(dxf.fill.fgColor)) if cell and dxf.fill is not None else \
            _xl_color(dxf.font.color) if not cell and dxf.font is not None else None
        return {"type": "color", "what": "bg" if cell else "fg", "color": color}
    return None


def _xl_filter_column(col_id, crit, book):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.styles.differential import DifferentialStyle
    from openpyxl.worksheet.filters import (ColorFilter, CustomFilter, CustomFilters, DynamicFilter, FilterColumn,
                                            Filters, Top10)

    if not isinstance(crit, dict):
        return FilterColumn(colId=col_id, filters=Filters(filter=sorted(v for v in crit if v), blank="" in crit or None))
    kind = crit["type"]
    if kind == "custom":
        items = []
        for op, val in crit["conds"][:2]:  # Excel keeps two conditions
            val = str(val)
            negate = op.startswith("n") and op[1:] in _TEXT_OPS
            if (op[1:] if negate else op) in _TEXT_OPS:
                val, op = _TEXT_OPS[op[1:] if negate else op].format(val), "<>" if negate else "="
            else:
                lit = parse_literal(val)[0] if val.strip() else ""
                val = number_text(lit) if is_number(lit) else val
            items.append(CustomFilter(operator=_XL_FILTER_OPS[op], val=val))
        return FilterColumn(colId=col_id, customFilters=CustomFilters(customFilter=items,
                                                                      _and=crit.get("join", "and") == "and" or None))
    if kind == "top":
        return FilterColumn(colId=col_id, top10=Top10(top=not crit.get("bottom") or False, val=int(crit.get("n", 10)),
                                                      percent=bool(crit.get("percent")) or None))
    if kind == "average":
        return FilterColumn(colId=col_id, dynamicFilter=DynamicFilter(
            type="aboveAverage" if crit.get("above", True) else "belowAverage"))
    if kind == "dynamic":
        return FilterColumn(colId=col_id, dynamicFilter=DynamicFilter(type=_XL_DYNAMIC[crit["kind"]]))
    color = _xl_hex(crit.get("color")) or "FFFFFFFF"
    cell = crit.get("what", "bg") == "bg"
    dxf = (DifferentialStyle(fill=PatternFill(fill_type="solid", fgColor=color, bgColor=color)) if cell
           else DifferentialStyle(font=Font(color=color)))
    return FilterColumn(colId=col_id, colorFilter=ColorFilter(dxfId=book._differential_styles.add(dxf),
                                                              cellColor=None if cell else False))


_XL_PAPER = {9: "A4", 8: "A3", 11: "A5", 1: "Letter", 5: "Legal"}


def _page_setup_from_xlsx(ws):
    """The printing options PyExcel keeps (see printing.DEFAULT_SETUP), or None when all are default."""
    ps, setup = ws.page_setup, {}
    if ps.orientation == "landscape":
        setup["orientation"] = "landscape"
    paper = _XL_PAPER.get(int(ps.paperSize)) if ps.paperSize else None
    if paper and paper != "A4":
        setup["paper"] = paper
    m = ws.page_margins
    margins = [m.top, m.bottom, m.left, m.right, m.header, m.footer]
    if margins != [0.75, 0.75, 0.7, 0.7, 0.3, 0.3] and all(v is not None for v in margins):
        setup["margins"] = [float(v) for v in margins]
    props = ws.sheet_properties.pageSetUpPr
    if props is not None and props.fitToPage:
        setup["fit"] = "page" if ps.fitToHeight not in (0, "0") else "width"
    elif ps.scale and int(ps.scale) != 100:
        setup["scale"] = int(ps.scale)
    area = ws.print_area
    if area:
        first = area if isinstance(area, str) else area[0]
        rect = parse_range(first.split("!")[-1].replace("$", ""))
        if rect:
            setup["area"] = rect
    titles = ws.print_title_rows
    if titles:
        parts = titles.replace("$", "").split(":")
        if len(parts) == 2 and parts[0].isdigit() and parts[1].isdigit():
            setup["titles"] = (int(parts[0]) - 1, int(parts[1]) - 1)
    opts = ws.print_options
    for key, on in (("gridlines", opts.gridLines), ("center_h", opts.horizontalCentered),
                    ("center_v", opts.verticalCentered)):
        if on:
            setup[key] = True
    for key, hf in (("header", ws.oddHeader), ("footer", ws.oddFooter)):
        if hf is not None and hf.center is not None and hf.center.text:
            setup[key] = hf.center.text
    for key, brks in (("breaks", ws.row_breaks), ("col_breaks", ws.col_breaks)):
        ids = sorted(b.id for b in brks.brk if b.id)  # break after row/column `id`: our 0-based first of the next
        if ids:
            setup[key] = ids
    return setup or None


def _xl_page_setup(ws, setup):
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.page import PageMargins

    ps = ws.page_setup
    ps.orientation = setup.get("orientation", "portrait")
    ps.paperSize = {v: k for k, v in _XL_PAPER.items()}.get(setup.get("paper", "A4"), 9)
    if setup.get("margins"):
        t, b, lft, rgt, h, f = setup["margins"]
        ws.page_margins = PageMargins(left=lft, right=rgt, top=t, bottom=b, header=h, footer=f)
    if setup.get("fit"):
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ps.fitToWidth, ps.fitToHeight = 1, (1 if setup["fit"] == "page" else 0)
    elif setup.get("scale", 100) != 100:
        ps.scale = int(setup["scale"])
    if setup.get("area"):
        ws.print_area = range_name(*setup["area"])
    if setup.get("titles"):
        ws.print_title_rows = f"{setup['titles'][0] + 1}:{setup['titles'][1] + 1}"
    ws.print_options.gridLines = bool(setup.get("gridlines")) or None
    ws.print_options.horizontalCentered = bool(setup.get("center_h")) or None
    ws.print_options.verticalCentered = bool(setup.get("center_v")) or None
    if setup.get("header"):
        ws.oddHeader.center.text = setup["header"]
    if setup.get("footer"):
        ws.oddFooter.center.text = setup["footer"]
    for r in setup.get("breaks", []):
        ws.row_breaks.append(Break(id=r))
    for c in setup.get("col_breaks", []):
        ws.col_breaks.append(Break(id=c))


def _freeze_from_xlsx(ws):
    view = ws.sheet_view
    pane = view.pane
    if pane is None or pane.state not in ("frozen", "frozenSplit"):
        return None
    origin = parse_address(view.topLeftCell or "A1")
    r0, c0 = origin if origin and len(origin) == 2 else (0, 0)
    ys, xs = int(pane.ySplit or 0), int(pane.xSplit or 0)
    if not (ys or xs) or r0 + ys >= MAX_ROWS or c0 + xs >= MAX_COLS:
        return None
    return (r0, r0 + ys, c0, c0 + xs)


def _names_from_xlsx(book):
    """[(name, refers-to without '=', comment)] of the workbook's defined names (not Excel's own _xlnm ones)."""
    try:
        items = list(book.defined_names.values())  # openpyxl 3.1
    except AttributeError:
        items = list(getattr(book.defined_names, "definedName", []))
    for ws in book.worksheets:  # names scoped to a sheet are used book-wide here
        try:
            items += list(ws.defined_names.values())
        except AttributeError:
            pass
    return [(d.name, d.attr_text, d.comment or "") for d in items
            if d.name and not d.name.startswith("_xlnm") and d.attr_text and not d.hidden]


def _xl_names(book, wb):
    from openpyxl.workbook.defined_name import DefinedName

    for d in wb.names.values():
        item = DefinedName(d["name"], attr_text=d["ref"][1:], comment=d.get("comment") or None)
        try:
            book.defined_names[d["name"]] = item  # openpyxl 3.1
        except TypeError:
            book.defined_names.append(item)


def _load_xlsx(path):
    import openpyxl

    book = openpyxl.load_workbook(path)  # formulas are kept as text
    wb = Workbook(())
    store = book[PIVOT_STORE] if PIVOT_STORE in book.sheetnames else None
    for ws in book.worksheets:
        if ws is store:
            continue
        name = wb.add_sheet(ws.title if wb.check_sheet_name(ws.title) is None else None)
        sh = wb.sheets[name]
        for row in ws.iter_rows(max_row=min(ws.max_row, MAX_ROWS), max_col=min(ws.max_column, MAX_COLS)):
            for cell in row:
                try:
                    pos = (cell.row - 1, cell.column - 1)
                    raw = _raw_from_xlsx(cell)
                    fmt = _format_from_xlsx(cell)
                except (AttributeError, TypeError, ValueError):
                    continue  # exotic styles
                if raw:
                    sh.cells[pos] = raw
                if fmt:
                    sh.formats[pos] = fmt
                if cell.comment is not None and cell.comment.text:
                    sh.notes[pos] = cell.comment.text.strip()
                if cell.hyperlink is not None:
                    link = _link_from_xlsx(cell.hyperlink)
                    if link:
                        sh.links[pos] = link
        for key, dim in ws.column_dimensions.items():
            lo = dim.min or (col_index(key) + 1 if re.fullmatch(r"[A-Z]{1,3}", key) else None)
            if lo is None:
                continue
            for c in range(lo - 1, min(dim.max or lo, MAX_COLS)):
                if dim.width:
                    sh.col_widths[c] = max(MIN_COL_WIDTH, round(dim.width * 7 + 5))
                if dim.hidden:
                    sh.hidden_cols.add(c)
                if dim.outlineLevel:
                    sh.col_levels[c] = min(MAX_OUTLINE, int(dim.outlineLevel))
                if dim.collapsed:
                    sh.collapsed_cols.add(c)
        for idx, dim in ws.row_dimensions.items():
            r = idx - 1
            if r >= MAX_ROWS:
                continue
            if dim.height is not None:
                sh.row_heights[r] = max(1, round(dim.height * 4 / 3))
                sh.custom_rows.add(r)
            if dim.hidden:
                sh.hidden_rows.add(r)
            if dim.outlineLevel:
                sh.row_levels[r] = min(MAX_OUTLINE, int(dim.outlineLevel))
            if dim.collapsed:
                sh.collapsed_rows.add(r)
        sh.set_merges((rg.min_row - 1, rg.min_col - 1, rg.max_row - 1, rg.max_col - 1)
                      for rg in ws.merged_cells.ranges if rg.max_row <= MAX_ROWS and rg.max_col <= MAX_COLS)
        sh.tab_color = _xl_color(ws.sheet_properties.tabColor)
        sh.freeze = _freeze_from_xlsx(ws)
        sh.validations = _validations_from_xlsx(ws)
        sh.charts = _charts_from_xlsx(ws, sh) + _pictures_from_xlsx(ws)
        sh.page_setup = _page_setup_from_xlsx(ws)
        rules = []
        for cf in ws.conditional_formatting:
            for xl_rule in cf.rules:
                spec = _cf_rule_from_xlsx(xl_rule)
                if spec:
                    for rg in cf.sqref.ranges:
                        rect = (rg.min_row - 1, rg.min_col - 1,
                                min(rg.max_row, MAX_ROWS) - 1, min(rg.max_col, MAX_COLS) - 1)
                        rules.append((xl_rule.priority or 0, dict(spec, range=rect)))
        sh.cond_formats = [spec for _, spec in sorted(rules, key=lambda p: -p[0])]  # highest priority last
        af = ws.auto_filter
        rect = parse_range(af.ref.replace("$", "")) if af is not None and af.ref else None
        if rect:
            criteria = {}
            for fc in af.filterColumn:
                crit = _filter_from_xlsx(fc, book)
                if crit is not None:
                    criteria[rect[1] + fc.colId] = crit
            sh.filter = {"range": rect, "criteria": criteria, "hidden": set()}
    if not wb.sheets:
        wb.add_sheet()
    for name, sh in wb.sheets.items():
        if sh.filter:
            wb.apply_filter(name)
            sh.hidden_rows -= sh.filter["hidden"]  # Excel stores filtered-out rows as hidden rows
    if store is not None:
        try:
            saved = json.loads(store["A1"].value or "{}")
        except (TypeError, ValueError):
            saved = {}
        for title, pivots in saved.items():
            if title in wb.sheets:
                wb.sheets[title].pivots = [p for p in map(_pivot_from_dict, pivots) if p]
    for name, ref, comment in _names_from_xlsx(book):
        if wb.check_name(name) is None:
            try:
                wb.define_name(name, "=" + ref, comment)
            except ValueError:
                pass
    titles = [ws.title for ws in book.worksheets if ws is not store]
    names = list(wb.sheets)
    if book.active is not None and book.active.title in titles:
        wb.active = names[titles.index(book.active.title)]
    return wb


def _xl_affix(s):
    """Literal text of an Excel number format, quoting characters Excel can't take bare."""
    out, i = [], 0
    while i < len(s):
        if s[i] == '"':
            j = s.find('"', i + 1)
            j = len(s) - 1 if j < 0 else j
            out.append(s[i:j + 1])
            i = j + 1
        else:
            out.append(s[i] if s[i] in " $%-+()/:" else f'"{s[i]}"')
            i += 1
    return "".join(out)


def _xl_number_format(fmt):
    if fmt == "@":
        return "@"
    out = []
    for section in split_sections(fmt):
        if is_date_format(section) or "?/" in section or _SCI_RE.search(section):
            out.append(section)
            continue
        m = _NUM_CORE_RE.search(section)
        out.append(_xl_affix(section) if not m else
                   _xl_affix(section[:m.start()]) + m.group(0) + _xl_affix(section[m.end():]))
    return ";".join(out) or "General"


def _xl_cf_rule(rule):
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, Rule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.styles.differential import DifferentialStyle

    style = rule.get("style") or {}
    fill = None
    if style.get("bg"):
        color = _xl_hex(style["bg"])
        fill = PatternFill(fill_type="solid", start_color=color, end_color=color, bgColor=color)
    font = Font(color=_xl_hex(style.get("fg")), bold=style.get("bold")) if style.get("fg") or style.get("bold") else None
    kind = rule["type"]

    def operand(text):
        text = str(text)
        if is_formula(text):
            return text[1:]
        value = parse_literal(text)[0]
        return number_text(value) if is_number(value) else '"' + str(value).replace('"', '""') + '"'

    if kind == "cell":
        formulas = [operand(rule["v1"])] + ([operand(rule.get("v2", ""))] if "between" in rule["op"] else [])
        return CellIsRule(operator=_XL_OPS[rule["op"]], formula=formulas, fill=fill, font=font)
    dxf = DifferentialStyle(fill=fill, font=font)
    if kind == "text":
        text = str(rule.get("text", "")).replace('"', '""')
        first = cell_name(rule["range"][0], rule["range"][1])
        return Rule(type="containsText", operator="containsText", text=rule.get("text", ""), dxf=dxf,
                    formula=[f'NOT(ISERROR(SEARCH("{text}",{first})))'])
    if kind in ("top", "bottom"):
        return Rule(type="top10", rank=int(rule.get("n", 10)), percent=bool(rule.get("percent")) or None,
                    bottom=(kind == "bottom") or None, dxf=dxf)
    if kind in ("above", "below"):
        return Rule(type="aboveAverage", aboveAverage=(kind == "above"), dxf=dxf)
    if kind in ("dup", "unique"):
        return Rule(type="duplicateValues" if kind == "dup" else "uniqueValues", dxf=dxf)
    if kind == "databar":
        return DataBarRule(start_type="min", end_type="max", color=_xl_hex(rule.get("color", "#638ec6")))
    if kind == "scale":
        colors = [_xl_hex(c) for c in rule["colors"]]
        if len(colors) == 3:
            return ColorScaleRule(start_type="min", start_color=colors[0], mid_type="percentile", mid_value=50,
                                  mid_color=colors[1], end_type="max", end_color=colors[2])
        return ColorScaleRule(start_type="min", start_color=colors[0], end_type="max", end_color=colors[1])
    return None


def _save_xlsx(wb, path, active=None):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    book = openpyxl.Workbook()
    book.remove(book.active)
    from openpyxl.comments import Comment

    for sh in wb.sheets.values():
        ws = book.create_sheet(sh.name)
        for (r, c), text in sh.notes.items():
            comment = Comment(text, FILE_TAG)
            comment.width, comment.height = 200, max(60, 20 * (text.count("\n") + 2))
            ws.cell(row=r + 1, column=c + 1).comment = comment
        for (r, c), link in sh.links.items():
            _xl_link(ws.cell(row=r + 1, column=c + 1), link)
        for r, c in set(sh.cells) | set(sh.formats):
            cell = ws.cell(row=r + 1, column=c + 1)
            raw = sh.cells.get((r, c))
            fmt = sh.formats.get((r, c)) or {}
            implied = None
            if raw is not None:
                as_text = fmt.get("num") == "@"
                if as_text:
                    value = raw[1:] if raw.startswith("'") else raw
                elif is_formula(raw):
                    value = raw
                else:
                    value, implied = parse_literal(raw)
                cell.value = value
                if isinstance(value, str) and (as_text or not is_formula(raw)):
                    cell.data_type = "s"
            numfmt = fmt.get("num") or implied
            if numfmt:
                cell.number_format = _xl_number_format(numfmt)
            if any(fmt.get(k) for k in ("font", "size", "bold", "italic", "underline", "strike", "fg")):
                underline = fmt.get("underline")
                cell.font = Font(name=fmt.get("font") or "Calibri", size=fmt.get("size") or 11,
                                 bold=bool(fmt.get("bold")), italic=bool(fmt.get("italic")),
                                 underline="double" if underline == "double" else ("single" if underline else None),
                                 strike=bool(fmt.get("strike")), color=_xl_hex(fmt.get("fg")))
            if fmt.get("bg"):
                cell.fill = PatternFill("solid", fgColor=_xl_hex(fmt["bg"]))
            if any(fmt.get(k) for k in ("align", "valign", "wrap", "indent", "rotate")):
                rot = fmt.get("rotate") or 0
                cell.alignment = Alignment(horizontal=fmt.get("align"), vertical=fmt.get("valign"),
                                           wrap_text=bool(fmt.get("wrap")), indent=int(fmt.get("indent") or 0),
                                           text_rotation=rot if rot >= 0 else 90 - rot)
            if any(fmt.get(k) for k in BORDER_KEYS):
                sides = [Side(style=v[0], color=_xl_hex(v[1])) if v else Side() for v in map(fmt.get, BORDER_KEYS)]
                cell.border = Border(top=sides[0], bottom=sides[1], left=sides[2], right=sides[3])
        for c, width in sh.col_widths.items():
            ws.column_dimensions[col_name(c)].width = max(1.0, (width - 5) / 7)
        for c in sh.hidden_cols:
            ws.column_dimensions[col_name(c)].hidden = True
        for r, height in sh.row_heights.items():
            ws.row_dimensions[r + 1].height = round(height * 0.75, 2)
        for r in sh.hidden_rows | (sh.filter["hidden"] if sh.filter else set()):
            ws.row_dimensions[r + 1].hidden = True
        for r, level in sh.row_levels.items():
            ws.row_dimensions[r + 1].outlineLevel = level
        for r in sh.collapsed_rows:
            ws.row_dimensions[r + 1].collapsed = True
        for c, level in sh.col_levels.items():
            ws.column_dimensions[col_name(c)].outlineLevel = level
        for c in sh.collapsed_cols:
            ws.column_dimensions[col_name(c)].collapsed = True
        if sh.row_levels:
            ws.sheet_format.outlineLevelRow = max(sh.row_levels.values())
        if sh.col_levels:
            ws.sheet_format.outlineLevelCol = max(sh.col_levels.values())
        for m in sh.merges:
            ws.merge_cells(start_row=m[0] + 1, start_column=m[1] + 1, end_row=m[2] + 1, end_column=m[3] + 1)
        if sh.tab_color:
            ws.sheet_properties.tabColor = _xl_hex(sh.tab_color)
        if sh.freeze:
            r0, r1, c0, c1 = sh.freeze
            ws.freeze_panes = cell_name(r1 - r0, c1 - c0)  # sets up the panes as if frozen from A1...
            ws.sheet_view.pane.topLeftCell = cell_name(r1, c1)
            if (r0, c0) != (0, 0):
                ws.sheet_view.topLeftCell = cell_name(r0, c0)  # ...then moves them to where they start
        if sh.filter:
            r1, c1, r2, c2 = sh.filter["range"]
            ws.auto_filter.ref = range_name(r1, c1, r2, c2)
            for col, crit in sorted(sh.filter["criteria"].items()):
                ws.auto_filter.filterColumn.append(_xl_filter_column(col - c1, crit, book))
        for rule in sh.validations:
            if rule.get("type", "any") != "any":
                ws.add_data_validation(_xl_validation(rule))
        for chart in sh.charts:
            if chart["type"] == PICTURE:
                picture = _xl_picture(chart)
                if picture is not None:
                    ws.add_image(picture)
            else:
                ws.add_chart(_xl_chart(wb, sh.name, ws, chart))
        if sh.page_setup:
            _xl_page_setup(ws, sh.page_setup)
        for rule in reversed(sh.cond_formats):  # Excel: the first rule added gets the top priority
            xl_rule = _xl_cf_rule(rule)
            if xl_rule is not None:
                ws.conditional_formatting.add(range_name(*rule["range"]), xl_rule)
    pivots = {sh.name: [_pivot_to_dict(p) for p in sh.pivots] for sh in wb.sheets.values() if sh.pivots}
    if pivots:  # Excel shows the reports as plain cells; PyExcel reads the definitions back from here
        store = book.create_sheet(PIVOT_STORE)
        store["A1"] = json.dumps(pivots, ensure_ascii=False)
        store.sheet_state = "veryHidden"
    _xl_names(book, wb)
    if active in wb.sheets:
        book.active = list(wb.sheets).index(active)
    book.save(path)
