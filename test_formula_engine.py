"""Unit tests for the PyExcel formula engine.  Run:  python -m unittest -v"""

import importlib.util
import os
import sys
import tempfile
import unittest

import formula_engine as fe
from formula_engine import CellError, Workbook


def make(cells, sheet="Sheet1", wb=None):
    wb = wb or Workbook()
    for addr, raw in cells.items():
        r, c = fe.parse_address(addr)
        wb.set_raw(sheet, r, c, raw)
    return wb


def val(wb, addr, sheet="Sheet1"):
    r, c = fe.parse_address(addr)
    return wb.get_value(sheet, r, c)


class AddressTests(unittest.TestCase):
    def test_column_names(self):
        self.assertEqual(fe.col_name(0), "A")
        self.assertEqual(fe.col_name(25), "Z")
        self.assertEqual(fe.col_name(26), "AA")
        self.assertEqual(fe.col_index("AB"), 27)
        self.assertEqual(fe.parse_address("b3"), (2, 1))
        self.assertEqual(fe.parse_address("C4:A1"), (0, 0, 3, 2))
        self.assertIsNone(fe.parse_address("A0"))


class LiteralTests(unittest.TestCase):
    def test_literals(self):
        self.assertEqual(fe.parse_literal("42"), (42.0, None))
        self.assertEqual(fe.parse_literal("1,234.5"), (1234.5, "#,##0.00"))
        self.assertEqual(fe.parse_literal("50%"), (0.5, "0%"))
        self.assertEqual(fe.parse_literal("true"), (True, None))
        self.assertEqual(fe.parse_literal("'123"), ("123", None))
        self.assertEqual(fe.parse_literal("hello"), ("hello", None))
        self.assertEqual(fe.parse_literal("nan"), ("nan", None))
        value, fmt = fe.parse_literal("21/10/2026")
        self.assertEqual(fmt, "dd/mm/yyyy")
        self.assertEqual(fe.format_value(value, fmt), "21/10/2026")


class FormulaTests(unittest.TestCase):
    def setUp(self):
        self.wb = make({"A1": "10", "A2": "20", "A3": "30", "B1": "abc", "B2": "", "C1": "2"})

    def ev(self, formula):
        self.wb.set_raw("Sheet1", 99, 0, formula)
        return self.wb.get_value("Sheet1", 99, 0)

    def test_arithmetic_and_precedence(self):
        self.assertEqual(self.ev("=1+2*3"), 7)
        self.assertEqual(self.ev("=(1+2)*3"), 9)
        self.assertEqual(self.ev("=-2^2"), 4)  # Excel: negation binds tighter than ^
        self.assertEqual(self.ev("=2^3^2"), 64)
        self.assertEqual(self.ev("=50%*A1"), 5)
        self.assertEqual(self.ev("=A1/C1+A2"), 25)
        self.assertAlmostEqual(self.ev("=0.1+0.2"), 0.3)

    def test_text_and_comparison(self):
        self.assertEqual(self.ev('="x"&A1&"y"'), "x10y")
        self.assertIs(self.ev("=A1<A2"), True)
        self.assertIs(self.ev('=B1="ABC"'), True)
        self.assertIs(self.ev("=A1<>10"), False)
        self.assertIs(self.ev("=Z50=0"), True)  # blank compares as 0

    def test_errors(self):
        self.assertEqual(self.ev("=A1/0"), CellError(fe.DIV0))
        self.assertEqual(self.ev("=B1+1"), CellError(fe.VALUE))
        self.assertEqual(self.ev("=FOO(1)"), CellError(fe.NAME))
        self.assertEqual(self.ev("=1+"), CellError(fe.SYNTAX))
        self.assertEqual(self.ev("=Nope!A1"), CellError(fe.REF))
        self.assertEqual(self.ev("=SQRT(-1)"), CellError(fe.NUM))
        self.assertEqual(self.ev("=(1/0)+1"), CellError(fe.DIV0))

    def test_aggregates(self):
        self.assertEqual(self.ev("=SUM(A1:A3)"), 60)
        self.assertEqual(self.ev("=SUM(A1:B3, 5)"), 65)  # text in ranges is ignored
        self.assertEqual(self.ev("=AVERAGE(A1:A3)"), 20)
        self.assertEqual(self.ev("=MIN(A1:A3)"), 10)
        self.assertEqual(self.ev("=MAX(A1:A3)"), 30)
        self.assertEqual(self.ev("=COUNT(A1:C3)"), 4)
        self.assertEqual(self.ev("=COUNTA(A1:C3)"), 5)
        self.assertEqual(self.ev("=COUNTBLANK(A1:C3)"), 4)
        self.assertEqual(self.ev("=PRODUCT(A1,C1)"), 20)
        self.assertEqual(self.ev("=MEDIAN(A1:A3)"), 20)
        self.assertEqual(self.ev("=SUM(A1:A99,C1:C10000)"), 62)  # large sparse range
        self.assertEqual(self.ev("=SUM(A1:A10000)"), CellError(fe.CIRC))  # includes its own cell

    def test_math(self):
        self.assertEqual(self.ev("=ROUND(2.675,2)"), 2.68)
        self.assertEqual(self.ev("=ROUND(-2.5,0)"), -3)
        self.assertEqual(self.ev("=ROUND(1234,-2)"), 1200)
        self.assertEqual(self.ev("=ROUNDUP(1.21,1)"), 1.3)
        self.assertEqual(self.ev("=ROUNDDOWN(1.29,1)"), 1.2)
        self.assertEqual(self.ev("=INT(-1.5)"), -2)
        self.assertEqual(self.ev("=MOD(-7,3)"), 2)
        self.assertEqual(self.ev("=POWER(2,10)"), 1024)
        self.assertEqual(self.ev("=ABS(-3)"), 3)

    def test_logic(self):
        self.assertEqual(self.ev('=IF(A1>5,"lớn","nhỏ")'), "lớn")
        self.assertEqual(self.ev("=IF(A1>50,1)"), False)
        self.assertEqual(self.ev("=IF(TRUE,,1)"), 0)
        self.assertEqual(self.ev('=IFERROR(A1/0,"lỗi")'), "lỗi")
        self.assertIs(self.ev("=AND(A1>5,A2>5)"), True)
        self.assertIs(self.ev("=OR(A1>50,A2>50)"), False)
        self.assertIs(self.ev("=NOT(FALSE)"), True)
        self.assertIs(self.ev("=ISERROR(1/0)"), True)
        self.assertIs(self.ev("=ISBLANK(Z1)"), True)
        self.assertIs(self.ev("=ISNUMBER(A1)"), True)
        self.assertEqual(self.ev('=IFS(A1>20,"a",A1>5,"b")'), "b")
        self.assertEqual(self.ev("=IF(A1>5;1;2)"), 1)  # ';' separators work too

    def test_text_functions(self):
        self.assertEqual(self.ev('=CONCATENATE("a",A1,"b")'), "a10b")
        self.assertEqual(self.ev('=CONCAT(A1:A3)'), "102030")
        self.assertEqual(self.ev('=TEXTJOIN("-",TRUE,A1:B2)'), "10-abc-20")
        self.assertEqual(self.ev('=LEN("Xin chào")'), 8)
        self.assertEqual(self.ev('=UPPER(B1)'), "ABC")
        self.assertEqual(self.ev('=LEFT("Hello",2)&RIGHT("Hello",3)&MID("Hello",2,3)'), "Helloell")
        self.assertEqual(self.ev('=TRIM("  a   b ")'), "a b")
        self.assertEqual(self.ev('=SUBSTITUTE("a-b-c","-","+")'), "a+b+c")
        self.assertEqual(self.ev('=FIND("l","Hello")'), 3)
        self.assertEqual(self.ev('=TEXT(1234.5,"#,##0.00")'), "1,234.50")
        self.assertEqual(self.ev('=VALUE("12.5")'), 12.5)
        self.assertEqual(self.ev('="Total: "&1/4'), "Total: 0.25")

    def test_dates(self):
        self.assertEqual(self.ev("=DATE(2026,10,21)"), fe.parse_literal("21/10/2026")[0])
        self.assertEqual(self.ev("=YEAR(DATE(2026,10,21))"), 2026)
        self.assertEqual(self.ev("=MONTH(DATE(2026,14,1))"), 2)
        self.assertEqual(self.ev("=WEEKDAY(DATE(2026,10,18))"), 1)  # a Sunday
        self.wb.set_raw("Sheet1", 99, 0, "=DATE(2026,10,21)")
        self.assertEqual(fe.format_value(*self.wb.cell_display("Sheet1", 99, 0)), "21/10/2026")

    def test_lookup(self):
        wb = make({"A1": "apple", "B1": "1", "A2": "banana", "B2": "2", "A3": "cherry", "B3": "3",
                   "D1": "10", "D2": "20", "D3": "30"})
        wb.set_raw("Sheet1", 9, 0, '=VLOOKUP("banana",A1:B3,2,FALSE)')
        self.assertEqual(val(wb, "A10"), 2)
        wb.set_raw("Sheet1", 9, 0, "=VLOOKUP(25,D1:D3,1)")
        self.assertEqual(val(wb, "A10"), 20)
        wb.set_raw("Sheet1", 9, 0, '=VLOOKUP("kiwi",A1:B3,2,FALSE)')
        self.assertEqual(val(wb, "A10"), CellError(fe.NA))
        wb.set_raw("Sheet1", 9, 0, '=INDEX(A1:B3,3,2)')
        self.assertEqual(val(wb, "A10"), 3)
        wb.set_raw("Sheet1", 9, 0, '=MATCH("cherry",A1:A3,0)')
        self.assertEqual(val(wb, "A10"), 3)

    def test_conditional(self):
        wb = make({"A1": "x", "B1": "5", "A2": "y", "B2": "15", "A3": "x", "B3": "25", "A4": "xyz", "B4": "1"})
        cases = {
            '=COUNTIF(B1:B4,">10")': 2,
            '=SUMIF(A1:A4,"x",B1:B4)': 30,
            '=SUMIF(A1:A4,"x*",B1:B4)': 31,
            '=AVERAGEIF(A1:A4,"x",B1:B4)': 15,
            '=COUNTIFS(A1:A4,"x",B1:B4,">10")': 1,
            '=SUMIFS(B1:B4,A1:A4,"<>y")': 31,
            '=SUMPRODUCT(B1:B2,B3:B4)': 140,
        }
        for formula, expected in cases.items():
            wb.set_raw("Sheet1", 9, 0, formula)
            self.assertEqual(val(wb, "A10"), expected, formula)

    def test_more_math_and_statistics(self):
        wb = make({"A1": "3", "A2": "1", "A3": "4", "A4": "1", "A5": "5", "A6": "x"})
        cases = {
            "=LARGE(A1:A6,2)": 4, "=SMALL(A1:A6,1)": 1, "=LARGE(A1:A6,9)": CellError(fe.NUM),
            "=RANK(4,A1:A5)": 2, "=RANK(4,A1:A5,1)": 4, "=RANK.AVG(1,A1:A5,1)": 1.5, "=RANK(9,A1:A5)": CellError(fe.NA),
            "=MODE(A1:A5)": 1, "=PERCENTILE(A1:A5,0.5)": 3, "=QUARTILE(A1:A5,1)": 1, "=QUARTILE(A1:A5,4)": 5,
            "=VARP(A1:A5)": 2.56, "=SUMSQ(A1:A3)": 26, "=MROUND(7,3)": 6, "=MROUND(-7,3)": CellError(fe.NUM),
            "=QUOTIENT(-7,2)": -3, "=GCD(12,18)": 6, "=LCM(4,6)": 12, "=FACT(5)": 120, "=COMBIN(5,2)": 10,
            "=EVEN(1.5)": 2, "=EVEN(-1)": -2, "=ODD(2)": 3, "=ODD(0)": 1, "=DEGREES(PI())": 180,
            "=SIN(PI()/2)": 1, "=ATAN2(1,1)*4": 3.141592653589793,
        }
        for formula, expected in cases.items():
            wb.set_raw("Sheet1", 9, 0, formula)
            if isinstance(expected, CellError):
                self.assertEqual(val(wb, "A10"), expected, formula)
            else:
                self.assertAlmostEqual(val(wb, "A10"), expected, msg=formula)

    def test_conditional_ifs(self):
        wb = make({"A1": "x", "B1": "5", "A2": "y", "B2": "15", "A3": "x", "B3": "25", "A4": "x", "B4": "1"})
        cases = {
            '=AVERAGEIFS(B1:B4,A1:A4,"x",B1:B4,">2")': 15,
            '=MAXIFS(B1:B4,A1:A4,"x")': 25, '=MINIFS(B1:B4,A1:A4,"x")': 1, '=MAXIFS(B1:B4,A1:A4,"z")': 0,
            '=AVERAGEIFS(B1:B4,A1:A4,"z")': CellError(fe.DIV0),
        }
        for formula, expected in cases.items():
            wb.set_raw("Sheet1", 9, 0, formula)
            self.assertEqual(val(wb, "A10"), expected, formula)

    def test_subtotal_skips_hidden_rows_and_subtotals(self):
        wb = make({"A1": "1", "A2": "2", "A3": "4", "A4": "=SUBTOTAL(9,A1:A3)", "A5": "8"})
        wb.set_raw("Sheet1", 9, 0, "=SUBTOTAL(9,A1:A5)")
        self.assertEqual(val(wb, "A10"), 15)  # A4 is a subtotal, so it is not counted twice
        sh = wb.sheets["Sheet1"]
        sh.hidden_rows.add(1)
        wb.invalidate()
        self.assertEqual(val(wb, "A10"), 15)  # code 9 keeps rows hidden by hand
        wb.set_raw("Sheet1", 9, 0, "=SUBTOTAL(109,A1:A5)")
        self.assertEqual(val(wb, "A10"), 13)
        sh.filter = {"range": (0, 0, 4, 0), "criteria": {}, "hidden": {2}}
        wb.set_raw("Sheet1", 9, 0, "=SUBTOTAL(9,A1:A5)")
        self.assertEqual(val(wb, "A10"), 11)  # rows hidden by a filter are always skipped
        wb.set_raw("Sheet1", 9, 0, "=SUBTOTAL(2,A1:A5)")
        self.assertEqual(val(wb, "A10"), 3)

    def test_xlookup_and_friends(self):
        wb = make({"A1": "apple", "B1": "10", "C1": "a", "A2": "banana", "B2": "20", "C2": "b",
                   "A3": "cherry", "B3": "30", "C3": "c", "A4": "banana", "B4": "40", "C4": "d",
                   "E1": "1", "F1": "2", "G1": "3", "E2": "x", "F2": "y", "G2": "z"})
        cases = {
            '=XLOOKUP("banana",A1:A4,B1:B4)': 20,
            '=XLOOKUP("banana",A1:A4,B1:B4,,0,-1)': 40,
            '=XLOOKUP("kiwi",A1:A4,B1:B4)': CellError(fe.NA),
            '=XLOOKUP("kiwi",A1:A4,B1:B4,"không có")': "không có",
            '=XLOOKUP(25,B1:B4,A1:A4,,1)': "cherry",
            '=XLOOKUP(25,B1:B4,A1:A4,,-1)': "banana",
            '=XLOOKUP("ch*",A1:A4,C1:C4,,2)': "c",
            '=SUM(XLOOKUP("cherry",A1:A4,B1:C4))': 30,  # a whole row comes back
            "=XLOOKUP(2,E1:G1,E2:G2)": "y",
            "=XLOOKUP(2,E1:G1,E2:F2)": CellError(fe.VALUE),
            '=XMATCH("cherry",A1:A4)': 3, '=XMATCH("banana",A1:A4,0,-1)': 4,
            "=SUM(OFFSET(A1,1,1,2,1))": 50, "=OFFSET(B1,3,0)": 40, "=OFFSET(A1,-1,0)": CellError(fe.REF),
            '=INDIRECT("B"&2)': 20, '=SUM(INDIRECT("B1:B3"))': 60, '=INDIRECT("R3C2",FALSE)': 30,
            '=INDIRECT("nope")': CellError(fe.REF),
            '=SWITCH(C2,"a",1,"b",2)': 2, '=SWITCH(C4,"a",1,"b",2,"khác")': "khác", '=SWITCH(9,1,2)': CellError(fe.NA),
            '=ISFORMULA(A1)': False, '=ISEVEN(4)': True, '=ISODD(-3)': True, '=ISNONTEXT(B1)': True,
            '=IFNA(NA(),"ok")': "ok",
            '=TEXTBEFORE("a-b-c","-")': "a", '=TEXTAFTER("a-b-c","-")': "b-c", '=TEXTBEFORE("a-b-c","-",-1)': "a-b",
            '=TEXTAFTER("a-b-c","-",2)': "c", '=TEXTAFTER("a-b","+")': CellError(fe.NA),
        }
        for formula, expected in cases.items():
            wb.set_raw("Sheet1", 9, 0, formula)
            self.assertEqual(val(wb, "A10"), expected, formula)
        wb.set_raw("Sheet1", 9, 1, '=INDIRECT("A10")')  # follows the formula above
        self.assertEqual(val(wb, "B10"), CellError(fe.NA))

    def test_more_dates(self):
        d = fe.parse_literal

        def ev(formula):
            self.wb.set_raw("Sheet1", 99, 0, formula)
            return self.wb.get_value("Sheet1", 99, 0)

        def text(formula):
            ev(formula)
            return fe.format_value(*self.wb.cell_display("Sheet1", 99, 0))

        self.assertEqual(text("=EDATE(DATE(2026,1,31),1)"), "28/02/2026")
        self.assertEqual(text("=EDATE(DATE(2026,3,15),-3)"), "15/12/2025")
        self.assertEqual(text("=EOMONTH(DATE(2024,1,10),1)"), "29/02/2024")
        self.assertEqual(text("=WORKDAY(DATE(2026,10,23),1)"), "26/10/2026")  # Friday -> Monday
        self.assertEqual(ev("=WORKDAY(DATE(2026,10,26),-1)"), d("23/10/2026")[0])
        self.assertEqual(ev("=NETWORKDAYS(DATE(2026,10,1),DATE(2026,10,31))"), 22)
        self.assertEqual(ev("=NETWORKDAYS(DATE(2026,10,31),DATE(2026,10,1))"), -22)
        self.wb.set_raw("Sheet1", 0, 4, "02/10/2026")
        self.assertEqual(ev("=NETWORKDAYS(DATE(2026,10,1),DATE(2026,10,31),E1)"), 21)
        self.assertEqual(ev("=WORKDAY(DATE(2026,10,1),1,E1)"), d("05/10/2026")[0])
        self.assertEqual(ev('=DATEDIF(DATE(2000,5,20),DATE(2026,10,21),"Y")'), 26)
        self.assertEqual(ev('=DATEDIF(DATE(2000,5,20),DATE(2026,10,21),"M")'), 317)
        self.assertEqual(ev('=DATEDIF(DATE(2000,5,20),DATE(2026,10,21),"YM")'), 5)
        self.assertEqual(ev('=DATEDIF(DATE(2000,5,20),DATE(2026,10,21),"MD")'), 1)
        self.assertEqual(ev('=DATEDIF(DATE(2026,1,1),DATE(2026,3,1),"D")'), 59)
        self.assertEqual(ev('=DATEDIF(DATE(2026,3,1),DATE(2026,1,1),"D")'), CellError(fe.NUM))
        self.assertEqual(ev('=DATEDIF(DATE(2025,12,25),DATE(2026,1,5),"YD")'), 11)
        self.assertEqual(text("=TIME(14,30,0)"), "14:30")
        self.assertEqual(ev("=SECOND(TIME(1,2,3))"), 3)
        self.assertEqual(ev('=DATEVALUE("21/10/2026")'), d("21/10/2026")[0])
        self.assertEqual(ev('=DATEVALUE("abc")'), CellError(fe.VALUE))
        self.assertEqual(ev("=WEEKNUM(DATE(2026,1,3))"), 1)   # Saturday of the first week
        self.assertEqual(ev("=WEEKNUM(DATE(2026,1,4))"), 2)   # Sunday starts week 2
        self.assertEqual(ev("=WEEKNUM(DATE(2026,1,4),2)"), 1)  # weeks starting on Monday
        self.assertEqual(ev("=ISOWEEKNUM(DATE(2026,10,21))"), 43)

    def test_every_function_is_documented(self):
        from function_info import CATEGORIES, FUNCTION_INFO
        self.assertEqual(set(fe.function_names()) - set(FUNCTION_INFO), set())
        self.assertTrue(all(info[0] in CATEGORIES for info in FUNCTION_INFO.values()))

    def test_circular(self):
        wb = make({"A1": "=B1+1", "B1": "=A1+1", "C1": "=C1"})
        self.assertEqual(val(wb, "A1"), CellError(fe.CIRC))
        self.assertEqual(val(wb, "B1"), CellError(fe.CIRC))
        self.assertEqual(val(wb, "C1"), CellError(fe.CIRC))

    def test_recalculation_after_edit(self):
        wb = make({"A1": "1", "A2": "=A1*10", "A3": "=A2+1"})
        self.assertEqual(val(wb, "A3"), 11)
        wb.set_raw("Sheet1", 0, 0, "5")
        self.assertEqual(val(wb, "A3"), 51)

    def test_long_chain(self):
        wb = Workbook()
        wb.set_raw("Sheet1", 0, 0, "1")
        for r in range(1, 1500):
            wb.set_raw("Sheet1", r, 0, f"=A{r}+1")
        self.assertEqual(wb.get_value("Sheet1", 1499, 0), 1500)

    def test_cross_sheet(self):
        wb = Workbook(("Sheet1", "Data Sheet"))
        make({"A1": "7"}, sheet="Data Sheet", wb=wb)
        make({"A1": "='Data Sheet'!A1*2", "A2": "=SUM('data sheet'!A1:A5)"}, wb=wb)
        self.assertEqual(val(wb, "A1"), 14)
        self.assertEqual(val(wb, "A2"), 7)
        wb.rename_sheet("Data Sheet", "Data")
        self.assertEqual(wb.get_raw("Sheet1", 0, 0), "=Data!A1*2")
        self.assertEqual(val(wb, "A1"), 14)


class RewriteTests(unittest.TestCase):
    def test_normalize(self):
        self.assertEqual(fe.normalize_formula("=sum(a1:b2;3"), "=SUM(A1:B2,3)")
        self.assertEqual(fe.normalize_formula('="abc"&a1'), '="abc"&A1')
        self.assertEqual(fe.normalize_formula("hello"), "hello")

    def test_shift(self):
        self.assertEqual(fe.shift_formula("=A1+$B$1+B$2+$C3", 2, 1), "=B3+$B$1+C$2+$C5")
        self.assertEqual(fe.shift_formula("=SUM(A1:A3)", 1, 0), "=SUM(A2:A4)")
        self.assertEqual(fe.shift_formula("=A1", -1, 0), "=#REF!")
        self.assertEqual(fe.shift_formula('="A1"&A1', 0, 1), '="A1"&B1')
        self.assertEqual(fe.shift_formula("=Sheet2!A1", 1, 0), "=Sheet2!A2")

    def test_insert_delete_rows(self):
        wb = make({"A1": "1", "A2": "2", "A3": "3", "B1": "=SUM(A1:A3)", "B2": "=A3*2"})
        wb.insert_rows("Sheet1", 1)
        self.assertEqual(wb.get_raw("Sheet1", 0, 1), "=SUM(A1:A4)")
        self.assertEqual(wb.get_raw("Sheet1", 2, 1), "=A4*2")
        self.assertEqual(val(wb, "B1"), 6)
        wb.delete_rows("Sheet1", 3)  # removes the row holding 3
        self.assertEqual(wb.get_raw("Sheet1", 0, 1), "=SUM(A1:A3)")
        self.assertEqual(wb.get_raw("Sheet1", 2, 1), "=#REF!*2")
        self.assertEqual(val(wb, "B1"), 3)

    def test_insert_delete_cols(self):
        wb = make({"A1": "1", "B1": "2", "C1": "=A1+B1"})
        wb.insert_cols("Sheet1", 1)
        self.assertEqual(wb.get_raw("Sheet1", 0, 3), "=A1+C1")
        wb.delete_cols("Sheet1", 0)
        self.assertEqual(wb.get_raw("Sheet1", 0, 2), "=#REF!+B1")

    def test_fill_series(self):
        self.assertEqual(fe.fill_series(["1", "3"])(4), "9")
        self.assertIsNone(fe.fill_series(["5"]))
        self.assertEqual(fe.fill_series(["Item 1"])(2), "Item 3")
        self.assertEqual(fe.fill_series(["21/10/2026"])(11), "01/11/2026")
        self.assertIsNone(fe.fill_series(["=A1"]))

    def test_formats(self):
        self.assertEqual(fe.format_number(1234567.891, "#,##0.00"), "1,234,567.89")
        self.assertEqual(fe.format_number(0.256, "0.0%"), "25.6%")
        self.assertEqual(fe.format_number(-1500, "#,##0 ₫"), "-1,500 ₫")
        self.assertEqual(fe.format_number(1.5, "$#,##0.00"), "$1.50")
        self.assertEqual(fe.format_general(1 / 3), "0.3333333333")
        self.assertEqual(fe.format_general(1e20), "1E+20")
        self.assertEqual(fe.adjust_decimals(None, 1.5, 1), "0.00")
        self.assertEqual(fe.adjust_decimals("#,##0.00", 1, -1), "#,##0.0")


class FileTests(unittest.TestCase):
    def roundtrip(self, ext):
        wb = Workbook(("Sheet1", "Khác"))
        make({"A1": "Tên", "B1": "12.5", "A2": "=B1*2", "A3": "21/10/2026", "A4": "'007"}, wb=wb)
        wb.set_format("Sheet1", 0, 0, {"bold": True, "bg": "#ffff00", "align": "center"})
        wb.sheets["Sheet1"].col_widths[0] = 150
        make({"A1": "=Sheet1!B1+1"}, sheet="Khác", wb=wb)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "test" + ext)
            fe.save_workbook(wb, path, active="Khác")
            return fe.load_workbook(path)

    def test_json(self):
        wb = self.roundtrip(".pyxl")
        self.assertEqual(list(wb.sheets), ["Sheet1", "Khác"])
        self.assertEqual(wb.active, "Khác")
        self.assertEqual(val(wb, "A2"), 25)
        self.assertEqual(val(wb, "A1", "Khác"), 13.5)
        self.assertEqual(wb.get_format("Sheet1", 0, 0)["bg"], "#ffff00")
        self.assertEqual(wb.sheets["Sheet1"].col_widths[0], 150)

    @unittest.skipUnless(fe.HAS_XLSX, "openpyxl not installed")
    def test_xlsx(self):
        wb = self.roundtrip(".xlsx")
        self.assertEqual(list(wb.sheets), ["Sheet1", "Khác"])
        self.assertEqual(wb.active, "Khác")
        self.assertEqual(wb.get_raw("Sheet1", 1, 0), "=B1*2")
        self.assertEqual(val(wb, "A2"), 25)
        self.assertEqual(val(wb, "A1", "Khác"), 13.5)
        self.assertEqual(val(wb, "A4"), "007")
        self.assertEqual(fe.format_value(*wb.cell_display("Sheet1", 2, 0)), "21/10/2026")
        fmt = wb.get_format("Sheet1", 0, 0)
        self.assertTrue(fmt["bold"])
        self.assertEqual(fmt["bg"], "#ffff00")
        self.assertEqual(fmt["align"], "center")
        self.assertEqual(wb.sheets["Sheet1"].col_widths[0], 150)

    def test_csv(self):
        wb = self.roundtrip(".csv")
        self.assertEqual(list(wb.sheets), ["Sheet1"])  # CSV holds only the active sheet's values
        self.assertEqual(val(wb, "A1"), 13.5)


class NumberFormatTests(unittest.TestCase):
    def test_sections_and_special_formats(self):
        acc = '#,##0 ₫;-#,##0 ₫;"-" ₫'
        self.assertEqual(fe.format_number(1500, acc), "1,500 ₫")
        self.assertEqual(fe.format_number(-1500, acc), "-1,500 ₫")
        self.assertEqual(fe.format_number(0, acc), "- ₫")
        self.assertEqual(fe.format_number(-3, "#,##0;(#,##0)"), "(3)")
        self.assertEqual(fe.format_number(12345.678, "0.00E+00"), "1.23E+04")
        self.assertEqual(fe.format_number(0.00012, "0.0E+00"), "1.2E-04")
        self.assertEqual(fe.format_number(1.25, "# ?/?"), "1 1/4")
        self.assertEqual(fe.format_number(0.5, "# ?/?"), "1/2")
        serial = fe.parse_literal("21/10/2026")[0]
        self.assertEqual(fe.format_number(serial, 'dddd, "ngày" dd "tháng" mm "năm" yyyy'),
                         "Thứ Tư, ngày 21 tháng 10 năm 2026")
        self.assertEqual(fe.format_number(serial + 0.5, "hh:mm:ss"), "12:00:00")
        self.assertEqual(fe.adjust_decimals('#,##0 ₫;-#,##0 ₫;"-" ₫', 1, 1), '#,##0.0 ₫;-#,##0.0 ₫;"-" ₫')

    def test_text_format_keeps_typed_text(self):
        wb = make({"A1": "007", "A2": "=1+1"})
        wb.set_format("Sheet1", 0, 0, {"num": "@"})
        wb.set_format("Sheet1", 1, 0, {"num": "@"})
        self.assertEqual(val(wb, "A1"), "007")
        self.assertEqual(val(wb, "A2"), "=1+1")
        wb.set_format("Sheet1", 0, 0, None)
        self.assertEqual(val(wb, "A1"), 7)

    def test_value_to_raw(self):
        self.assertEqual(fe.value_to_raw(3.0), "3")
        self.assertEqual(fe.value_to_raw("12"), "'12")
        self.assertEqual(fe.value_to_raw("abc"), "abc")
        self.assertEqual(fe.value_to_raw(True), "TRUE")
        self.assertIsNone(fe.value_to_raw(""))


class LayoutTests(unittest.TestCase):
    def test_merges_follow_row_and_column_changes(self):
        wb = make({"A1": "Tiêu đề"})
        sh = wb.sheets["Sheet1"]
        sh.set_merges([(0, 0, 0, 2), (4, 0, 5, 1)])
        self.assertEqual(sh.merge_at(0, 1), (0, 0, 0, 2))
        wb.insert_rows("Sheet1", 1)
        self.assertEqual(sh.merges, [(0, 0, 0, 2), (5, 0, 6, 1)])
        wb.insert_cols("Sheet1", 1)
        self.assertEqual(sh.merges, [(0, 0, 0, 3), (5, 0, 6, 2)])
        wb.delete_rows("Sheet1", 5, 2)
        self.assertEqual(sh.merges, [(0, 0, 0, 3)])

    def test_insert_and_delete_cells_shift_only_the_band(self):
        wb = make({"A1": "1", "A2": "2", "B1": "10", "B2": "20", "C1": "=SUM(A1:A2)", "C2": "=B2", "D1": "=A2"})
        wb.insert_cells("Sheet1", 0, 0, 0, 0, "down")  # only column A moves down
        self.assertEqual(wb.get_raw("Sheet1", 1, 0), "1")
        self.assertEqual(wb.get_raw("Sheet1", 0, 1), "10")  # column B untouched
        self.assertEqual(wb.get_raw("Sheet1", 0, 2), "=SUM(A2:A3)")
        self.assertEqual(wb.get_raw("Sheet1", 1, 2), "=B2")
        self.assertEqual(wb.get_raw("Sheet1", 0, 3), "=A3")
        wb.delete_cells("Sheet1", 0, 1, 0, 1, "left")  # B1 removed, C1.. shift left on row 1 only
        self.assertEqual(wb.get_raw("Sheet1", 0, 1), "=SUM(A2:A3)")
        self.assertEqual(wb.get_raw("Sheet1", 1, 1), "20")
        self.assertEqual(val(wb, "B1"), 3)

    def test_sizes_and_hidden_follow_rows(self):
        wb = Workbook()
        sh = wb.sheets["Sheet1"]
        sh.row_heights, sh.custom_rows, sh.hidden_rows = {3: 40}, {3}, {5}
        sh.hidden_cols, sh.col_widths = {2}, {2: 120}
        wb.insert_rows("Sheet1", 0, 2)
        wb.delete_cols("Sheet1", 0)
        self.assertEqual((sh.row_heights, sh.custom_rows, sh.hidden_rows), ({5: 40}, {5}, {7}))
        self.assertEqual((sh.hidden_cols, sh.col_widths), ({1}, {1: 120}))

    def test_freeze_follows_row_and_column_changes(self):
        wb = Workbook()
        sh = wb.sheets["Sheet1"]
        sh.freeze = (0, 2, 0, 1)
        wb.insert_rows("Sheet1", 1)        # inside the frozen rows: one more row frozen
        self.assertEqual(sh.freeze, (0, 3, 0, 1))
        wb.insert_rows("Sheet1", 5)        # below them: no change
        self.assertEqual(sh.freeze, (0, 3, 0, 1))
        wb.insert_rows("Sheet1", 0)        # at the very top: the new row is frozen too
        self.assertEqual(sh.freeze, (0, 4, 0, 1))
        wb.delete_rows("Sheet1", 0)
        self.assertEqual(sh.freeze, (0, 3, 0, 1))
        wb.delete_cols("Sheet1", 0)        # the only frozen column is gone
        self.assertEqual(sh.freeze, (0, 3, 0, 0))
        wb.delete_rows("Sheet1", 0, 3)
        self.assertIsNone(sh.freeze)

    def test_snapshot_restore_and_copy_sheet(self):
        wb = make({"A1": "5", "B1": "=A1*2"})
        sh = wb.sheets["Sheet1"]
        sh.set_merges([(2, 0, 2, 1)])
        before = wb.snapshot()
        wb.delete_rows("Sheet1", 0, 5)
        self.assertNotEqual(wb.snapshot(), before)
        wb.restore(before)
        self.assertEqual(val(wb, "B1"), 10)
        self.assertEqual(wb.sheets["Sheet1"].merge_at(2, 1), (2, 0, 2, 1))
        new = wb.copy_sheet("Sheet1")
        self.assertEqual(new, "Sheet1 (2)")
        wb.set_raw(new, 0, 0, "7")
        self.assertEqual((val(wb, "B1"), val(wb, "B1", new)), (10, 14))


class ConditionalFormatTests(unittest.TestCase):
    def setUp(self):
        self.wb = make({f"A{i}": str(v) for i, v in enumerate([5, 15, 25, 35, 15, 0], 1)})
        self.wb.set_raw("Sheet1", 6, 0, "chuối")
        red = {"bg": "#ffc7ce", "fg": "#9c0006"}
        self.style = red

    def styles(self, rule):
        rule = dict(rule, range=(0, 0, 6, 0), style=self.style)
        self.wb.set_cond_formats("Sheet1", [rule])
        return [bool(self.wb.cond_format("Sheet1", r, 0)) for r in range(7)]

    def test_rules(self):
        T, F = True, False
        self.assertEqual(self.styles({"type": "cell", "op": ">", "v1": "20"}), [F, F, T, T, F, F, T])
        self.assertEqual(self.styles({"type": "cell", "op": "between", "v1": "10", "v2": "30"}),
                         [F, T, T, F, T, F, F])
        self.assertEqual(self.styles({"type": "text", "text": "UỐ"}), [F, F, F, F, F, F, T])
        self.assertEqual(self.styles({"type": "dup"}), [F, T, F, F, T, F, F])
        self.assertEqual(self.styles({"type": "top", "n": 2}), [F, F, T, T, F, F, F])
        self.assertEqual(self.styles({"type": "bottom", "n": 1}), [F, F, F, F, F, T, F])
        self.assertEqual(self.styles({"type": "above"}), [F, F, T, T, F, F, F])  # average 15.83
        bar = dict(type="databar", color="#638ec6", range=(0, 0, 6, 0))
        self.wb.set_cond_formats("Sheet1", [bar])
        self.assertEqual(self.wb.cond_format("Sheet1", 3, 0)["databar"], (1.0, "#638ec6"))
        self.assertAlmostEqual(self.wb.cond_format("Sheet1", 0, 0)["databar"][0], 5 / 35)
        scale = dict(type="scale", colors=["#f8696b", "#ffeb84", "#63be7b"], range=(0, 0, 6, 0))
        self.wb.set_cond_formats("Sheet1", [scale])
        self.assertEqual(self.wb.cond_format("Sheet1", 5, 0)["bg"], "#f8696b")
        self.assertEqual(self.wb.cond_format("Sheet1", 3, 0)["bg"], "#63be7b")

    def test_rules_follow_values(self):
        self.styles({"type": "cell", "op": ">", "v1": "=A1*4"})  # threshold 20 from a formula
        self.wb.set_raw("Sheet1", 0, 0, "1")                  # threshold now 4
        self.assertTrue(self.wb.cond_format("Sheet1", 1, 0))


class FilterTests(unittest.TestCase):
    def test_filter_hides_rows(self):
        wb = make({"A1": "Loại", "B1": "Số", "A2": "x", "B2": "1", "A3": "y", "B3": "2", "A4": "x", "B4": "3", "A5": ""})
        sh = wb.sheets["Sheet1"]
        sh.filter = {"range": (0, 0, 4, 1), "criteria": {0: {"x"}}, "hidden": set()}
        wb.apply_filter("Sheet1")
        self.assertEqual(sh.filter["hidden"], {2, 4})
        self.assertTrue(sh.row_hidden(2))
        self.assertEqual(wb.filter_values("Sheet1", 0), ["x", "y", ""])
        wb.insert_rows("Sheet1", 0)
        self.assertEqual((sh.filter["range"], sh.filter["hidden"]), ((1, 0, 5, 1), {3, 5}))


class ValidationTests(unittest.TestCase):
    def setUp(self):
        self.wb = make({"E1": "Hà Nội", "E2": "Huế", "E3": "", "E4": "Đà Nẵng", "F1": "10"})
        self.sh = self.wb.sheets["Sheet1"]

    def allowed(self, rule, raw, addr="A1"):
        self.sh.validations = [dict(rule, range=(0, 0, 9, 0))]
        r, c = fe.parse_address(addr)
        return self.wb.validate("Sheet1", r, c, raw) is None

    def test_list_from_text_and_range(self):
        rule = {"type": "list", "source": "Có, Không"}
        self.assertTrue(self.allowed(rule, "có"))
        self.assertFalse(self.allowed(rule, "Có lẽ"))
        rule = {"type": "list", "source": "=$E$1:$E$4"}
        self.assertEqual(self.wb.validation_options("Sheet1", 0, 0, dict(rule, range=(0, 0, 9, 0))),
                         ["Hà Nội", "Huế", "Đà Nẵng"])
        self.assertTrue(self.allowed(rule, "Huế"))
        self.assertFalse(self.allowed(rule, "Sài Gòn"))
        self.assertTrue(self.allowed({"type": "list", "source": "1,2,3"}, "2"))

    def test_numbers_dates_and_length(self):
        whole = {"type": "whole", "op": "between", "v1": "1", "v2": "=F1"}
        self.assertTrue(self.allowed(whole, "10"))
        self.assertFalse(self.allowed(whole, "11"))
        self.assertFalse(self.allowed(whole, "2.5"))
        self.assertFalse(self.allowed(whole, "abc"))
        self.assertTrue(self.allowed(whole, ""))  # blanks are allowed unless "blank" is off
        self.assertFalse(self.allowed(dict(whole, blank=False), ""))
        self.assertTrue(self.allowed({"type": "decimal", "op": ">", "v1": "0"}, "=F1/4"))
        date = {"type": "date", "op": ">=", "v1": "01/01/2026"}
        self.assertTrue(self.allowed(date, "21/10/2026"))
        self.assertFalse(self.allowed(date, "31/12/2025"))
        self.assertTrue(self.allowed({"type": "length", "op": "<=", "v1": "3"}, "abc"))
        self.assertFalse(self.allowed({"type": "length", "op": "<=", "v1": "3"}, "abcd"))

    def test_custom_formula_is_relative_to_each_cell(self):
        rule = {"type": "custom", "v1": "=COUNTIF($A$1:$A$10,A1)=1"}  # no duplicates in the column
        self.wb.set_raw("Sheet1", 0, 0, "x")
        self.assertFalse(self.allowed(rule, "x", "A2"))
        self.assertTrue(self.allowed(rule, "y", "A2"))
        self.assertEqual(self.wb.get_raw("Sheet1", 1, 0), "")  # the trial value is not kept

    def test_rect_minus(self):
        self.assertEqual(fe.rect_minus((0, 0, 9, 0), (20, 0, 30, 0)), [(0, 0, 9, 0)])
        self.assertEqual(fe.rect_minus((0, 0, 9, 0), (3, 0, 4, 5)), [(0, 0, 2, 0), (5, 0, 9, 0)])
        self.assertEqual(fe.rect_minus((0, 0, 2, 2), (1, 1, 1, 1)),
                         [(0, 0, 0, 2), (2, 0, 2, 2), (1, 0, 1, 0), (1, 2, 1, 2)])
        self.assertEqual(fe.rect_minus((0, 0, 2, 2), (0, 0, 5, 5)), [])

    def test_rules_and_notes_follow_inserted_rows(self):
        self.sh.validations = [{"range": (2, 0, 4, 0), "type": "list", "source": "a,b"}]
        self.sh.notes = {(3, 1): "ghi chú"}
        self.wb.insert_rows("Sheet1", 0, 2)
        self.assertEqual(self.sh.validations[0]["range"], (4, 0, 6, 0))
        self.assertEqual(self.sh.notes, {(5, 1): "ghi chú"})
        self.assertIsNotNone(self.wb.validation_at("Sheet1", 5, 0))
        self.assertIsNone(self.wb.validation_at("Sheet1", 2, 0))

    def roundtrip(self, ext):
        self.sh.validations = [
            {"range": (0, 0, 9, 0), "type": "list", "source": "Có,Không", "dropdown": True, "blank": True,
             "style": "stop", "error_title": "Sai", "error": "Chọn Có hoặc Không", "prompt_title": "", "prompt": ""},
            {"range": (0, 1, 9, 1), "type": "list", "source": "=$E$1:$E$4", "dropdown": True, "blank": True,
             "style": "warning", "error_title": "", "error": "", "prompt_title": "Thành phố", "prompt": "Chọn từ danh sách"},
            {"range": (0, 2, 4, 2), "type": "whole", "op": "between", "v1": "1", "v2": "=F1", "blank": False,
             "style": "info", "error_title": "", "error": "", "prompt_title": "", "prompt": ""},
            {"range": (0, 3, 4, 3), "type": "date", "op": ">=", "v1": "01/01/2026", "blank": True,
             "style": "stop", "error_title": "", "error": "", "prompt_title": "", "prompt": ""},
        ]
        self.sh.notes = {(0, 0): "Ghi chú đầu tiên\nhai dòng", (7, 7): "Ô trống có ghi chú"}
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "dv" + ext)
            fe.save_workbook(self.wb, path)
            wb = fe.load_workbook(path)
        sh = wb.sheets["Sheet1"]
        self.assertEqual(sh.notes, {(0, 0): "Ghi chú đầu tiên\nhai dòng", (7, 7): "Ô trống có ghi chú"})
        by_col = {rule["range"][1]: rule for rule in sh.validations}
        self.assertEqual(sorted(by_col), [0, 1, 2, 3])
        self.assertEqual((by_col[0]["source"], by_col[0]["error"], by_col[0]["style"]),
                         ("Có,Không", "Chọn Có hoặc Không", "stop"))
        self.assertEqual((by_col[1]["source"], by_col[1]["style"], by_col[1]["prompt"]),
                         ("=$E$1:$E$4", "warning", "Chọn từ danh sách"))
        self.assertEqual((by_col[2]["op"], by_col[2]["v1"], by_col[2]["v2"], by_col[2]["blank"]),
                         ("between", "1", "=F1", False))
        self.assertEqual((by_col[3]["op"], by_col[3]["v1"]), (">=", "01/01/2026"))
        self.assertTrue(by_col[0]["dropdown"])
        self.assertIsNotNone(wb.validate("Sheet1", 0, 2, "11"))
        self.assertIsNone(wb.validate("Sheet1", 0, 1, "Huế"))

    def test_json_roundtrip(self):
        self.roundtrip(".pyxl")

    @unittest.skipUnless(fe.HAS_XLSX, "openpyxl not installed")
    def test_xlsx_roundtrip(self):
        self.roundtrip(".xlsx")


class ChartTests(unittest.TestCase):
    def setUp(self):
        self.wb = make({"A1": "Tháng", "B1": "Doanh thu", "C1": "Chi phí",
                        "A2": "T1", "B2": "10", "C2": "7", "A3": "T2", "B3": "14", "C3": "x", "A4": "T3", "B4": "9"})
        self.chart = {"type": "column", "range": (0, 0, 3, 2), "title": "Kết quả", "anchor": (1, 4, 5, 3),
                      "size": (400, 250)}

    def test_series_in_columns_with_names_and_categories(self):
        labels, series = self.wb.chart_data("Sheet1", self.chart)
        self.assertEqual(labels, ["T1", "T2", "T3"])
        self.assertEqual(series, [("Doanh thu", [10, 14, 9]), ("Chi phí", [7, None, None])])

    def test_series_in_rows_and_plain_numbers(self):
        wb = make({"A1": "", "B1": "2024", "C1": "2025", "D1": "2026", "A2": "Bắc", "B2": "1", "C2": "2", "D2": "3"})
        labels, series = wb.chart_data("Sheet1", {"range": (0, 0, 1, 3)})
        self.assertEqual((labels, series), (["2024", "2025", "2026"], [("Bắc", [1, 2, 3])]))
        wb = make({"A1": "Tên", "B1": "Nhóm", "C1": "Điểm", "A2": "An", "B2": "x", "C2": "7",
                   "A3": "Bình", "B3": "y", "C3": "9"})
        self.assertEqual(wb.chart_data("Sheet1", {"range": (0, 0, 2, 2)}), (["An", "Bình"], [("Điểm", [7, 9])]))
        wb = make({"A1": "5", "A2": "6", "B1": "7", "B2": "8", "B3": "9"})
        labels, series = wb.chart_data("Sheet1", {"range": (0, 0, 2, 1)})
        self.assertEqual((labels, series), (["1", "2", "3"], [("Chuỗi 1", [5, 6, None]), ("Chuỗi 2", [7, 8, 9])]))

    def test_charts_follow_rows_and_columns(self):
        sh = self.wb.sheets["Sheet1"]
        sh.charts = [self.chart]
        self.wb.insert_rows("Sheet1", 0, 2)
        self.assertEqual((sh.charts[0]["range"], sh.charts[0]["anchor"]), ((2, 0, 5, 2), (3, 4, 5, 3)))
        self.wb.delete_cols("Sheet1", 3, 2)  # the anchor column goes: the chart moves to where it was
        self.assertEqual(sh.charts[0]["anchor"], (3, 3, 5, 3))
        self.wb.delete_rows("Sheet1", 2, 4)  # all its data goes: so does the chart
        self.assertEqual(sh.charts, [])

    def test_axis_ticks(self):
        import charts
        self.assertEqual(charts.nice_ticks(9, 14), [0, 5, 10, 15])
        self.assertEqual(charts.nice_ticks(-20, 210), [-50, 0, 50, 100, 150, 200, 250])
        self.assertEqual(charts.nice_ticks(0, 0), [0, 0.2, 0.4, 0.6, 0.8, 1])
        self.assertEqual(charts.tick_label(12500, 2500), "12,500")
        self.assertEqual(charts.tick_label(2.5, 2.5), "2.5")

    def roundtrip(self, ext):
        sh = self.wb.sheets["Sheet1"]
        sh.charts = [self.chart, {"type": "pie", "range": (0, 0, 3, 1), "title": "", "anchor": (20, 1, 0, 0),
                                  "size": (300, 200)}]
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "chart" + ext)
            fe.save_workbook(self.wb, path)
            return fe.load_workbook(path).sheets["Sheet1"].charts

    def test_json_roundtrip(self):
        self.assertEqual(self.roundtrip(".pyxl"), self.wb.sheets["Sheet1"].charts)

    @unittest.skipUnless(fe.HAS_XLSX, "openpyxl not installed")
    def test_xlsx_roundtrip(self):
        charts = self.roundtrip(".xlsx")
        self.assertEqual(charts, self.wb.sheets["Sheet1"].charts)


class RichClipboardTests(unittest.TestCase):
    TABLE = {"cols": [80, 120], "rows": [20, 30, 20],
             "cells": [[{"text": "Bảng giá", "align": "center",
                         "fmt": {"bold": True, "bg": "#ddebf7", "bb": ["thin", "#000000"]}},
                        {"text": "", "align": "left", "fmt": {}}],
                       [{"text": "Thép <D10>", "align": "left", "fmt": {"font": "Arial", "size": 10}},
                        {"text": "1,250", "align": "right", "fmt": {"fg": "#ff0000", "br": ["double", "#000000"]}}],
                       [{"text": "Dòng 1\nDòng 2", "align": "left", "fmt": {"wrap": True}},
                        {"text": "", "align": "left", "fmt": {}}]],
             "merges": [(0, 0, 0, 1)]}

    def test_html(self):
        import richclip
        out = richclip.html_table(self.TABLE)
        self.assertIn('<td colspan=2 style="', out)
        self.assertEqual(out.count("<td"), 5)  # the merged-away cell is left out
        self.assertIn("Thép &lt;D10&gt;", out)
        self.assertIn("Dòng 1<br>Dòng 2", out)
        self.assertIn("border-bottom:0.5pt solid #000000", out)
        self.assertIn("font-weight:700", out)
        data = richclip.html_clipboard(self.TABLE)
        header = data[:data.index(b"<html>")].decode("ascii")
        offsets = dict(line.split(":") for line in header.split())
        fragment = data[int(offsets["StartFragment"]):int(offsets["EndFragment"])].decode("utf-8")
        self.assertTrue(fragment.startswith("<table") and fragment.endswith("</table>"))
        self.assertEqual(int(offsets["EndHTML"]), len(data))

    def test_rtf(self):
        import richclip
        out = richclip.rtf_table(self.TABLE)
        out.encode("ascii")  # RTF is 7-bit; Vietnamese goes through \u escapes
        self.assertIn("\\clmgf", out)
        self.assertIn("\\clmrg", out)
        self.assertIn("\\cellx1200\\", out.replace("\n", "\\"))  # 80 px = 1200 twips
        self.assertIn("B\\u7843?ng gi\\u225?", out)  # "Bảng giá"
        self.assertIn("\\line ", out)
        self.assertEqual(out.count("\\row"), 3)
        self.assertIn("{\\f1\\fnil\\fcharset0 Arial;}", out)

    def test_windows_clipboard_calls(self):
        """Run the Win32 code path against a stand-in for user32/gdi32/kernel32 that records calls."""
        import ctypes
        import ctypes.wintypes
        import types
        from unittest import mock

        import richclip
        calls, handles = [], iter(range(1000, 100000))
        caps = {88: 96, 90: 96, 4: 344, 8: 1920, 6: 194, 10: 1080}

        def func(name):
            def call(*args):
                calls.append((name, args))
                if name == "GetDeviceCaps":
                    return caps[args[1]]
                if name == "RegisterClipboardFormatW":
                    return {"HTML Format": 49330, "Rich Text Format": 49331}[args[0]]
                return 1 if name in ("OpenClipboard", "EmptyClipboard", "CloseClipboard") else next(handles)
            return call

        class FakeDll:
            def __init__(self, *a, **k):
                self.funcs = {}

            def __getattr__(self, name):
                return self.funcs.setdefault(name, mock.Mock(side_effect=func(name)))

        fake = types.SimpleNamespace(WinDLL=FakeDll, wintypes=ctypes.wintypes, c_void_p=ctypes.c_void_p, c_int=ctypes.c_int,
                                     c_size_t=ctypes.c_size_t, POINTER=ctypes.POINTER, byref=lambda x: x,
                                     memmove=lambda *a: calls.append(("memmove", a)))
        with mock.patch.object(richclip.sys, "platform", "win32"), mock.patch.dict("sys.modules", {"ctypes": fake}):
            self.assertTrue(richclip.copy_to_clipboard(123, "a\tb\nc", self.TABLE))
        names = [n for n, _ in calls]
        formats = [a[0] for n, a in calls if n == "SetClipboardData"]
        self.assertEqual(formats, [13, 49330, 49331, 14, 2])
        self.assertLess(names.index("CloseEnhMetaFile"), names.index("OpenClipboard"))
        self.assertEqual(names[-1], "CloseClipboard")
        texts = [a[1] for n, a in calls if n == "DrawTextW"]
        self.assertEqual(texts, ["Bảng giá", "Thép <D10>", "1,250", "Dòng 1\r\nDòng 2"])
        payloads = [a[1] for n, a in calls if n == "memmove"]
        self.assertEqual(payloads[0], "a\tb\r\nc\0".encode("utf-16-le"))
        self.assertEqual(sum(1 for n in names if n == "CreatePen"), 2)
        self.assertFalse(richclip.copy_to_clipboard(123, "x", self.TABLE) if sys.platform != "win32" else False)


@unittest.skipUnless(importlib.util.find_spec("tkinter"), "tkinter not installed")
class MenuIconTests(unittest.TestCase):
    def test_every_icon_is_in_the_sprite_sheets(self):
        """Fails when an icon was added without rerunning tools/make_menu_icons.py."""
        import excel_app
        from ribbon import ICONS, MenuIcons
        folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), "icons")
        with open(os.path.join(folder, "menu_icons.txt"), encoding="utf-8") as f:
            names = [n for n in f.read().split("\n") if n]
        self.assertEqual(set(ICONS) - set(names), set())
        self.assertEqual(set(excel_app.MENU_ICONS.values()) - set(ICONS), set())
        for size in MenuIcons.SIZES:
            with open(os.path.join(folder, f"menu_{size}.png"), "rb") as f:
                header = f.read(24)
            width, height = int.from_bytes(header[16:20], "big"), int.from_bytes(header[20:24], "big")
            self.assertEqual((width, height), (MenuIcons.COLUMNS * size, -(-len(names) // MenuIcons.COLUMNS) * size))


class AutoFilterTests(unittest.TestCase):
    def setUp(self):
        import datetime as dt
        today = dt.date.today()
        self.wb = make({"A1": "Tên", "B1": "Số", "C1": "Ngày"})
        rows = [("Nguyễn An", "5", today - dt.timedelta(days=1)), ("Trần Bình", "15", today),
                ("Lê Chi", "25", today + dt.timedelta(days=40)), ("Nguyễn Dũng", "35", dt.date(today.year - 1, 6, 1)),
                ("Phạm Em", "x", today)]
        for i, (name, num, day) in enumerate(rows, 2):
            make({f"A{i}": name, f"B{i}": num, f"C{i}": day.strftime("%d/%m/%Y")}, wb=self.wb)
        self.sh = self.wb.sheets["Sheet1"]
        self.sh.filter = {"range": (0, 0, 5, 2), "criteria": {}, "hidden": set()}
        self.wb.set_format("Sheet1", 2, 0, {"bg": "#ffff00"})
        self.wb.set_format("Sheet1", 4, 0, {"bg": "#ffff00", "fg": "#ff0000"})

    def shown(self, col, crit):
        self.sh.filter["criteria"] = {col: crit}
        self.wb.apply_filter("Sheet1")
        return [r for r in range(1, 6) if r not in self.sh.filter["hidden"]]

    def test_custom_number_and_text(self):
        self.assertEqual(self.shown(1, {"type": "custom", "join": "and", "conds": [[">", "10"], ["<=", "25"]]}), [2, 3])
        self.assertEqual(self.shown(1, {"type": "custom", "join": "or", "conds": [["<", "10"], [">", "30"]]}), [1, 4])
        self.assertEqual(self.shown(0, {"type": "custom", "join": "and", "conds": [["begins", "nguyễn"]]}), [1, 4])
        self.assertEqual(self.shown(0, {"type": "custom", "join": "and", "conds": [["ncontains", "n"]]}), [3, 5])
        self.assertEqual(self.shown(0, {"type": "custom", "join": "and", "conds": [["=", "*ũng"]]}), [4])
        self.assertEqual(self.shown(0, {"type": "custom", "join": "and", "conds": [["ends", "chi"]]}), [3])
        self.assertEqual(self.shown(1, {"type": "custom", "join": "and", "conds": [["=", "15"]]}), [2])

    def test_top_average_and_dates(self):
        self.assertEqual(self.shown(1, {"type": "top", "n": 2}), [3, 4])
        self.assertEqual(self.shown(1, {"type": "top", "n": 1, "bottom": True}), [1])
        self.assertEqual(self.shown(1, {"type": "top", "n": 50, "percent": True}), [3, 4])
        self.assertEqual(self.shown(1, {"type": "average", "above": True}), [3, 4])  # average 20
        self.assertEqual(self.shown(1, {"type": "average", "above": False}), [1, 2])
        self.assertEqual(self.shown(2, {"type": "dynamic", "kind": "today"}), [2, 5])
        self.assertEqual(self.shown(2, {"type": "dynamic", "kind": "yesterday"}), [1])
        self.assertEqual(self.shown(2, {"type": "dynamic", "kind": "lastyear"}), [4])
        self.assertEqual(self.shown(2, {"type": "custom", "join": "and",
                                        "conds": [[">", fe.format_date(fe.parse_literal(self.wb.display_text("Sheet1", 2, 2))[0], "dd/mm/yyyy")]]}), [3])
        self.assertEqual(self.wb.filter_kind("Sheet1", 1), "number")
        self.assertEqual(self.wb.filter_kind("Sheet1", 2), "date")
        self.assertEqual(self.wb.filter_kind("Sheet1", 0), "text")

    def test_colors(self):
        self.assertEqual(self.wb.filter_colors("Sheet1", 0), [None, "#ffff00"])
        self.assertEqual(self.wb.filter_colors("Sheet1", 1), [])  # one colour: nothing to filter by
        self.assertEqual(self.shown(0, {"type": "color", "what": "bg", "color": "#ffff00"}), [2, 4])
        self.assertEqual(self.shown(0, {"type": "color", "what": "bg", "color": None}), [1, 3, 5])
        self.assertEqual(self.shown(0, {"type": "color", "what": "fg", "color": "#ff0000"}), [4])
        self.sh.cond_formats = [{"range": (1, 1, 5, 1), "type": "cell", "op": ">", "v1": "20", "style": {"bg": "#ffc7ce"}}]
        self.wb.invalidate()
        # like Excel, the "greater than 20" rule also colours the text "x" (text sorts after numbers)
        self.assertEqual(self.shown(1, {"type": "color", "what": "bg", "color": "#ffc7ce"}), [3, 4, 5])

    def test_dynamic_spans(self):
        import datetime as dt
        wed = dt.date(2026, 10, 21)
        self.assertEqual(fe.dynamic_date_span("thisweek", wed), (dt.date(2026, 10, 18), dt.date(2026, 10, 24)))
        self.assertEqual(fe.dynamic_date_span("lastmonth", dt.date(2026, 1, 15)), (dt.date(2025, 12, 1), dt.date(2025, 12, 31)))
        self.assertEqual(fe.dynamic_date_span("nextyear", wed), (dt.date(2027, 1, 1), dt.date(2027, 12, 31)))

    def roundtrip(self, ext):
        criteria = {0: {"type": "custom", "join": "or", "conds": [["begins", "Nguyễn"], ["contains", "Chi"]]},
                    1: {"type": "top", "n": 3, "bottom": True, "percent": False}}
        for crit2 in ({"type": "dynamic", "kind": "thismonth"}, {"type": "color", "what": "bg", "color": "#ffff00"},
                      {"type": "average", "above": False}, {"type": "color", "what": "fg", "color": "#ff0000"},
                      {"type": "custom", "join": "and", "conds": [[">=", "01/01/2026"], ["<", "31/12/2026"]]}):
            col = 2 if crit2["type"] in ("dynamic", "custom") else 0 if crit2["type"] == "color" else 1
            self.sh.filter["criteria"] = {**criteria, col: crit2}
            self.wb.apply_filter("Sheet1")
            with tempfile.TemporaryDirectory() as tmp:
                path = os.path.join(tmp, "f" + ext)
                fe.save_workbook(self.wb, path)
                loaded = fe.load_workbook(path).sheets["Sheet1"].filter
            got = loaded["criteria"][col]
            if crit2["type"] == "custom" and ext == ".xlsx":  # dates are stored as serial numbers
                self.assertEqual([op for op, _ in got["conds"]], [">=", "<"])
                self.assertEqual(fe.parse_literal(got["conds"][0][1])[0], fe.parse_literal("01/01/2026")[0])
            else:
                self.assertEqual(got, crit2)
            self.assertEqual(loaded["hidden"], self.sh.filter["hidden"])
            if col != 0:
                self.assertEqual(loaded["criteria"][0], criteria[0])
            if col != 1:
                self.assertEqual(loaded["criteria"][1], criteria[1])

    def test_json_roundtrip(self):
        self.roundtrip(".pyxl")

    @unittest.skipUnless(fe.HAS_XLSX, "openpyxl not installed")
    def test_xlsx_roundtrip(self):
        self.roundtrip(".xlsx")


class PivotTests(unittest.TestCase):
    def setUp(self):
        import pivot as pv
        self.pv = pv
        self.wb = Workbook(("Data", "Pivot"))
        rows = [("Khu vực", "Sản phẩm", "Ngày", "Số lượng", "Doanh thu"),
                ("Bắc", "Thép", "05/01/2026", "10", "1000"), ("Bắc", "Xi măng", "10/02/2026", "5", "250"),
                ("Nam", "Thép", "15/01/2026", "8", "800"), ("Nam", "Thép", "20/02/2026", "2", "200"),
                ("Bắc", "Thép", "25/02/2026", "4", "400"), ("Trung", "Xi măng", "01/03/2026", "6", "=50*6"),
                ("", "", "", "", "")]
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                self.wb.set_raw("Data", r, c, v)
        self.wb.set_format("Data", 1, 4, {"num": "#,##0"})
        self.p = pv.new_pivot("PivotTable1", "Data", (0, 0, 7, 4), (2, 0))
        self.wb.sheets["Pivot"].pivots.append(self.p)

    def grid(self, report):
        return [[report.cells.get((r, c)) for c in range(report.width)] for r in range(report.height)]

    def test_rows_with_subtotals_and_grand_total(self):
        self.p.update(rows=["Khu vực", "Sản phẩm"], values=[{"field": "Doanh thu", "func": "sum"}])
        rep = self.pv.compute(self.wb, self.p)
        self.assertEqual(self.grid(rep), [
            ["Nhãn Hàng", "Tổng của Doanh thu"], ["Bắc", 1650], ["Thép", 1400], ["Xi măng", 250], ["Nam", 1000],
            ["Thép", 1000], ["Trung", 300], ["Xi măng", 300], ["Tổng Cuối", 2950]])
        self.assertEqual(rep.formats[(2, 0)], {"indent": 1})
        self.assertTrue(rep.formats[(1, 1)]["bold"])           # a parent line
        self.assertEqual(rep.formats[(2, 1)]["num"], "#,##0")  # the source column's number format
        self.assertEqual(rep.sources[(2, 1)], (("Bắc", "Thép"), ()))

    def test_columns_grouped_dates_filters_and_two_values(self):
        self.p.update(rows=["Khu vực"], cols=["Ngày"], group={"Ngày": "month"}, filters=["Sản phẩm"],
                      items={"Sản phẩm": {"Thép"}},
                      values=[{"field": "Doanh thu", "func": "sum"}, {"field": "Số lượng", "func": "average"}])
        g = self.grid(self.pv.compute(self.wb, self.p))
        self.assertEqual(g[0][:2], ["Sản phẩm", "Thép"])
        self.assertEqual(g[2][1], "Nhãn Cột")
        self.assertEqual([g[3][c] for c in (1, 3, 5)], ["Tháng 01/2026", "Tháng 02/2026", "Tổng Cuối"])
        self.assertEqual(g[4][:3], ["Nhãn Hàng", "Tổng của Doanh thu", "Trung bình của Số lượng"])
        self.assertEqual(g[5], ["Bắc", 1000, 10, 400, 4, 1400, 7])
        self.assertEqual(g[7], ["Tổng Cuối", 1800, 9, 600, 3, 2400, 6])

    def test_tabular_percentages_counts_and_sorting(self):
        self.p.update(rows=["Khu vực", "Sản phẩm"], layout="tabular", sort={"Khu vực": "desc"},
                      values=[{"field": "Doanh thu", "func": "sum", "show": "pct_total"},
                              {"field": "Khu vực", "func": "count"}])
        rep = self.pv.compute(self.wb, self.p)
        g = self.grid(rep)
        self.assertEqual(g[0], ["Khu vực", "Sản phẩm", "Tổng của Doanh thu", "Đếm của Khu vực"])
        self.assertEqual([row[0] for row in g[1:]], ["Trung", "Trung Tổng", "Nam", "Nam Tổng", "Bắc", None,
                                                    "Bắc Tổng", "Tổng Cuối"])
        self.assertAlmostEqual(g[-1][2], 1.0)
        self.assertAlmostEqual(g[5][2], 1400 / 2950)
        self.assertEqual(g[5][3], 2)
        self.assertEqual(rep.formats[(1, 2)]["num"], "0.00%")

    def test_column_subtotals_and_no_rows(self):
        self.p.update(cols=["Khu vực", "Sản phẩm"], values=[{"field": "Số lượng", "func": "sum"}])
        g = self.grid(self.pv.compute(self.wb, self.p))
        self.assertEqual(g[1][1:], ["Bắc", None, "Bắc Tổng", "Nam", "Nam Tổng", "Trung", "Trung Tổng", "Tổng Cuối"])
        self.assertEqual(g[2][1:3], ["Thép", "Xi măng"])
        self.assertEqual(g[3][1:], [14, 5, 19, 10, 10, 6, 6, 35])

    def test_place_refresh_drill_and_placeholder(self):
        pv = self.pv
        report = pv.refresh(self.wb, "Pivot", self.p)  # nothing chosen yet: the placeholder box
        self.assertEqual(self.wb.get_raw("Pivot", 3, 0), "PivotTable1")
        self.p.update(rows=["Sản phẩm"], values=[{"field": "Doanh thu", "func": "sum"}])
        report = pv.refresh(self.wb, "Pivot", self.p)
        self.assertEqual(self.p["rect"], (2, 0, 5, 1))
        self.assertEqual(self.wb.get_raw("Pivot", 3, 0), "Thép")  # the placeholder is gone
        self.assertEqual(self.wb.get_value("Pivot", 5, 1), 2950)
        self.assertIs(pv.pivot_at(self.wb, "Pivot", 4, 1), self.p)
        self.assertIsNone(pv.pivot_at(self.wb, "Pivot", 7, 0))
        self.wb.set_raw("Data", 1, 4, "2000")  # the report is static until refreshed, like Excel's
        self.assertEqual(self.wb.get_value("Pivot", 5, 1), 2950)
        pv.refresh(self.wb, "Pivot", self.p)
        self.assertEqual(self.wb.get_value("Pivot", 5, 1), 3950)
        names, records = pv.drill_down(self.wb, self.p, *report.sources[(1, 1)])  # the "Thép" cell
        self.assertEqual(names[0], "Khu vực")
        self.assertEqual([rec[0][1] for rec in records], ["Bắc", "Nam", "Nam", "Bắc"])
        self.wb.set_raw("Pivot", 20, 0, "x")
        self.p["anchor"] = (19, 0)
        self.assertTrue(pv.overlaps(self.wb, "Pivot", self.p, report))

    def test_source_follows_rows_renames_and_files(self):
        self.p.update(rows=["Sản phẩm"], values=[{"field": "Doanh thu", "func": "sum"}], items={"Sản phẩm": {"Thép"}})
        self.pv.refresh(self.wb, "Pivot", self.p)
        self.wb.insert_rows("Data", 0, 2)
        self.assertEqual(self.p["source"]["range"], (2, 0, 9, 4))
        self.wb.insert_rows("Pivot", 0)
        self.assertEqual((self.p["anchor"], self.p["rect"]), ((3, 0), (3, 0, 5, 1)))
        self.wb.rename_sheet("Data", "Dữ liệu")
        self.assertEqual(self.p["source"]["sheet"], "Dữ liệu")
        self.assertEqual(self.pv.compute(self.wb, self.p).cells[(1, 1)], 2400)
        for ext in (".pyxl",) + ((".xlsx",) if fe.HAS_XLSX else ()):
            with tempfile.TemporaryDirectory() as tmp:
                path = os.path.join(tmp, "pivot" + ext)
                fe.save_workbook(self.wb, path)
                wb = fe.load_workbook(path)
            self.assertEqual(list(wb.sheets), ["Dữ liệu", "Pivot"])  # the store sheet stays hidden
            loaded = wb.sheets["Pivot"].pivots[0]
            self.assertEqual(loaded, self.p)
            self.assertEqual(wb.get_value("Pivot", 4, 1), 2400)


class NaturalOrderTests(unittest.TestCase):
    def test_numbers_and_codes_sort_by_value(self):
        order = lambda items: sorted(items, key=fe.natural_key)  # noqa: E731
        self.assertEqual(order(["Mục 10", "Mục 2", "Mục 1"]), ["Mục 1", "Mục 2", "Mục 10"])
        self.assertEqual(order(["10", "9", "1", "100"]), ["1", "9", "10", "100"])  # numbers typed as text
        self.assertEqual(order([10.0, "9", 1.0, "2"]), [1.0, "2", "9", 10.0])     # mixed with real numbers
        self.assertEqual(order(["BCR-1000x50", "BCR-300x30", "BCR-400x40"]), ["BCR-300x30", "BCR-400x40", "BCR-1000x50"])
        self.assertEqual(order(["SP10", "sp2", "SP1"]), ["SP1", "sp2", "SP10"])
        self.assertEqual(order(["15/09/2026", "02/10/2026", "01/01/2025"]), ["01/01/2025", "15/09/2026", "02/10/2026"])

    def test_vietnamese_alphabet(self):
        names = ["Đức", "Dũng", "Ánh", "Anh", "Bình", "Zung", "Ân", "Ăn", "Ệ", "Ê", "Em"]
        self.assertEqual(sorted(names, key=fe.natural_key),
                         ["Anh", "Ánh", "Ăn", "Ân", "Bình", "Dũng", "Đức", "Em", "Ê", "Ệ", "Zung"])

    def test_filter_lists_and_pivot_items_use_it(self):
        import pivot as pv
        wb = make({"A1": "Mã", "B1": "Số", "A2": "SP10", "B2": "1", "A3": "SP2", "B3": "2", "A4": "SP1", "B4": "3",
                   "A5": "'10", "B5": "4", "A6": "'9", "B6": "5"})
        wb.sheets["Sheet1"].filter = {"range": (0, 0, 5, 1), "criteria": {}, "hidden": set()}
        self.assertEqual(wb.filter_values("Sheet1", 0), ["9", "10", "SP1", "SP2", "SP10"])
        p = pv.new_pivot("P", "Sheet1", (0, 0, 5, 1), (0, 3))
        p.update(rows=["Mã"], values=[{"field": "Số", "func": "sum"}])
        report = pv.compute(wb, p)
        self.assertEqual([report.cells.get((r, 0)) for r in range(1, 6)], ["9", "10", "SP1", "SP2", "SP10"])
        self.assertIs(fe.compare("<", "10", "9"), True)  # formulas still compare text the Excel way


class LayoutFileTests(unittest.TestCase):
    def build(self):
        wb = make({"A1": "Báo cáo", "A2": "10", "A3": "20", "B2": "=A2*2", "A4": "x", "A5": "y"})
        sh = wb.sheets["Sheet1"]
        sh.set_merges([(0, 0, 0, 2)])
        wb.set_format("Sheet1", 0, 0, {"font": "Times New Roman", "size": 16, "bold": True, "align": "center",
                                       "valign": "center", "underline": "double"})
        wb.set_format("Sheet1", 1, 0, {"wrap": True, "rotate": -45, "indent": 2, "strike": True,
                                       "bt": ["thin", "#000000"], "bb": ["double", "#ff0000"],
                                       "num": '#,##0 ₫;-#,##0 ₫;"-" ₫'})
        sh.row_heights, sh.custom_rows = {0: 40}, {0}
        sh.hidden_cols = {5}
        sh.tab_color = "#ff0000"
        sh.cond_formats = [{"range": (1, 0, 2, 0), "type": "cell", "op": ">", "v1": "15",
                            "style": {"bg": "#ffc7ce", "fg": "#9c0006"}},
                           {"range": (1, 0, 2, 0), "type": "scale", "colors": ["#ffffff", "#63be7b"]}]
        sh.filter = {"range": (2, 0, 4, 0), "criteria": {0: {"x"}}, "hidden": set()}
        wb.apply_filter("Sheet1")
        sh.freeze = (0, 1, 0, 2)
        name = wb.add_sheet("Khác")  # add_sheet replaces wb.sheets, so look up afterwards
        wb.sheets[name].freeze = (10, 12, 0, 0)  # frozen while scrolled: rows 11-12 stay on screen
        return wb

    def check(self, wb):
        sh = wb.sheets["Sheet1"]
        self.assertEqual(sh.merges, [(0, 0, 0, 2)])
        f0, f1 = wb.get_format("Sheet1", 0, 0), wb.get_format("Sheet1", 1, 0)
        self.assertEqual((f0["font"], f0["size"], f0["underline"], f0["valign"]), ("Times New Roman", 16, "double", "center"))
        self.assertEqual((f1["wrap"], f1["rotate"], f1["indent"], f1["strike"]), (True, -45, 2, True))
        self.assertEqual((f1["bt"], f1["bb"]), (["thin", "#000000"], ["double", "#ff0000"]))
        self.assertEqual(fe.format_value(*wb.cell_display("Sheet1", 1, 0)), "10 ₫")
        self.assertEqual(sh.row_heights.get(0), 40)
        self.assertEqual((sh.hidden_cols, sh.tab_color), ({5}, "#ff0000"))
        self.assertEqual([r["type"] for r in sh.cond_formats], ["cell", "scale"])
        self.assertTrue(wb.cond_format("Sheet1", 2, 0)["fg"] == "#9c0006")
        self.assertEqual((sh.filter["range"], sh.filter["hidden"]), ((2, 0, 4, 0), {4}))
        self.assertNotIn(4, sh.hidden_rows)
        self.assertEqual(sh.freeze, (0, 1, 0, 2))
        self.assertEqual(wb.sheets["Khác"].freeze, (10, 12, 0, 0))

    def roundtrip(self, ext):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "layout" + ext)
            fe.save_workbook(self.build(), path)
            return fe.load_workbook(path)

    def test_json_layout(self):
        self.check(self.roundtrip(".pyxl"))

    @unittest.skipUnless(fe.HAS_XLSX, "openpyxl not installed")
    def test_xlsx_layout(self):
        self.check(self.roundtrip(".xlsx"))


class PrintingTests(unittest.TestCase):
    def build(self, rows=60):
        wb = make({"A1": "Bảng khối lượng", "A2": "STT", "B2": "Tên"})
        for i in range(rows):
            wb.set_raw("Sheet1", i + 2, 0, str(i + 1))
            wb.set_raw("Sheet1", i + 2, 1, f"Hạng mục {i + 1}")
        return wb

    def layout(self, wb, **setup):
        import printing
        sh = wb.sheets["Sheet1"]
        sh.page_setup = setup
        return printing.Layout(wb, "Sheet1", printing.setup_of(sh))

    def test_pages_repeat_title_rows_below_the_top(self):
        lay = self.layout(self.build(), titles=(1, 1))
        self.assertEqual(len(lay.pages), 2)
        (first, _), (second, _) = lay.pages
        self.assertEqual(first[:3], [0, 1, 2])         # the title row prints once, in place, on page 1
        self.assertEqual(second[:2], [1, first[-1] + 1])  # and is repeated above the rows of page 2
        self.assertEqual(lay.leads, [0, 1])
        self.assertEqual(lay.breaks(), ({first[-1] + 1}, set()))

    def test_manual_breaks_scaling_and_fit(self):
        wb = self.build()
        lay = self.layout(wb, breaks=[10], col_breaks=[1])  # 48 rows fit an A4 page; pages go down, then over
        self.assertEqual([(rows[0], rows[-1], cols) for rows, cols in lay.pages],
                         [(0, 9, [0]), (10, 57, [0]), (58, 61, [0]), (0, 9, [1]), (10, 57, [1]), (58, 61, [1])])
        self.assertEqual(len(self.layout(wb, scale=50).pages), 1)
        lay = self.layout(wb, fit="page")
        self.assertEqual(len(lay.pages), 1)
        self.assertLess(lay.scale, 1)
        self.assertEqual(len(self.layout(wb, area=(0, 0, 4, 1)).pages), 1)
        wb.sheets["Sheet1"].col_widths[1] = 2000
        lay = self.layout(wb, fit="width")
        self.assertEqual(len({tuple(cols) for _, cols in lay.pages}), 1)

    def test_setup_follows_inserted_rows_and_files(self):
        import printing
        wb = self.build(5)
        sh = wb.sheets["Sheet1"]
        sh.page_setup = {"area": (0, 0, 6, 1), "titles": (1, 1), "breaks": [4], "col_breaks": [1],
                         "orientation": "landscape", "paper": "A3", "fit": "width", "gridlines": True,
                         "center_h": True, "header": "&A", "footer": "Trang &P / &N",
                         "margins": printing.MARGINS["narrow"]}
        wb.insert_rows("Sheet1", 0)
        expected = dict(sh.page_setup, area=(1, 0, 7, 1), titles=(2, 2), breaks=[5])
        self.assertEqual(sh.page_setup, expected)
        exts = [".pyxl"] + ([".xlsx"] if fe.HAS_XLSX else [])
        for ext in exts:
            with tempfile.TemporaryDirectory() as tmp:
                path = os.path.join(tmp, "in" + ext)
                fe.save_workbook(wb, path)
                got = fe.load_workbook(path).sheets["Sheet1"].page_setup
            self.assertEqual(printing.setup_of(type("S", (), {"page_setup": got})), printing.setup_of(sh), ext)

    def test_title_rows_text_and_header_codes(self):
        import printing
        self.assertEqual(printing.parse_title_rows("$2:$3"), (1, 2))
        self.assertEqual(printing.parse_title_rows("4"), (3, 3))
        self.assertIsNone(printing.parse_title_rows(" "))
        self.assertIs(printing.parse_title_rows("A1"), False)
        self.assertEqual(printing.title_rows_text((1, 2)), "$2:$3")
        self.assertEqual(printing.header_text("&A - Trang &P / &N", 2, 5, "Thép", "a.xlsx"), "Thép - Trang 2 / 5")

    def test_pdf_export(self):
        import printing
        import ttf
        library = ttf.FontLibrary()
        if library.find("Arial") is None:
            self.skipTest("no TrueType fonts on this machine")
        wb = self.build()
        wb.sheets["Sheet1"].page_setup = {"titles": (1, 1), "footer": "Trang &P / &N"}
        wb.sheets["Sheet1"].charts = [{"type": "column", "range": (1, 0, 6, 0), "title": "Biểu đồ",
                                       "anchor": (2, 3, 0, 0), "size": (300, 200)}]
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "out.pdf")
            pages = printing.export_pdf(wb, ["Sheet1"], path, library, "out.xlsx")
            with open(path, "rb") as f:
                data = f.read()
        self.assertEqual(pages, 2)
        self.assertTrue(data.startswith(b"%PDF-") and data.rstrip().endswith(b"%%EOF"))
        self.assertEqual(data.count(b"/Type /Page\n") + data.count(b"/Type /Page "), 2)
        self.assertIn(b"/ToUnicode", data)
        self.assertIn(b"/FontFile2", data)


class FormulaAssistTests(unittest.TestCase):
    def test_completion_and_call_context(self):
        import formula_assist as fa
        names = fe.function_names()
        self.assertEqual(fa.completion_prefix("=A1+vl", 6), (4, "vl"))
        self.assertIn("VLOOKUP", fa.matching_functions("vl", names))
        for text in ('="SU', "=Sheet2!A", "=SUM(", "abc"):
            self.assertIsNone(fa.completion_prefix(text, len(text)), text)
        self.assertEqual(fa.call_context("=IF(A1>2,SUM(B1,", 17), ("SUM", 1))
        self.assertEqual(fa.call_context('=IF(A1,"a,b",', 14), ("IF", 2))
        self.assertEqual(fa.call_context("=SUM((A1+B1),", 13), ("SUM", 1))
        self.assertIsNone(fa.call_context("=SUM(A1)", 8))
        args = fa.signature_args("SUM(số1, [số2], ...)")[1]
        self.assertEqual((fa.current_arg(args, 1), fa.current_arg(args, 6)), (1, 2))
        self.assertIsNone(fa.current_arg(fa.signature_args("ABS(số)")[1], 1))

    def test_f4_cycles_references(self):
        import formula_assist as fa
        text, steps = "=A1", []
        for _ in range(4):
            text, pos = fa.toggle_reference(text, len(text))
            steps.append(text)
        self.assertEqual(steps, ["=$A$1", "=A$1", "=$A1", "=A1"])
        self.assertEqual(fa.toggle_reference("=SUM(A1:B2)", 7), ("=SUM($A$1:$B$2)", 14))
        self.assertEqual(fa.toggle_reference("=LOG10(B3)+C4", 9)[0], "=LOG10($B$3)+C4")
        self.assertEqual(fa.toggle_reference('="A1"&B2', 8)[0], '="A1"&$B$2')
        self.assertEqual(fa.toggle_reference("=Sheet2!C3*2", 10)[0], "=Sheet2!$C$3*2")
        self.assertIsNone(fa.toggle_reference("=1+2", 4))


class OutlineTests(unittest.TestCase):
    def table(self):
        wb = make({"A1": "Khu vực", "B1": "Tiền", "A2": "Bắc", "B2": "10", "A3": "Bắc", "B3": "20",
                   "A4": "Nam", "B4": "5", "A5": "Nam", "B5": "7", "A6": "Trung", "B6": "1", "A8": "dưới"})
        return wb

    def test_groups_collapse_expand_and_levels(self):
        wb = Workbook()
        wb.group("Sheet1", 0, 1, 8)
        wb.group("Sheet1", 0, 2, 4)
        sh = wb.sheets["Sheet1"]
        self.assertEqual(fe.outline_groups(sh.row_levels), [(1, 1, 8), (2, 2, 4)])
        wb.set_collapsed("Sheet1", 0, 2, 2, 4, True)
        wb.set_collapsed("Sheet1", 0, 1, 1, 8, True)
        self.assertEqual(sh.hidden_rows, set(range(1, 9)))
        wb.set_collapsed("Sheet1", 0, 1, 1, 8, False)  # the inner group stays collapsed
        self.assertEqual(sh.hidden_rows, {2, 3, 4})
        wb.show_level("Sheet1", 0, 3)
        self.assertEqual((sh.hidden_rows, sh.collapsed_rows), (set(), set()))
        wb.insert_rows("Sheet1", 3)  # inside the inner group: joins it
        self.assertEqual(sh.row_levels[3], 2)
        self.assertTrue(wb.group("Sheet1", 0, 1, 9, -1))
        self.assertEqual(fe.outline_groups(sh.row_levels), [(1, 2, 5)])

    def test_subtotals(self):
        wb = self.table()
        rect = wb.add_subtotals("Sheet1", (0, 0, 5, 1), 0, 9, [1])
        self.assertEqual(rect, (0, 0, 9, 1))
        col = [(wb.display_text("Sheet1", r, 0), wb.display_text("Sheet1", r, 1)) for r in range(10)]
        self.assertEqual(col[3], ("Bắc Tổng", "30"))
        self.assertEqual(col[6], ("Nam Tổng", "12"))
        self.assertEqual(col[9], ("Tổng Cuối", "43"))  # SUBTOTAL skips the subtotals above it
        self.assertEqual(wb.get_raw("Sheet1", 11, 0), "dưới")  # rows were inserted, not overwritten
        sh = wb.sheets["Sheet1"]
        self.assertEqual([sh.row_levels.get(r, 0) for r in range(10)], [0, 2, 2, 1, 2, 2, 1, 2, 1, 0])
        rect = wb.add_subtotals("Sheet1", rect, 0, 1, [1])  # replaces the sums with averages
        self.assertEqual(wb.display_text("Sheet1", 9, 1), "8.6")
        self.assertEqual(wb.remove_subtotals("Sheet1", rect), (0, 0, 5, 1))
        self.assertEqual((sh.row_levels, wb.get_raw("Sheet1", 7, 0)), ({}, "dưới"))

    def test_outline_files(self):
        wb = self.table()
        wb.add_subtotals("Sheet1", (0, 0, 5, 1), 0, 9, [1])
        wb.group("Sheet1", 1, 1, 1)
        wb.show_level("Sheet1", 0, 2)
        sh = wb.sheets["Sheet1"]
        for ext in [".pyxl"] + ([".xlsx"] if fe.HAS_XLSX else []):
            with tempfile.TemporaryDirectory() as tmp:
                path = os.path.join(tmp, "o" + ext)
                fe.save_workbook(wb, path)
                got = fe.load_workbook(path).sheets["Sheet1"]
            self.assertEqual((got.row_levels, got.col_levels, got.collapsed_rows, got.hidden_rows),
                             (sh.row_levels, sh.col_levels, sh.collapsed_rows, sh.hidden_rows), ext)


class DefinedNameTests(unittest.TestCase):
    def build(self):
        wb = make({"B1": "10", "B2": "20", "B3": "30", "B4": "40", "B5": "50"})
        wb.define_name("ĐơnGiá", "=Sheet1!$B$1:$B$5", "giá")
        wb.define_name("VAT", "=0.1")
        wb.set_raw("Sheet1", 0, 3, fe.normalize_formula("=sum(đơngiá)*(1+vat)", wb.names))
        return wb

    def test_names_in_formulas(self):
        wb = self.build()
        self.assertEqual(wb.get_raw("Sheet1", 0, 3), "=SUM(ĐơnGiá)*(1+VAT)")  # written as defined
        self.assertAlmostEqual(val(wb, "D1"), 165)
        wb.set_raw("Sheet1", 1, 3, "=KhongCo+1")
        self.assertEqual(val(wb, "D2"), CellError("#NAME?"))
        self.assertEqual(wb.name_target("đơngiá"), ("Sheet1", (0, 1, 4, 1)))
        self.assertEqual(wb.name_value("VAT"), "0.1")
        self.assertEqual(wb.name_value("ĐơnGiá"), "{10; 20; 30; 40; 50}")
        wb.delete_name("VAT")
        self.assertEqual(val(wb, "D1"), CellError("#NAME?"))

    def test_checks_and_labels(self):
        wb = self.build()
        for bad in ("", "A1", "kl12", "R1C1", "ab c", "1abc", "vat", "TRUE"):
            self.assertIsNotNone(wb.check_name(bad), bad)
        self.assertIsNone(wb.check_name("Tổng_KL"))
        self.assertEqual(fe.name_from_label("Đơn giá (đ)"), "Đơn_giá_đ")
        self.assertEqual(fe.name_from_label("2024"), "_2024")
        self.assertEqual(fe.absolute_ref("Bảng 1", (0, 1, 4, 2)), "'Bảng 1'!$B$1:$C$5")

    def test_names_follow_rows_sheets_undo_and_files(self):
        wb = self.build()
        wb.insert_rows("Sheet1", 0, 2)
        self.assertEqual(wb.names["ĐƠNGIÁ"]["ref"], "=Sheet1!$B$3:$B$7")
        wb.rename_sheet("Sheet1", "Dữ liệu")
        self.assertEqual(wb.names["ĐƠNGIÁ"]["ref"], "='Dữ liệu'!$B$3:$B$7")
        snap = wb.snapshot()
        wb.delete_name("ĐơnGiá")
        wb.restore(snap)
        self.assertAlmostEqual(wb.get_value("Dữ liệu", 2, 3), 165)
        for ext in [".pyxl"] + ([".xlsx"] if fe.HAS_XLSX else []):
            with tempfile.TemporaryDirectory() as tmp:
                path = os.path.join(tmp, "n" + ext)
                fe.save_workbook(wb, path)
                got = fe.load_workbook(path)
            self.assertEqual(got.names, wb.names, ext)
            self.assertAlmostEqual(got.get_value("Dữ liệu", 2, 3), 165)


def _png(width, height, rows):
    """A tiny RGBA PNG; rows are (filter type, filtered bytes)."""
    import struct
    import zlib

    def chunk(kind, body):
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body))
    raw = b"".join(bytes([f]) + bytes(data) for f, data in rows)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 6, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


class LinkAndPictureTests(unittest.TestCase):
    PIXELS = [[10, 20, 30, 255, 15, 25, 35, 128], [12, 22, 32, 255, 20, 30, 40, 0]]

    def picture_png(self):
        a, b = self.PIXELS
        sub = a[:4] + [(a[i] - a[i - 4]) % 256 for i in range(4, 8)]  # filter 1 (Sub)
        up = [(b[i] - a[i]) % 256 for i in range(8)]                  # filter 2 (Up)
        return _png(2, 2, [(1, sub), (2, up)])

    def test_png_decoding_and_pdf_data(self):
        import images
        data = self.picture_png()
        self.assertEqual(images.image_info(data), ("png", 2, 2))
        w, h, rgb, alpha = images.decode_png(data)
        a, b = self.PIXELS
        self.assertEqual((w, h), (2, 2))
        self.assertEqual(list(rgb), a[0:3] + a[4:7] + b[0:3] + b[4:7])
        self.assertEqual(list(alpha), [255, 128, 255, 0])
        info = images.pdf_image(data, "png")
        self.assertEqual((info["w"], info["h"], info["space"]), (2, 2, "DeviceRGB"))
        self.assertIsNotNone(info["alpha"])

    def test_links_pictures_follow_rows_and_files(self):
        import base64
        wb = make({"A1": "Google", "A2": '=HYPERLINK("https://x.vn","Trang X")'})
        sh = wb.sheets["Sheet1"]
        sh.links[(0, 0)] = {"target": "https://www.google.com", "tip": "Tìm"}
        self.assertEqual(wb.display_text("Sheet1", 1, 0), "Trang X")
        self.assertEqual(wb.hyperlink_at("Sheet1", 1, 0)["target"], "https://x.vn")
        picture = {"type": fe.PICTURE, "image": base64.b64encode(self.picture_png()).decode(), "format": "png",
                   "id": "p1", "name": "ảnh", "anchor": (3, 1, 2, 2), "size": (40, 40)}
        sh.charts.append(picture)
        wb.insert_rows("Sheet1", 0)
        self.assertEqual((list(sh.links), sh.charts[0]["anchor"]), ([(1, 0)], (4, 1, 2, 2)))
        for ext in [".pyxl"] + ([".xlsx"] if fe.HAS_XLSX else []):
            with tempfile.TemporaryDirectory() as tmp:
                path = os.path.join(tmp, "l" + ext)
                fe.save_workbook(wb, path)
                got = fe.load_workbook(path).sheets["Sheet1"]
            self.assertEqual(got.links, sh.links, ext)
            if ext == ".xlsx" and importlib.util.find_spec("PIL") is None:
                continue  # openpyxl reads pictures back only with Pillow
            self.assertEqual([(c["type"], c["anchor"], c["size"], c["image"]) for c in got.charts],
                             [(fe.PICTURE, (4, 1, 2, 2), (40, 40), picture["image"])], ext)

    def test_pictures_in_pdf(self):
        import base64
        import printing
        import ttf
        library = ttf.FontLibrary()
        if library.find("Arial") is None:
            self.skipTest("no TrueType fonts on this machine")
        wb = make({"A1": "x"})
        wb.sheets["Sheet1"].charts.append({"type": fe.PICTURE, "image": base64.b64encode(self.picture_png()).decode(),
                                           "format": "png", "id": "p1", "anchor": (1, 1, 0, 0), "size": (60, 60)})
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "p.pdf")
            printing.export_pdf(wb, ["Sheet1"], path, library)
            with open(path, "rb") as f:
                data = f.read()
        self.assertIn(b"/Subtype /Image", data)
        self.assertIn(b"/SMask", data)


class SheetExtrasTests(unittest.TestCase):
    def test_hidden_sheets(self):
        wb = Workbook(("A", "B"))
        wb.hide_sheet("B")
        self.assertEqual(wb.visible_sheets(), ["A"])
        with self.assertRaises(ValueError):
            wb.hide_sheet("A")  # the last visible sheet stays
        with self.assertRaises(ValueError):
            wb.delete_sheet("A")
        for ext in [".pyxl"] + ([".xlsx"] if fe.HAS_XLSX else []):
            with tempfile.TemporaryDirectory() as tmp:
                path = os.path.join(tmp, "h" + ext)
                fe.save_workbook(wb, path, active="B")  # a hidden sheet is never the active one
                got = fe.load_workbook(path)
            self.assertEqual((got.visible_sheets(), got.sheets["B"].hidden), (["A"], True), ext)
            if ext == ".xlsx":
                self.assertEqual(got.active, "A")

    def test_special_cells(self):
        wb = make({"A1": "x", "B1": "5", "A2": "=B1*2", "C3": "TRUE", "B3": "=1/0"})
        everything = (0, 0, fe.MAX_ROWS - 1, fe.MAX_COLS - 1)
        self.assertEqual(wb.special_cells("Sheet1", everything, "constants"), [(0, 0), (0, 1), (2, 2)])
        self.assertEqual(wb.special_cells("Sheet1", everything, "constants", ("number",)), [(0, 1)])
        self.assertEqual(wb.special_cells("Sheet1", everything, "formulas", ("error",)), [(2, 1)])
        self.assertEqual(wb.special_cells("Sheet1", everything, "blanks"), [(0, 2), (1, 1), (1, 2), (2, 0)])
        self.assertEqual(wb.special_cells("Sheet1", everything, "last"), [(2, 2)])
        self.assertEqual(fe.cells_to_rects([(0, 0), (0, 1), (1, 0), (1, 1), (3, 0), (3, 2), (4, 2)]),
                         [(0, 0, 1, 1), (3, 0, 3, 0), (3, 2, 4, 2)])


class ProtectionAndSplitTests(unittest.TestCase):
    def test_split_text(self):
        self.assertEqual(fe.split_text('An, 0912, "Hà Nội, VN"', ","), ["An", " 0912", " Hà Nội, VN"])
        self.assertEqual(fe.split_text("a  b   c", " ", merge=True), ["a", "b", "c"])
        self.assertEqual(fe.split_text("a;;b", ";"), ["a", "", "b"])
        self.assertEqual(fe.split_text("THEP001  Cột 12", breaks=[7, 9]), ["THEP001", "", "Cột 12"])

    def test_password_hash_matches_excel(self):
        self.assertEqual(fe.password_hash("abc"), "CC1A")
        self.assertEqual(fe.password_hash(""), "")

    def test_protection_rules(self):
        wb = make({"A1": "5", "B1": "x"})
        wb.set_format("Sheet1", 0, 1, {"unlocked": True})
        sh = wb.sheets["Sheet1"]
        before = sh.copy_state()
        sh.cells[(0, 1)] = "y"
        self.assertIsNone(fe.protection_problem(before, sh.copy_state()))  # unlocked cell
        sh.cells[(0, 0)] = "6"
        self.assertEqual(fe.protection_problem(before, sh.copy_state()), "cells")
        sh.cells[(0, 0)] = "5"
        sh.formats[(0, 0)] = {"bold": True}
        self.assertEqual(fe.protection_problem(before, sh.copy_state()), "format")
        self.assertIsNone(fe.protection_problem(before, sh.copy_state(), ["format_cells"]))
        del sh.formats[(0, 0)]
        sh.row_levels = {3: 1}
        before = sh.copy_state()
        sh.hidden_rows.add(3)
        self.assertEqual(fe.protection_problem(before, sh.copy_state()), "rows")
        self.assertIsNone(fe.protection_problem(before, sh.copy_state(), ["outline"]))

    def test_protection_files(self):
        wb = make({"A1": "5"})
        wb.set_format("Sheet1", 1, 0, {"unlocked": True, "hide_formula": True})
        wb.sheets["Sheet1"].protection = {"hash": fe.password_hash("abc"), "allow": ["format_cells"]}
        for ext in [".pyxl"] + ([".xlsx"] if fe.HAS_XLSX else []):
            with tempfile.TemporaryDirectory() as tmp:
                path = os.path.join(tmp, "p" + ext)
                fe.save_workbook(wb, path)
                got = fe.load_workbook(path).sheets["Sheet1"]
            self.assertEqual(got.protection, {"hash": "CC1A", "allow": ["format_cells"]}, ext)
            self.assertEqual(got.formats[(1, 0)], {"unlocked": True, "hide_formula": True}, ext)


class SessionTests(unittest.TestCase):
    def test_recent_files_and_settings(self):
        import session
        with tempfile.TemporaryDirectory() as tmp:
            settings = session.Settings(tmp)
            a = os.path.join(tmp, "a.xlsx")
            open(a, "w").close()
            settings.add_recent(a)
            settings.add_recent(os.path.join(tmp, "gone.xlsx"))
            settings.add_recent(a)  # moves to the top, no duplicate
            settings.autosave_minutes = 3
            again = session.Settings(tmp)
            self.assertEqual(again.recent(), [a])
            self.assertEqual(again.autosave_minutes, 3)

    def test_toolbar_and_ribbon_settings(self):
        import session
        with tempfile.TemporaryDirectory() as tmp:
            settings = session.Settings(tmp)
            self.assertEqual((settings.qat, settings.ribbon), (["save", "undo", "redo"], "full"))
            self.assertEqual(settings.custom_title_bar, sys.platform == "win32")  # Windows: drawn like Excel
            settings.qat, settings.ribbon, settings.custom_title_bar = ["new", "save"], "auto", False
            again = session.Settings(tmp)
            self.assertEqual((again.qat, again.ribbon, again.custom_title_bar), (["new", "save"], "auto", False))

    def test_autorecover(self):
        import time
        import session
        with tempfile.TemporaryDirectory() as tmp:
            crashed = session.Recovery(tmp)
            wb = make({"A1": "chưa lưu"})
            crashed.beat(path="D:/bang.xlsx")
            self.assertTrue(crashed.save(wb))
            now = session.Recovery(tmp)
            half_written = crashed.file + ".tmp.pyxl"  # the next copy was being written when it crashed
            open(half_written, "w").close()
            self.assertEqual(now.orphans(), [])  # still beating: another PyExcel, not a crash
            self.assertTrue(os.path.exists(half_written))
            found = now.orphans(now=time.time() + 600)
            self.assertFalse(os.path.exists(half_written))
            self.assertEqual([e["path"] for e in found], ["D:/bang.xlsx"])
            self.assertEqual(fe.load_workbook(found[0]["file"]).get_raw("Sheet1", 0, 0), "chưa lưu")
            now.forget(found[0])
            crashed.discard()
            self.assertEqual(now.orphans(now=time.time() + 600), [])


def cp1258(text):
    """Text as Windows writes it in code page 1258: letters it has as they are, the others as a letter
    plus a separate tone mark."""
    import unicodedata
    out = b""
    for ch in text:
        try:
            out += ch.encode("cp1258")
        except UnicodeEncodeError:
            parts = unicodedata.normalize("NFD", ch)
            out += unicodedata.normalize("NFC", parts[:-1]).encode("cp1258") + parts[-1].encode("cp1258")
    return out


class TextImportTests(unittest.TestCase):
    VN = "Họ tên;Số tiền\nNguyễn Văn Ánh;1.234,5\nTrần Thị Bích Ngọc;-12,75\n"

    def test_encodings_found_like_excel(self):
        cases = [(b"\xef\xbb\xbf" + self.VN.encode("utf-8"), "utf-8"), (self.VN.encode("utf-8"), "utf-8"),
                 (self.VN.encode("utf-16"), "utf-16"), (self.VN.encode("utf-16-le"), "utf-16-le"),
                 (cp1258(self.VN), "cp1258")]
        for data, encoding in cases:
            self.assertEqual(fe.decode_text(data), (self.VN, encoding), encoding)
        self.assertEqual(fe.decode_text("Café;12\n".encode("cp1252"))[0], "Café;12\n")

    def test_japanese_windows(self):
        from unittest import mock
        japanese = "種類,ﾌﾞﾗｹｯﾄ,ﾃﾞｰﾀ,スプライス\n1,ﾌﾞﾚｰｽ,大梁,PL-16x570x100\n"
        with mock.patch.object(fe, "ansi_code_page", return_value="cp932"):
            self.assertEqual(fe.decode_text(japanese.encode("cp932")), (japanese, "cp932"))
            self.assertEqual(fe.decode_text(cp1258(self.VN)), (self.VN, "cp1258"))  # still Vietnamese
        self.assertEqual(fe.decode_text(japanese.encode("cp932"), "cp932")[0], japanese)

    def test_separate_tone_marks_are_composed(self):
        self.assertEqual(fe.decode_text("Nguyễn".encode("utf-8"))[0], "Nguyễn")

    def test_old_vietnamese_encodings(self):
        self.assertEqual(len(fe._TCVN3_FROM), len(fe._TCVN3_TO))
        self.assertEqual(len(set(fe._TCVN3_FROM)), len(fe._TCVN3_FROM))
        self.assertEqual(fe.decode_text(b"H\xb5 N\xe9i, Ti\xd5ng Vi\xd6t \xa7\xb7", "tcvn3")[0], "Hà Nội, Tiếng Việt Đã")
        self.assertEqual(fe.decode_text("Haø Noäi, Ñaø Naüng, Vieät Nam".encode("latin-1"), "vni")[0],
                         "Hà Nội, Đà Nẵng, Việt Nam")

    def test_delimiter_and_decimal_comma(self):
        rows, encoding, delimiter, comma = fe.read_text_table(cp1258(self.VN))
        self.assertEqual((encoding, delimiter, comma), ("cp1258", ";", True))
        self.assertEqual(rows[1:], [["Nguyễn Văn Ánh", "1234.5"], ["Trần Thị Bích Ngọc", "-12.75"]])
        rows = fe.read_text_table(self.VN.encode("utf-8"), decimal_comma=False)[0]
        self.assertEqual(rows[1][1], "1.234,5")
        rows, _, delimiter, comma = fe.read_text_table(b'x,y\n"1,234.5",2.5\n')
        self.assertEqual((rows[1], delimiter, comma), (["1,234.5", "2.5"], ",", False))
        rows, _, delimiter, _ = fe.read_text_table("Tên\tTuổi\nAn\t30\n".encode("utf-16"))
        self.assertEqual((rows, delimiter), ([["Tên", "Tuổi"], ["An", "30"]], "\t"))

    def test_load_csv(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "ansi.csv")
            with open(path, "wb") as f:
                f.write(cp1258(self.VN))
            wb = fe.load_workbook(path)
        self.assertEqual((val(wb, "A2"), val(wb, "B2"), val(wb, "B3")), ("Nguyễn Văn Ánh", 1234.5, -12.75))
        self.assertEqual(wb.import_info["encoding"], "cp1258")
        self.assertFalse(wb.import_info["truncated"])


class BigSheetTests(unittest.TestCase):
    def test_excel_size(self):
        self.assertEqual((fe.MAX_ROWS, fe.col_name(fe.MAX_COLS - 1)), (1048576, "XFD"))
        wb = make({"XFD1048576": "góc"})
        self.assertEqual(wb.used_extent("Sheet1"), (fe.MAX_ROWS, fe.MAX_COLS))

    def test_cell_store(self):
        cells = fe.CellStore({(0, 0): "a", (5, 2): "b"})
        self.assertEqual(cells.extent(), (6, 3))
        cells[(9, 1)] = "c"
        self.assertEqual(cells.extent(), (10, 3))
        del cells[(9, 1)]
        self.assertEqual(cells.extent(), (6, 3))
        copy = cells.frozen()
        self.assertIs(cells.frozen(), copy)  # unchanged: undo reuses the copy
        cells.pop((7, 7), None)
        self.assertIs(cells.frozen(), copy)
        cells[(0, 1)] = "x"
        self.assertIsNot(cells.frozen(), copy)
        self.assertNotIn((0, 1), copy)
        sh = fe.Sheet("S")
        sh.cells = {(1, 1): "y"}
        self.assertIsInstance(sh.cells, fe.CellStore)
        self.assertEqual(sh.cells.extent(), (2, 2))

    def test_bounds_and_memo(self):
        cells = fe.CellStore({(3, 2): "a", (7, 1): "b"})
        self.assertEqual((cells.first(), cells.extent()), ((3, 1), (8, 3)))
        cells[(1, 5)] = "c"
        self.assertEqual((cells.first(), cells.extent()), ((1, 1), (8, 6)))
        del cells[(1, 5)]
        self.assertEqual(cells.first(), (3, 1))
        formats = fe.TrackedDict({(0, 0): {"bold": True}})
        calls = []
        count = lambda d: calls.append(1) or len(d)  # noqa: E731
        self.assertEqual((formats.memo("n", count), formats.memo("n", count)), (1, 1))
        formats[(1, 1)] = {"bg": "#ffff00"}
        self.assertEqual((formats.memo("n", count), len(calls)), (2, 2))

    def test_snapshots_share_unchanged_cells(self):
        wb = make({"A1": "1", "B1": "2"})
        first = wb.snapshot()[0][0][1]
        second = wb.snapshot()[0][0][1]
        self.assertIs(first["cells"], second["cells"])
        wb.set_raw("Sheet1", 0, 0, "5")
        third = wb.snapshot()[0][0][1]
        self.assertIsNot(third["cells"], first["cells"])
        self.assertEqual(first["cells"][(0, 0)], "1")

    def test_write_cells(self):
        wb = make({"A1": "1", "C1": "old"})
        wb.write_cells([("Sheet1", 0, 0, "5", None), ("Sheet1", 0, 1, "=A1*2", {"bold": True}),
                        ("Sheet1", 0, 2, None, None)])
        self.assertEqual((val(wb, "B1"), wb.get_format("Sheet1", 0, 1)["bold"]), (10, True))
        self.assertEqual(wb.used_extent("Sheet1"), (1, 2))

    def test_text_table_cells(self):
        cells = fe.table_cells([["a", "", "b"], [], ["", "c"]], 2, 1)
        self.assertEqual(dict(cells), {(2, 1): "a", (2, 3): "b", (4, 2): "c"})


@unittest.skipUnless(fe.HAS_XLSX, "openpyxl not installed")
class FastXlsxTests(unittest.TestCase):
    """Big sheets are read and written by PyExcel's own XML code; it must give what openpyxl gives."""

    SHEET_DATA = (
        '<sheetData>'
        '<row r="1" spans="1:3" ht="30" customHeight="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c>'
        '<c r="C1" t="inlineStr"><is><t xml:space="preserve"> nội tuyến </t></is></c></row>'
        '<row r="2"><c r="A2"><v>1.5</v></c><c r="B2"><f t="shared" ref="B2:B4" si="0">A2*2</f><v>3</v></c>'
        '<c r="C2" t="b"><v>1</v></c></row>'
        '<row r="3" hidden="1"><c r="A3"><v>2</v></c><c r="B3"><f t="shared" si="0"/><v>4</v></c>'
        '<c r="C3" t="e"><v>#DIV/0!</v></c></row>'
        '<row r="4"><c r="A4"><v>3</v></c><c r="B4"><f t="shared" si="0"/><v>6</v></c>'
        '<c r="C4" t="str"><f>"a"&amp;"b"</f><v>ab</v></c><c r="D4" t="s"><v>0</v></c></row>'
        '</sheetData>')
    STRINGS = ('<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="3" uniqueCount="2">'
               '<si><t>Tên</t></si><si><r><t>đậm</t></r><r><rPr><b/></rPr><t xml:space="preserve"> thường</t></r>'
               '<rPh sb="0" eb="1"><t>x</t></rPh></si></sst>')

    @staticmethod
    def load_with_openpyxl(path):
        from unittest import mock
        with mock.patch.object(fe, "XlsxCells", side_effect=ValueError):
            return fe.load_workbook(path)

    def assertSameSheets(self, a, b):
        self.assertEqual(list(a.sheets), list(b.sheets))
        for name in a.sheets:
            for key in fe.Sheet._STATE:
                self.assertEqual(getattr(a.sheets[name], key), getattr(b.sheets[name], key), f"{name}.{key}")

    def test_reads_excel_cell_xml(self):
        import re
        import zipfile
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            plain, path = os.path.join(tmp, "plain.xlsx"), os.path.join(tmp, "excel.xlsx")
            book = openpyxl.Workbook()
            book.active["A1"] = "x"
            book.save(plain)
            with zipfile.ZipFile(plain) as src, zipfile.ZipFile(path, "w") as dst:  # as Excel writes it
                for item in src.infolist():
                    data = src.read(item.filename).decode("utf-8")
                    if item.filename == "xl/worksheets/sheet1.xml":
                        data = re.sub(r"<sheetData>.*</sheetData>|<sheetData/>", self.SHEET_DATA, data, flags=re.S)
                    elif item.filename == "xl/_rels/workbook.xml.rels":
                        data = data.replace("</Relationships>", '<Relationship Id="rId9" Target="sharedStrings.xml" Type='
                                            '"http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                                            'sharedStrings"/></Relationships>')
                    elif item.filename == "[Content_Types].xml":
                        data = data.replace("</Types>", '<Override PartName="/xl/sharedStrings.xml" ContentType="appl'
                                            'ication/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"'
                                            '/></Types>')
                    dst.writestr(item, data)
                dst.writestr("xl/sharedStrings.xml", self.STRINGS)
            fast, slow = fe.load_workbook(path), self.load_with_openpyxl(path)
        self.assertSameSheets(fast, slow)
        cells = fast.sheets["Sheet"].cells
        self.assertEqual([cells.get((0, c)) for c in range(3)], ["Tên", "đậm thường", " nội tuyến "])
        self.assertEqual([cells.get((r, 1)) for r in range(1, 4)], ["=A2*2", "=A3*2", "=A4*2"])
        self.assertEqual([cells.get((r, 2)) for r in range(1, 4)], ["TRUE", "#DIV/0!", '="a"&"b"'])
        self.assertEqual((fast.sheets["Sheet"].row_heights, fast.sheets["Sheet"].hidden_rows), ({0: 40}, {2}))

    def test_writes_like_openpyxl(self):
        from unittest import mock
        wb = make({"A1": "Tên", "B1": "12.5", "C1": "=B1*2", "D1": "TRUE", "E1": "21/10/2026", "F1": "'007",
                   "A2": "dòng 1\ndòng 2", "B2": "a\x01b", "C2": "<&>", "D2": " cách ", "A3": "1e5", "B3": "#N/A"})
        wb.set_format("Sheet1", 0, 0, {"bold": True, "bg": "#ffff00", "num": "#,##0", "wrap": True})
        wb.set_format("Sheet1", 2, 0, {"num": "@"})
        wb.set_format("Sheet1", 5, 5, {"bt": ["thin", "#000000"]})
        sh = wb.sheets["Sheet1"]
        sh.row_heights, sh.custom_rows, sh.hidden_rows = {1: 50}, {1}, {2}
        sh.row_levels = {3: 1, 4: 2}
        with tempfile.TemporaryDirectory() as tmp:
            slow_path, fast_path = os.path.join(tmp, "slow.xlsx"), os.path.join(tmp, "fast.xlsx")
            fe.save_workbook(wb, slow_path)
            with mock.patch.object(fe, "FAST_SAVE_CELLS", 0):
                fe.save_workbook(wb, fast_path)
            slow, fast = fe.load_workbook(slow_path), fe.load_workbook(fast_path)
            self.assertSameSheets(fast, slow)
            other = self.load_with_openpyxl(fast_path)
        self.assertEqual(fast.get_raw("Sheet1", 1, 1), "a\x01b")
        for wb_ in (fast, other):  # openpyxl leaves Excel's escape for control characters (_x0001_) as it is
            wb_.sheets["Sheet1"].cells.pop((1, 1))
        self.assertSameSheets(fast, other)
        self.assertEqual(val(fast, "C1"), 25)
        self.assertEqual(fast.sheets["Sheet1"].hidden_rows, {2})

@unittest.skipUnless(importlib.util.find_spec("tkinter"), "tkinter not installed")
class UndoNameTests(unittest.TestCase):
    """The names the Undo button's list shows for each step."""

    def test_names_of_cell_changes(self):
        import excel_app
        name = excel_app.SpreadsheetApp._step_name
        sel = ("Sheet1", (0, 0), (0, 0))
        typed = ("cells", [(("Sheet1", 0, 0), None, "Hà Nội", None, None)], sel)
        self.assertEqual(name(typed), 'Nhập "Hà Nội" vào A1')
        long_text = ("cells", [(("Sheet1", 2, 1), None, "x" * 40, None, None)], sel)
        self.assertEqual(name(long_text), f'Nhập "{"x" * 23}…" vào B3')
        cleared = ("cells", [(("Sheet1", 0, 0), "1", None, None, None), (("Sheet1", 1, 0), "2", None, None, None)], sel)
        self.assertEqual(name(cleared), "Xóa nội dung")
        bold = ("cells", [(("Sheet1", 0, 0), "1", "1", None, {"bold": True})], sel)
        self.assertEqual(name(bold), "In đậm")
        border = ("cells", [(("Sheet1", 0, 0), None, None, {"bt": ["thin", "#000000"]}, {"bb": ["thin", "#000000"]})], sel)
        self.assertEqual(name(border), "Viền")
        mixed = ("cells", [(("Sheet1", 0, 0), None, None, None, {"bold": True, "bg": "#ffff00"})], sel)
        self.assertEqual(name(mixed), "Định dạng ô")

    def test_outer_command_names_the_step(self):
        import excel_app

        class Commands:
            _undo_name = None

            def __init__(self):
                self.seen = []

            @excel_app.undo_name("Chèn hàng")
            def insert_rows(self):
                self.finish_typing()
                self.seen.append(self._undo_name)
                self.format_cells({"bold": True})

            @excel_app.undo_name(lambda self, update: excel_app.format_name(update))
            def format_cells(self, update):
                self.seen.append(self._undo_name)

            @excel_app.content_named
            def finish_typing(self):
                self.seen.append(self._undo_name)

        cmd = Commands()
        cmd.insert_rows()
        self.assertEqual(cmd.seen, [None, "Chèn hàng", "Chèn hàng"])
        cmd.seen.clear()
        cmd.format_cells({"bold": True})
        self.assertEqual((cmd.seen, cmd._undo_name), (["In đậm"], None))

class ScreenTests(unittest.TestCase):
    """Popups and dialogs stay on the monitor PyExcel is on (Tk on Windows only knows the primary one)."""

    MONITORS = [(-1920, 0, 1920, 1040), (0, 0, 1920, 1040), (1920, -200, 2560, 1400)]

    def area(self, widget, point=None):
        x, y = point
        return next((m for m in self.MONITORS if m[0] <= x < m[0] + m[2] and m[1] <= y < m[1] + m[3]),
                    self.MONITORS[1])

    def test_popups_stay_on_their_monitor(self):
        from unittest import mock
        import screens
        with mock.patch.object(screens, "work_area", self.area):
            self.assertEqual(screens.keep_on_screen(None, 3000, 500, 300, 200, flip=30), (3000, 500))
            self.assertEqual(screens.keep_on_screen(None, 4400, 500, 300, 200, flip=30), (4176, 500))  # right edge
            self.assertEqual(screens.keep_on_screen(None, -1500, 100, 300, 200), (-1500, 100))  # left monitor
            self.assertEqual(screens.keep_on_screen(None, -1500, 950, 300, 200, flip=30), (-1500, 720))  # opens above
            self.assertEqual(screens.keep_on_screen(None, 100, 1000, 300, 200), (100, 840))

class ThreeDReferenceTests(unittest.TestCase):
    """Sheet1:Sheet3!A1 - the same cells on every sheet from the first to the last, as in Excel."""

    def build(self):
        wb = Workbook(("Tổng", "Tháng 1", "Tháng 2", "Tháng 3"))
        for i, name in enumerate(["Tháng 1", "Tháng 2", "Tháng 3"], 1):
            make({"A1": str(i * 10), "A2": str(i)}, sheet=name, wb=wb)
        return wb

    def test_functions(self):
        wb = self.build()
        for formula, expected in (("=SUM('Tháng 1:Tháng 3'!A1)", 60), ("=SUM('Tháng 1:Tháng 3'!A1:A2)", 66),
                                  ("=AVERAGE('Tháng 1:Tháng 3'!A1)", 20), ("=MAX('Tháng 3:Tháng 1'!A1:A2)", 30),
                                  ("=COUNT('Tháng 1:Tháng 2'!A1:A5)", 4), ("=SUM('Tháng 2:Tháng 2'!A1:A50000)", 22)):
            make({"B2": formula}, sheet="Tổng", wb=wb)
            self.assertEqual(val(wb, "B2", "Tổng"), expected, formula)

    def test_errors_and_unquoted_names(self):
        wb = self.build()
        make({"B1": "='Tháng 1:Tháng 3'!A1", "B2": "=SUM('Tháng 1:Không có'!A1)"}, sheet="Tổng", wb=wb)
        self.assertEqual((val(wb, "B1", "Tổng").code, val(wb, "B2", "Tổng").code), ("#VALUE!", "#REF!"))
        wb = Workbook(("BH", "形鋼", "ボルト"))
        for name in wb.sheets:
            make({"Q5": "2"}, sheet=name, wb=wb)
        make({"U3": "=SUM(BH:ボルト!Q5:Q50000)"}, sheet="BH", wb=wb)
        self.assertEqual(val(wb, "U3", "BH"), 6)
        self.assertEqual(fe.shift_formula("=SUM(BH:ボルト!Q5:Q50000)", 1, 0), "=SUM(BH:ボルト!Q6:Q50001)")

if __name__ == "__main__":
    unittest.main()
